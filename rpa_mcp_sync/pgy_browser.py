from __future__ import annotations

import csv
import json
import re
import subprocess
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

import requests

from .config_store import ROOT
from .creator_store import _threshold_from_text, sanitize_creator_type

CDP_URL = "http://127.0.0.1:9222/json/version"
PGY_KOL_URL = "https://pgy.xiaohongshu.com/solar/pre-trade/note/kol"
PGY_ALL_NON_LIVE_METRICS = "全部非直播指标"
PGY_DETAIL_SCREENSHOT_DIR = ROOT / "runtime" / "pgy_detail_screenshots"
PGY_BLOGGER_CATEGORY_TAXONOMY_PATH = ROOT / "config" / "pgy_blogger_category_taxonomy.json"
PGY_KOL_API_TEMPLATE_PATH = ROOT / "runtime" / "pgy_kol_api_template.json"
PGY_CHROME_PROFILE_DIR = ROOT / "runtime" / "chrome-pgy-profile"
PGY_DETAIL_COLLECT_CONCURRENCY = 3

PGY_DISPLAY_METRICS = [
    PGY_ALL_NON_LIVE_METRICS,
]

PGY_REQUIRED_NON_LIVE_METRICS = [
    "粉丝数",
    "粉丝量变化幅度",
    "活跃粉丝占比",
    "互动粉丝占比",
    "曝光中位数（日常）",
    "阅读中位数（日常）",
    "互动中位数（日常）",
    "图文阅读中位数（日常）",
    "视频阅读中位数（日常）",
    "曝光中位数（合作）",
    "阅读中位数（合作）",
    "互动中位数（合作）",
    "外溢进店单价",
    "图文预估阅读单价",
    "图文预估互动单价",
    "视频预估阅读单价",
    "视频预估互动单价",
    "邀约48h回复率",
]


def _now_text() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _duration_text(seconds: float | int | None) -> str:
    if seconds is None:
        return ""
    safe_seconds = max(0.0, float(seconds))
    if safe_seconds < 60:
        return f"{safe_seconds:.1f}秒" if safe_seconds < 10 else f"{round(safe_seconds)}秒"
    minutes = int(safe_seconds // 60)
    remainder = int(round(safe_seconds % 60))
    if remainder >= 60:
        minutes += 1
        remainder = 0
    return f"{minutes}分{remainder:02d}秒"


def _elapsed_timing(started_at: str, started_perf: float) -> dict[str, Any]:
    duration_seconds = round(max(0.0, time.perf_counter() - started_perf), 2)
    return {
        "started_at": started_at,
        "finished_at": _now_text(),
        "duration_seconds": duration_seconds,
        "duration_text": _duration_text(duration_seconds),
    }


PGY_EXPORT_FIELD_ALIASES = {
    "nickname": ["达人昵称", "博主昵称", "昵称", "博主名称"],
    "pgy_url": ["蒲公英链接", "详情页链接", "主页链接", "博主链接"],
    "xiaohongshu_id": ["小红书号", "小红书ID", "小红书id"],
    "followers_count": ["粉丝数", "粉丝量"],
    "quote_price": ["全部报价", "报价", "图文报价", "图文笔记一口价"],
    "video_quote_price": ["视频报价", "视频笔记一口价"],
    "liked_collected_count": ["获赞与收藏", "赞藏数", "赞藏量"],
    "female_fans_ratio": ["粉丝女性用户占比", "女性粉丝占比"],
    "male_fans_ratio": ["粉丝男性用户占比", "男性粉丝占比"],
    "fans_18_24_ratio": ["粉丝年龄18-24占比", "18-24占比"],
    "fans_25_34_ratio": ["粉丝年龄25-34占比", "25-34占比"],
    "fans_35_44_ratio": ["粉丝年龄35-44占比", "35-44占比"],
    "fans_44_plus_ratio": ["粉丝年龄44岁以上占比", "44岁以上占比"],
    "active_fans_ratio": ["活跃粉丝占比"],
    "fans_growth_ratio": ["粉丝量变化幅度"],
    "read_fans_ratio": ["阅读粉丝占比"],
    "interaction_fans_ratio": ["互动粉丝占比"],
    "order_fans_ratio": ["下单粉丝占比"],
    "daily_exposure_median": ["曝光中位数（日常）", "日常曝光中位数"],
    "daily_read_median": ["阅读中位数（日常）", "日常阅读中位数"],
    "daily_interaction_median": ["互动中位数（日常）", "日常互动中位数"],
    "daily_thousand_like_note_ratio": ["千赞笔记比例", "千赞笔记比例（日常）"],
    "daily_hundred_like_note_ratio": ["百赞笔记比例", "百赞笔记比例（日常）"],
    "image_daily_exposure_median": ["图文曝光中位数（日常）"],
    "image_daily_read_median": ["图文阅读中位数（日常）"],
    "image_daily_interaction_median": ["图文互动中位数（日常）"],
    "image_daily_thousand_like_note_ratio": ["图文千赞笔记比例"],
    "image_daily_hundred_like_note_ratio": ["图文百赞笔记比例"],
    "video_daily_exposure_median": ["视频曝光中位数（日常）"],
    "video_daily_read_median": ["视频阅读中位数（日常）"],
    "video_daily_interaction_median": ["视频互动中位数（日常）"],
    "video_daily_thousand_like_note_ratio": ["视频千赞笔记比例"],
    "video_daily_hundred_like_note_ratio": ["视频百赞笔记比例"],
    "video_completion_rate": ["视频完播率"],
    "cooperation_exposure_median": ["曝光中位数（合作）", "合作曝光中位数"],
    "cooperation_read_median": ["阅读中位数（合作）", "合作阅读中位数"],
    "cooperation_interaction_median": ["互动中位数（合作）", "合作互动中位数"],
    "overflow_store_median": ["外溢进店中位数"],
    "overflow_store_unit_price": ["外溢进店单价"],
    "image_cpm": ["图文预估CPM价格", "预估图文CPM"],
    "image_read_unit_price": ["图文预估阅读单价", "图文笔记阅读单价"],
    "image_interaction_unit_price": ["图文预估互动单价", "图文笔记互动单价"],
    "video_cpm": ["视频预估CPM价格", "预估视频CPM"],
    "video_read_unit_price": ["视频预估阅读单价", "视频笔记阅读单价"],
    "video_interaction_unit_price": ["视频预估互动单价", "视频笔记互动单价"],
    "reply_rate_48h": ["邀约48h回复率", "邀约48小时回复率"],
    "live_30d_count": ["近30天直播场次"],
    "live_avg_viewers": ["场均观看人数", "场均观播人数"],
    "live_avg_sales": ["场均销售额"],
}

PGY_TABLE_TO_PAYLOAD_FIELDS = {
    "粉丝数": "followers_count",
    "全部报价": "quote_price",
    "图文笔记一口价": "quote_price",
    "视频报价": "video_quote_price",
    "视频笔记一口价": "video_quote_price",
    "获赞与收藏": "liked_collected_count",
    "赞藏数": "liked_collected_count",
    "赞藏量": "liked_collected_count",
    "粉丝女性用户占比": "female_fans_ratio",
    "女性粉丝占比": "female_fans_ratio",
    "粉丝年龄25-34占比": "fans_25_34_ratio",
    "活跃粉丝占比": "active_fans_ratio",
    "粉丝量变化幅度": "fans_growth_ratio",
    "阅读粉丝占比": "read_fans_ratio",
    "互动粉丝占比": "interaction_fans_ratio",
    "下单粉丝占比": "order_fans_ratio",
    "曝光中位数（日常）": "daily_exposure_median",
    "阅读中位数（日常）": "daily_read_median",
    "互动中位数（日常）": "daily_interaction_median",
    "千赞笔记比例": "daily_thousand_like_note_ratio",
    "千赞笔记比例（日常）": "daily_thousand_like_note_ratio",
    "百赞笔记比例": "daily_hundred_like_note_ratio",
    "百赞笔记比例（日常）": "daily_hundred_like_note_ratio",
    "图文曝光中位数（日常）": "image_daily_exposure_median",
    "图文阅读中位数（日常）": "image_daily_read_median",
    "图文互动中位数（日常）": "image_daily_interaction_median",
    "图文千赞笔记比例": "image_daily_thousand_like_note_ratio",
    "图文百赞笔记比例": "image_daily_hundred_like_note_ratio",
    "视频曝光中位数（日常）": "video_daily_exposure_median",
    "视频阅读中位数（日常）": "video_daily_read_median",
    "视频互动中位数（日常）": "video_daily_interaction_median",
    "视频千赞笔记比例": "video_daily_thousand_like_note_ratio",
    "视频百赞笔记比例": "video_daily_hundred_like_note_ratio",
    "视频完播率": "video_completion_rate",
    "曝光中位数（合作）": "cooperation_exposure_median",
    "阅读中位数（合作）": "cooperation_read_median",
    "互动中位数（合作）": "cooperation_interaction_median",
    "外溢进店中位数": "overflow_store_median",
    "外溢进店单价": "overflow_store_unit_price",
    "图文预估CPM价格": "image_cpm",
    "预估图文CPM": "image_cpm",
    "图文预估阅读单价": "image_read_unit_price",
    "图文笔记阅读单价": "image_read_unit_price",
    "图文预估互动单价": "image_interaction_unit_price",
    "图文笔记互动单价": "image_interaction_unit_price",
    "视频预估CPM价格": "video_cpm",
    "预估视频CPM": "video_cpm",
    "视频预估阅读单价": "video_read_unit_price",
    "视频笔记阅读单价": "video_read_unit_price",
    "视频预估互动单价": "video_interaction_unit_price",
    "视频笔记互动单价": "video_interaction_unit_price",
    "邀约48h回复率": "reply_rate_48h",
    "邀约48小时回复率": "reply_rate_48h",
    "近30天直播场次": "live_30d_count",
    "场均观播人数": "live_avg_viewers",
    "场均观看人数": "live_avg_viewers",
    "场均销售额": "live_avg_sales",
}

PGY_DETAIL_CATEGORY_STOP_LINES = {
    "粉丝数",
    "获赞与收藏",
    "收藏",
    "邀约",
    "合作报价",
    "图文笔记一口价",
    "视频笔记一口价",
}

PGY_NON_LIVE_ROW_METRIC_FIELDS = [
    ("followers_count", "粉丝数"),
    ("fans_growth_ratio", "粉丝量变化幅度"),
    ("active_fans_ratio", "活跃粉丝占比"),
    ("interaction_fans_ratio", "互动粉丝占比"),
    ("daily_exposure_median", "曝光中位数（日常）"),
    ("daily_read_median", "阅读中位数（日常）"),
    ("daily_interaction_median", "互动中位数（日常）"),
    ("daily_thousand_like_note_ratio", "千赞笔记比例"),
    ("daily_hundred_like_note_ratio", "百赞笔记比例"),
    ("image_daily_exposure_median", "图文曝光中位数（日常）"),
    ("image_daily_read_median", "图文阅读中位数（日常）"),
    ("image_daily_interaction_median", "图文互动中位数（日常）"),
    ("image_daily_thousand_like_note_ratio", "图文千赞笔记比例"),
    ("image_daily_hundred_like_note_ratio", "图文百赞笔记比例"),
    ("video_daily_exposure_median", "视频曝光中位数（日常）"),
    ("video_daily_read_median", "视频阅读中位数（日常）"),
    ("video_daily_interaction_median", "视频互动中位数（日常）"),
    ("video_daily_thousand_like_note_ratio", "视频千赞笔记比例"),
    ("video_daily_hundred_like_note_ratio", "视频百赞笔记比例"),
    ("video_completion_rate", "视频完播率"),
    ("cooperation_exposure_median", "曝光中位数（合作）"),
    ("cooperation_read_median", "阅读中位数（合作）"),
    ("cooperation_interaction_median", "互动中位数（合作）"),
    ("overflow_store_median", "外溢进店中位数"),
    ("overflow_store_unit_price", "外溢进店单价"),
    ("image_cpm", "图文预估CPM价格"),
    ("image_read_unit_price", "图文预估阅读单价"),
    ("image_interaction_unit_price", "图文预估互动单价"),
    ("video_cpm", "视频预估CPM价格"),
    ("video_read_unit_price", "视频预估阅读单价"),
    ("video_interaction_unit_price", "视频预估互动单价"),
    ("reply_rate_48h", "邀约48h回复率"),
]

PGY_ROW_METRIC_FIELDS_WITH_LIVE = [
    *PGY_NON_LIVE_ROW_METRIC_FIELDS[:4],
    ("live_avg_viewers", "场均观播人数"),
    ("live_avg_sales", "场均销售额"),
    *PGY_NON_LIVE_ROW_METRIC_FIELDS[4:],
]


PGY_FILTER_ALIASES = {
    "北京/上海优先": ["北京", "上海"],
    "35岁以上优先": ["35-44岁", "35岁以上", "35-44"],
    "35岁以上≥40%": ["35-44岁", "35岁以上", "35-44"],
    "报价≤20000": ["1w-2w", "1万-2万", "2万以下", "≤2万"],
    "CPC<2/CPE<20": ["预估阅读单价", "预估互动单价"],
    "预估阅读/互动单价": ["预估阅读单价", "预估互动单价"],
}

PGY_METRIC_ALIASES = {
    "曝光中位数": ["曝光中位数（日常）", "曝光中位数"],
    "阅读中位数": ["阅读中位数（日常）", "阅读中位数"],
    "互动中位数": ["互动中位数（日常）", "互动中位数"],
    "全部报价": ["全部报价"],
}

PGY_COUNT_RANGE_OPTIONS = ["100万以上", "50万～100万", "10万～50万", "1万～10万", "0.5万～1万", "0.1万～0.5万"]
PGY_NOTE_COUNT_RANGE_OPTIONS = ["5万以上", "1万～5万", "0.5万～1万", "0.1万～0.5万"]
PGY_INTERACTION_RANGE_OPTIONS = ["2000以上", "1000～2000", "500～1000", "200～500", "100～200"]
PGY_RATE_RANGE_OPTIONS = ["40%以上", "30%～40%", "20%～30%", "10%～20%", "10%以下"]
PGY_PRICE_RANGE_OPTIONS = ["5万及以上", "1万～5万", "0.5万～1万", "0.1万～0.5万", "0.1万以下"]
PGY_UNIT_PRICE_OPTIONS = ["0.5以下", "0.5～1.0", "1.0～1.5", "1.5～2.0", "2.0以上"]
PGY_FAMILY_IDENTITY_GROUPS = [
    {"label": "家庭角色", "options": ["妈妈", "萌娃", "爸爸", "奶奶"]},
    {"label": "出镜人关系", "options": ["情侣", "夫妻", "家庭", "闺蜜", "兄弟"]},
    {"label": "母婴阶段", "options": ["备孕中", "孕期中", "0-6个月", "6-12个月", "1-3岁", "3-6岁", "6-12岁", "12岁以上"]},
]
PGY_CAREER_IDENTITY_GROUPS = [
    {"label": "传统行业", "options": ["工程师", "销售", "HR"]},
    {"label": "互联网", "options": ["主播", "运营", "产品经理", "程序员"]},
    {"label": "教育科研", "options": ["学生"]},
    {"label": "金融法律", "options": ["金融从业者"]},
    {"label": "企业创业", "options": ["创业者", "品牌创始人", "公益人"]},
    {"label": "时尚美妆", "options": ["模特", "化妆师", "造型师", "服装设计师", "珠宝设计师", "发型设计师"]},
    {"label": "食品饮料", "options": ["甜点师", "厨师", "咖啡师", "调酒师"]},
    {"label": "文化传媒", "options": ["编辑", "记者", "翻译", "作家", "娱评人", "影评人", "乐评人"]},
    {"label": "医疗健康", "options": ["营养师", "医生", "康复师"]},
    {"label": "艺术设计", "options": ["摄影师", "插画师", "室内设计师", "画家", "平面设计师", "建筑设计师", "非遗传承人", "涂鸦艺术家", "数字艺术家"]},
    {"label": "影视娱乐", "options": ["主持人", "导演", "制片人", "编剧", "经纪人", "真人秀嘉宾", "虚拟偶像", "rapper"]},
    {"label": "运动健身", "options": ["教练", "运动员", "舞蹈老师"]},
    {"label": "专业服务", "options": ["空乘", "花艺师", "整理师", "民宿主", "育婴师"]},
]
PGY_SPECIAL_BACKGROUND_GROUPS = [
    {"label": "生活背景", "options": ["留学背景", "海外华人", "铲屎官", "孕妈", "独居人群", "外国人", "混血儿"]},
    {"label": "备考经验", "options": ["考公过来人", "考研过来人", "法考过来人", "注会过来人"]},
    {"label": "兴趣爱好", "options": ["户外爱好者", "数码爱好者", "手账爱好者", "二次元人群", "汉服爱好者", "手办爱好者", "模型爱好者", "街舞爱好者", "骑行爱好者", "飞盘爱好者", "书法爱好者"]},
]
PGY_AUDIENCE_20_GROUPS = [
    {"label": "自在户外", "options": ["挑战极限者", "野趣探索家", "短逃离自愈派", "心灵远行客", "户外显眼包", "户外欢聚团"]},
    {"label": "自由畅行", "options": ["都市漫游家", "静奢新贵", "爆改浓人", "出行精算师"]},
    {"label": "运动焕活", "options": ["轻松健体派", "线条雕塑家", "寻乐运动派", "好动局内人", "自我超越者", "身心觉察师"]},
    {"label": "孕育学习", "options": ["科研育儿党", "松驰爸妈", "友伴式父母", "积进式父母", "好孕预备役", "稳孕选手"]},
    {"label": "娱乐放松", "options": ["放松乐子人", "沉浸式“戏”迷", "娱乐交友派", "真爱忠粉"]},
    {"label": "优奢享法", "options": ["奢派生活家", "悦己摘星人", "潮奢风格家", "静奢知识分子", "奢品入门人", "奢交体面人"]},
    {"label": "养身韧体", "options": ["爆肝青年", "高能青年", "娇宠彼得潘", "高消耗中年", "稳定守成中年", "探索人生的中年玩家", "熟龄悦己中年", "活力夕阳红"]},
    {"label": "虚拟人生", "options": ["审美收藏控", "高能“偷闲”客", "沉浸式畅“游”人", "竞技大神", "通关小机灵", "联结小“玩伴”", "“游”文化信徒"]},
    {"label": "美力加成", "options": ["美养佳人", "风格日抛党", "精养奢美族", "变美练习生", "气场精英", "美研尖子生"]},
    {"label": "心灵奇旅", "options": ["亲密学习父母", "恋爱修炼家", "实用信徒", "野生玄学家", "精进修心客"]},
    {"label": "文艺沉浸", "options": ["情绪捕手", "美学鉴赏家", "规律钻研党", "热门玩家", "世界狂想家"]},
    {"label": "数智未来", "options": ["效能领航员", "灵感创想客", "未来原住民", "品质感官控", "数码时髦精"]},
    {"label": "舌尖盛宴", "options": ["好味饕客", "精算稳妥人", "食饮养生族", "吃喝欢聚派", "逐潮尝新客", "拓圈商务客", "专味信徒", "“怪味”猎手", "“乐养”零食客", "囤粮“小馋猫”", "生活“增味”家", "“纵情”高压党"]},
    {"label": "看世界", "options": ["轻松舒心派", "热门追踪党", "同心群游党", "求索漫旅人", "野地探险家", "圣地巡礼者", "追爱忠粉", "山水避世客"]},
    {"label": "家有萌宠", "options": ["自然“动物学家”", "同行伙伴", "宠溺“爸妈”", "爱宠观赏派", "流浪动物保护党"]},
    {"label": "家生活", "options": ["游牧青年", "筑巢青年", "全能生活家", "居家策展人"]},
    {"label": "发现附近", "options": ["下楼享受派", "社区玩咖", "市井“街溜子”", "圈层专研人", "“速联”社交狂", "举家“撒欢”党", "城郊出走族"]},
    {"label": "成长进阶", "options": ["争渡“上岸”人", "资格证“卷王”", "进阶专业精英", "职场闯关人", "兴趣研学家", "精英培优家", "“社会人”教练", "因材施教师", "尽责陪练员"]},
    {"label": "时尚态度", "options": ["追新之乐", "弄潮先锋", "IP狂人", "街头潮客", "三坑玩家", "质感男士", "社会新鲜人", "气场大女主"]},
]
PGY_SKILLED_CONTENT_GROUPS = [
    {"label": "形式", "options": ["vlog", "探店", "测评", "ootd", "合集", "plog", "开箱", "教程", "成分解析", "彩妆试色", "仿妆", "沉浸式"]},
    {"label": "风格", "options": ["韩系", "日系", "欧美风", "氛围感", "纯欲", "甜酷", "复古", "高级感", "校园风", "中性风"]},
    {"label": "生活方式", "options": ["职场生活", "自律生活", "露营徒步", "极简主义", "低脂低卡"]},
    {"label": "肤质肤色", "options": ["油皮", "干皮", "混合肌", "敏感肌", "痘痘肌", "瑕疵皮", "白皮", "黄皮"]},
    {"label": "皮肤养护", "options": ["保湿补水", "美白", "淡斑", "祛黄", "抗氧化", "抗老", "祛皱", "抗炎", "修复", "祛痘祛闭口", "隔离防晒", "控油", "眼部护理"]},
]
PGY_CONTENT_SUBJECT_GROUPS = [
    {"label": "汽车特色", "options": ["沉浸式开车", "汽车美图"]},
    {"label": "通用", "options": ["大字报", "干货分享", "街头采访", "口播", "长文", "知识科普", "变装", "访谈", "梗图", "好物分享", "幽默搞笑", "挑战"]},
]
PGY_INDUSTRY_PORTRAIT_OPTIONS = ["家居家装", "日化家清", "多行业适用", "教育培训", "母婴", "汽车出行", "服饰鞋包", "美妆个护", "出行旅游", "珠宝配饰", "文玩娱乐", "奢侈品", "食品饮料", "到店综合", "宠物", "3C数码", "互联网", "家用电器", "运动户外", "本地生活", "行业通用人群", "家具", "家居百货", "灯饰光源", "装修设计与工程服务", "家居建材零售", "智能家居", "家装主材", "场景", "风格", "产品", "卧室兴趣人群", "餐厅兴趣人群", "客厅场景人群", "阳台兴趣用户", "厨房兴趣人群", "儿童房兴趣人群", "自我充电卧室", "多边形卧室", "高敏感卧室", "客厅兴趣人群", "浴室兴趣人群", "玄关兴趣人群", "造型卧室"]
PGY_CONSUMPTION_BEHAVIOR_GROUPS = [
    {"label": "预估车主作者", "options": ["Porsche", "ORA", "MINI", "MAZDA", "BYD", "萤火虫", "一汽红旗", "一汽奥迪", "小鹏", "五菱", "蔚来", "特斯拉", "坦克", "斯巴鲁", "上汽大众", "梅赛德斯-奔驰", "路虎", "领克", "铃木", "理想", "雷克萨斯", "兰博基尼", "捷途", "江铃福特", "极越", "极氪", "吉普", "吉利银河", "哈弗", "广汽丰田", "福特", "宾利", "宝马", "奥迪", "阿维塔"]},
]


def _extract_brands_after_labels(brief: str, labels: list[str]) -> list[str]:
    brands: list[str] = []
    stop_keywords = ["使用", "行业推荐", "推荐博主", "需要", "要求", "达人", "博主", "筛选", "人群目标", "找相似", "相似人群", "粉丝推荐"]
    for label in labels:
        for match in re.findall(rf"{label}[：: ]*([^\n。；;]{{2,120}})", brief):
            for part in re.split(r"[、,，/]+|\s+", match):
                brand = part.strip(" ：:，,、；;。")
                if not brand or any(keyword in brand for keyword in stop_keywords):
                    continue
                brands.append(brand)
    return list(dict.fromkeys(brands))

PGY_MARKETING_GOAL_GROUPS = [
    {"label": "曝光", "options": ["曝光表现", "阅读表现"]},
    {"label": "种草", "options": ["互动表现"]},
    {"label": "转化", "options": ["外溢进店表现"]},
]
PGY_MARKETING_GOAL_DEFAULT_METRIC = {
    "曝光": "曝光表现",
    "种草": "互动表现",
    "转化": "外溢进店表现",
}
PGY_MARKETING_GOAL_METRIC_PARENT = {
    metric: group["label"]
    for group in PGY_MARKETING_GOAL_GROUPS
    for metric in group["options"]
}


def _load_pgy_blogger_category_taxonomy() -> dict[str, list[str]]:
    try:
        payload = json.loads(PGY_BLOGGER_CATEGORY_TAXONOMY_PATH.read_text(encoding="utf-8"))
    except Exception:
        return {}
    result: dict[str, list[str]] = {}
    for item in payload.get("categories") or []:
        if not isinstance(item, dict):
            continue
        value = str(item.get("value") or "").strip()
        subcategories = [
            str(subcategory).strip()
            for subcategory in (item.get("subcategories") or [])
            if str(subcategory).strip()
        ]
        if value:
            result[value] = subcategories
    return result


PGY_BLOGGER_CATEGORY_TAXONOMY = _load_pgy_blogger_category_taxonomy()
PGY_BLOGGER_CATEGORY_OPTIONS = list(PGY_BLOGGER_CATEGORY_TAXONOMY) or [
    "美妆",
    "护肤",
    "个人护理",
    "母婴",
    "时尚",
    "美食",
    "家居家装",
    "影视综资讯",
    "运动健身",
    "宠物",
    "文化艺术",
    "兴趣爱好",
    "生活记录",
    "教育",
    "职场",
    "情感",
    "摄影",
    "游戏",
    "科技数码",
    "出行旅游",
    "音乐",
    "搞笑",
    "健康养生",
    "汽车",
    "婚嫁",
    "商业财经",
    "素材",
    "其他",
]

PGY_FILTER_CATALOG = [
    {
        "field": "营销目标",
        "control_type": "marketing_goal_metric",
        "parent_options": [group["label"] for group in PGY_MARKETING_GOAL_GROUPS],
        "option_groups": PGY_MARKETING_GOAL_GROUPS,
        "options": [value for group in PGY_MARKETING_GOAL_GROUPS for value in group["options"]],
        "notes": "先点击曝光/种草/转化父级，再在弹层内选择对应指标；低优先级筛选，不作为必备条件。",
    },
    {
        "field": "按博主粉丝推荐",
        "control_type": "brand_search_recommendation",
        "input_fields": ["brand_name_or_id", "competitor_brand_name_or_id"],
        "notes": "右上角“智能推荐博主/请选择您的合作品牌”入口；可搜索合作品牌，也可输入竞品/对标品牌，用该品牌粉丝画像推荐博主。",
    },
    {
        "field": "博主类目",
        "control_type": "tag_select_with_hover_subcategory",
        "options": PGY_BLOGGER_CATEGORY_OPTIONS,
        "option_groups": [
            {"label": value, "options": subcategories}
            for value, subcategories in PGY_BLOGGER_CATEGORY_TAXONOMY.items()
        ],
        "notes": "先选择主类目；有 sub_value 时悬停/展开主类目后选择对应二级类目。二级类目必须来自本地白名单。",
    },
    {
        "field": "家庭身份",
        "control_type": "checkbox_popover",
        "option_groups": PGY_FAMILY_IDENTITY_GROUPS,
        "options": [value for group in PGY_FAMILY_IDENTITY_GROUPS for value in group["options"]],
        "notes": "打开后是多选弹层，需要点确定。",
    },
    {
        "field": "职业身份",
        "control_type": "checkbox_popover",
        "option_groups": PGY_CAREER_IDENTITY_GROUPS,
        "options": [value for group in PGY_CAREER_IDENTITY_GROUPS for value in group["options"]],
        "notes": "打开后是多选弹层，需要点确定。",
    },
    {
        "field": "特色背景",
        "control_type": "checkbox_popover",
        "option_groups": PGY_SPECIAL_BACKGROUND_GROUPS,
        "options": [value for group in PGY_SPECIAL_BACKGROUND_GROUPS for value in group["options"]],
        "notes": "打开后是多选弹层，需要点确定。",
    },
    {
        "field": "性别",
        "control_type": "dropdown_single",
        "options": ["不限", "男", "女"],
        "notes": "博主性别，打开后选择单项。",
    },
    {
        "field": "地域",
        "control_type": "three_level_cascade_checkbox_popover",
        "levels": ["国家/地区", "省/直辖市", "城市/区"],
        "options": ["中国", "美国", "日本", "澳大利亚", "英国", "加拿大", "韩国", "法国", "德国", "新加坡", "其他"],
        "notes": "中国地域需先选国家，再选省/市，再按需选城市/区；外国国家/地区可在第一列直接勾选。支持多选。",
    },
    {
        "field": "二十大人群",
        "control_type": "checkbox_popover",
        "option_groups": PGY_AUDIENCE_20_GROUPS,
        "options": [value for group in PGY_AUDIENCE_20_GROUPS for value in group["options"]],
        "notes": "页面显示为“新 二十大人群”，打开后是两列级联多选，需要点确定。",
    },
    {
        "field": "行业特色画像",
        "control_type": "checkbox_popover",
        "options": PGY_INDUSTRY_PORTRAIT_OPTIONS,
        "notes": "页面显示为“新 行业特色画像”，打开后是行业级联多选；选项会随行业列继续展开。",
    },
    {
        "field": "预估消费行为",
        "control_type": "checkbox_popover",
        "option_groups": PGY_CONSUMPTION_BEHAVIOR_GROUPS,
        "options": [value for group in PGY_CONSUMPTION_BEHAVIOR_GROUPS for value in group["options"]],
        "notes": "页面显示为“新 预估消费行为”，当前可在汽车/预估车主作者下选择品牌。",
    },
    {
        "field": "签约情况",
        "control_type": "dropdown_single",
        "options": ["不限", "个人博主", "机构博主"],
        "notes": "打开后选择博主属性。",
    },
    {
        "field": "擅长内容",
        "control_type": "checkbox_popover",
        "option_groups": PGY_SKILLED_CONTENT_GROUPS,
        "options": [value for group in PGY_SKILLED_CONTENT_GROUPS for value in group["options"]],
        "notes": "打开后是多选弹层，需要点确定。",
    },
    {
        "field": "内容题材",
        "control_type": "checkbox_popover",
        "option_groups": PGY_CONTENT_SUBJECT_GROUPS,
        "options": [value for group in PGY_CONTENT_SUBJECT_GROUPS for value in group["options"]],
        "notes": "页面显示为“新 内容题材”，当前汽车类目下包含汽车特色与通用题材。",
    },
    {
        "field": "粉丝量",
        "control_type": "preset_or_number_range",
        "options": PGY_COUNT_RANGE_OPTIONS,
        "input_fields": ["min", "max"],
        "notes": "打开后可选预设粉丝量档位，也可填写自定义最小/最大值。",
    },
    {
        "field": "粉丝年龄",
        "control_type": "dropdown",
        "options": ["<18 占比高", "18～24 占比高", "25～34 占比高", "35～44 占比高", ">44 占比高"],
        "notes": "打开后是下拉项。",
    },
    {
        "field": "粉丝性别",
        "control_type": "dropdown",
        "options": ["男性占比高", "女性占比高"],
        "notes": "打开后是下拉项。",
    },
    {
        "field": "粉丝地域",
        "control_type": "three_level_cascade_checkbox_popover",
        "levels": ["国家/地区", "省/直辖市", "城市"],
        "options": ["中国", "美国", "日本", "澳大利亚", "英国", "加拿大", "韩国", "法国", "德国", "新加坡", "其他"],
        "notes": "打开后是国家/省/城市三级级联，可多选，需要点确定。",
    },
    {
        "field": "婚恋状态",
        "control_type": "dropdown_single",
        "options": ["不限", "未婚", "已婚"],
        "notes": "打开后选择单项。",
    },
    {
        "field": "消费水平",
        "control_type": "dropdown_single",
        "options": ["不限", "低消费", "中消费", "高消费"],
        "notes": "打开后选择单项。",
    },
    {
        "field": "母婴阶段",
        "control_type": "checkbox_popover",
        "options": ["备孕", "0-6月", "7-12月", "1-3岁", "4-6岁", "7-12岁", "孕早期", "孕晚期"],
        "notes": "打开后是多选弹层，需要点确定。",
    },
    {
        "field": "手机价格",
        "control_type": "checkbox_popover",
        "options": ["0-999", "1000-1999", "2000-2999", "3000-3999", "4000-4999", "5000-5999", "6000-6999", "7000-7999", "8000+"],
        "notes": "打开后是多选弹层，需要点确定。",
    },
    {
        "field": "手机品牌",
        "control_type": "checkbox_popover",
        "options": ["苹果", "华为", "OPPO", "VIVO", "荣耀", "小米", "一加", "魅族", "中兴", "联想"],
        "notes": "打开后是多选弹层，需要点确定。",
    },
    {
        "field": "曝光中位数",
        "control_type": "preset_or_number_range",
        "options": PGY_NOTE_COUNT_RANGE_OPTIONS,
        "input_fields": ["min", "max"],
        "notes": "打开后可选预设曝光档位，也可填写自定义区间。",
    },
    {
        "field": "阅读中位数",
        "control_type": "preset_or_number_range",
        "options": PGY_NOTE_COUNT_RANGE_OPTIONS,
        "input_fields": ["min", "max"],
        "notes": "打开后可选预设阅读档位，也可填写自定义区间。",
    },
    {
        "field": "互动中位数",
        "control_type": "preset_or_number_range",
        "options": PGY_INTERACTION_RANGE_OPTIONS,
        "input_fields": ["min", "max"],
        "notes": "打开后可选预设互动档位，也可填写自定义区间。",
    },
    {
        "field": "千赞笔记比例",
        "control_type": "preset_or_percent_range",
        "options": PGY_RATE_RANGE_OPTIONS,
        "input_fields": ["min_percent", "max_percent"],
        "notes": "打开后可选预设比例，也可填写百分比区间。",
    },
    {
        "field": "笔记类型",
        "control_type": "dropdown_single",
        "options": ["不限", "图文笔记为主", "视频笔记为主"],
        "notes": "打开后选择单项。",
    },
    {
        "field": "合作报价",
        "control_type": "subfield_preset_or_number_range",
        "sub_fields": ["图文笔记", "视频笔记"],
        "options": PGY_PRICE_RANGE_OPTIONS,
        "input_fields": ["min", "max"],
        "notes": "打开后分别选择图文/视频笔记，再在子下拉中选报价档位或填自定义区间。",
    },
    {
        "field": "合作信用度",
        "control_type": "subfield_preset_or_percent_range",
        "sub_fields": ["邀约48h回复率"],
        "input_fields": ["min_percent", "max_percent"],
        "notes": "打开后选择邀约48h回复率，再选预设或填百分比区间。",
    },
    {
        "field": "合作订单数",
        "control_type": "number_range",
        "input_fields": ["min", "max"],
        "notes": "打开后填写最小/最大订单数。",
    },
    {
        "field": "近期合作行业",
        "control_type": "dropdown",
        "options": ["美妆个护", "食品饮料", "母婴", "3c及电器", "日用百货", "服装配饰", "互联网", "生活服务", "家居建材", "汽车"],
        "notes": "打开后是近期合作行业下拉项。",
    },
    {
        "field": "近期合作品牌",
        "control_type": "searchable_multi_select_with_exclude",
        "input_fields": ["brands"],
        "min_items": 3,
        "options": ["剔除上述品牌已合作博主"],
        "notes": "打开后是可搜索品牌下拉，页面提示至少选择3个品牌，可勾选剔除上述品牌已合作博主。",
    },
    {
        "field": "传播规模",
        "control_type": "multi_subfield_preset_or_number_range",
        "sub_fields": ["曝光中位数", "阅读中位数", "互动中位数", "外溢进店中位数"],
        "input_fields": ["min", "max"],
        "notes": "打开后每个子字段都有下拉档位和自定义区间。",
    },
    {
        "field": "预估CPM",
        "control_type": "subfield_preset_or_number_range",
        "sub_fields": ["预估图文CPM", "预估视频CPM"],
        "input_fields": ["min", "max"],
        "notes": "打开后分别选择图文/视频CPM，再选预设或填自定义区间。",
    },
    {
        "field": "预估阅读单价",
        "control_type": "subfield_preset_or_number_range",
        "sub_fields": ["图文笔记阅读单价", "视频笔记阅读单价"],
        "options": PGY_UNIT_PRICE_OPTIONS,
        "input_fields": ["min", "max"],
        "notes": "打开后分别选择图文/视频阅读单价，再在子下拉中选档位或填自定义区间。",
    },
    {
        "field": "预估互动单价",
        "control_type": "subfield_preset_or_number_range",
        "sub_fields": ["图文笔记互动单价", "视频笔记互动单价"],
        "options": PGY_UNIT_PRICE_OPTIONS,
        "input_fields": ["min", "max"],
        "notes": "打开后分别选择图文/视频互动单价，再在子下拉中选档位或填自定义区间。",
    },
    {
        "field": "外溢进店单价",
        "control_type": "preset_or_number_range",
        "options": ["0.5以下", "0.5～1.0", "1.0～1.5", "1.5～2.5", "2.5～4.0"],
        "input_fields": ["min", "max"],
        "notes": "打开后可选预设档位，也可填写自定义区间。",
    },
    {
        "field": "近30天直播场次",
        "control_type": "checkbox_popover",
        "options": ["0次", "1～5次", "6～10次", "10次以上"],
        "notes": "打开后是多选档位，需要点确定。",
    },
    {
        "field": "场均观播人数",
        "control_type": "checkbox_popover",
        "options": ["0~5k", "5k~1w", "1w~10w", "10w~50w", "50w以上"],
        "notes": "打开后是多选档位，需要点确定。",
    },
    {
        "field": "场均销售额",
        "control_type": "checkbox_popover",
        "options": ["5千以下", "5千～1万", "1万～10万", "10万～50万", "50万～100万", "100万～200万", "200万～500万", "500万以上"],
        "notes": "打开后是多选档位，需要点确定。",
    },
    {
        "field": "平台推荐",
        "control_type": "checkbox",
        "options": ["明星", "优质博主", "新锐博主", "笔记+直播均可合作", "意向行业匹配"],
        "notes": "一层复选框，可直接勾选。",
    },
    {
        "field": "行业推荐博主",
        "control_type": "nested_select_popover",
        "sub_fields": ["我的行业"],
        "notes": "打开后仍有“我的行业 / 请选择”的子下拉选择，需要继续选择行业后确定。",
    },
    {
        "field": "常规剔除",
        "control_type": "checkbox",
        "options": ["剔除低活博主", "剔除掉粉博主", "剔除已合作博主", "剔除已邀约博主"],
        "notes": "一层复选框，可直接勾选。",
    },
]


def _number_from_text(value: str) -> float | None:
    text = str(value or "").strip()
    if not text or text in {"-", "--"}:
        return None
    match = re.search(r"-?\d[\d,，]*(?:\.\d+)?\s*(?:万|w|W)?", text)
    if not match:
        return None
    token = match.group(0).strip()
    multiplier = 10000 if re.search(r"(?:万|w)\s*$", token, flags=re.I) else 1
    text = re.sub(r"[^\d.\-]", "", token.replace(",", "").replace("，", ""))
    try:
        return float(text) * multiplier
    except ValueError:
        return None


def _looks_like_quote_text(value: Any) -> bool:
    text = _clean_text(value)
    if not text or text in {"-", "--"}:
        return False
    if "%" in text or "占比" in text or "比例" in text:
        return False
    if any(marker in text for marker in ["¥", "￥", "元", "万", "w", "W"]):
        return _number_from_text(text) is not None
    number = _number_from_text(text)
    return number is not None and number >= 1000


def _looks_like_count_text(value: Any) -> bool:
    text = _clean_text(value)
    if not text or text in {"-", "--"}:
        return False
    if "%" in text or "占比" in text or "比例" in text:
        return False
    number = _number_from_text(text)
    return number is not None and number >= 1000


def _format_filter_number(value: Any) -> str:
    if value in (None, ""):
        return ""
    try:
        number = float(str(value).replace(",", "").replace("，", ""))
        return str(int(number)) if number.is_integer() else f"{number:g}"
    except Exception:
        return str(value)


def _range_numbers_from_text(value: str) -> tuple[float | None, float | None]:
    text = str(value or "")
    text = text.split("：", 1)[-1]
    parts = [part.strip() for part in re.split(r"～|~|至|到|-", text) if part.strip()]
    if len(parts) >= 2:
        return _number_from_text(parts[0]), _number_from_text(parts[1])
    if re.search(r"(?:及以上|以上)$", text.strip()):
        return _number_from_text(text), None
    if re.search(r"(?:及以下|以下)$", text.strip()):
        return None, _number_from_text(text)
    max_match = re.search(r"[≤<]\s*([0-9][0-9,，]*(?:\.\d+)?\s*(?:万|w|W)?)", text)
    if max_match:
        return None, _number_from_text(max_match.group(1))
    min_match = re.search(r"[≥>]\s*([0-9][0-9,，]*(?:\.\d+)?\s*(?:万|w|W)?)", text)
    if min_match:
        return _number_from_text(min_match.group(1)), None
    return None, None


PGY_MIN_ONLY_RANGE_FIELDS = {"粉丝量", "曝光中位数", "阅读中位数", "互动中位数", "合作订单数"}
PGY_MIN_ONLY_SUBFIELD_RANGE_FIELDS = {"传播规模", "合作信用度"}
PGY_MIN_ONLY_SUBFIELDS = {"曝光中位数", "阅读中位数", "互动中位数", "外溢进店中位数", "邀约48h回复率"}
PGY_BOUNDED_RANGE_FIELDS = {"合作报价"}
PGY_MAX_ONLY_RANGE_FIELDS = {"外溢进店单价"}
PGY_MAX_ONLY_SUBFIELD_RANGE_FIELDS = {"预估阅读单价", "预估互动单价", "预估CPM"}


def _range_policy_for_item(item: dict[str, Any], sub_field: str = "") -> str:
    field = str(item.get("field") or "")
    explicit_policy = str(item.get("range_policy") or item.get("rangePolicy") or "").strip()
    if explicit_policy in {"min_only", "max_only", "bounded", "default"}:
        return explicit_policy
    if field in PGY_BOUNDED_RANGE_FIELDS:
        return "bounded"
    if field in PGY_MAX_ONLY_RANGE_FIELDS or field in PGY_MAX_ONLY_SUBFIELD_RANGE_FIELDS:
        return "max_only"
    if field in PGY_MIN_ONLY_RANGE_FIELDS:
        return "min_only"
    if field in PGY_MIN_ONLY_SUBFIELD_RANGE_FIELDS and sub_field in PGY_MIN_ONLY_SUBFIELDS:
        return "min_only"
    return "default"


def _apply_range_policy(item: dict[str, Any], min_value: Any, max_value: Any, sub_field: str = "") -> tuple[Any, Any]:
    policy = _range_policy_for_item(item, sub_field)
    if policy == "min_only":
        if min_value in (None, "") and max_value not in (None, ""):
            min_value = max_value
        return min_value, ""
    if policy == "max_only":
        if max_value in (None, "") and min_value not in (None, ""):
            max_value = min_value
        return "", max_value
    if policy == "bounded":
        return min_value, max_value
    return min_value, max_value


def _split_filter_values(value: str) -> list[str]:
    return [part.strip() for part in re.split(r"[;；\n]+", str(value or "")) if part.strip()]


def _subfield_names_from_item(item: dict[str, Any], default: str | list[str] = "图文笔记") -> list[str]:
    raw_values: list[Any] = []
    for key in ("sub_fields", "subFields"):
        value = item.get(key)
        if isinstance(value, list):
            raw_values.extend(value)
        elif value:
            raw_values.append(value)
    for key in ("sub_field", "subField"):
        if item.get(key):
            raw_values.append(item.get(key))
    names: list[str] = []
    for raw in raw_values:
        for part in re.split(r"[、,，/;；|]+", str(raw)):
            part = part.strip()
            if part:
                names.append(part)
    text = str(item.get("value") or "")
    if not names:
        if "图文笔记" in text:
            names.append("图文笔记")
        if "视频笔记" in text:
            names.append("视频笔记")
    if not names and default:
        if isinstance(default, list):
            names.extend(default)
        else:
            names.append(default)
    return list(dict.fromkeys(names))


def _range_for_subfield(item: dict[str, Any], sub_field: str) -> tuple[Any, Any]:
    ranges = item.get("sub_ranges") or item.get("subRanges") or {}
    if isinstance(ranges, dict):
        config = ranges.get(sub_field)
        if isinstance(config, dict):
            return _apply_range_policy(item, config.get("min", item.get("min", "")), config.get("max", item.get("max", "")), sub_field)
        if isinstance(config, (list, tuple)) and len(config) >= 2:
            return _apply_range_policy(item, config[0], config[1], sub_field)
    text = str(item.get("value") or "")
    for segment in _split_filter_values(text):
        if sub_field in segment:
            parsed_min, parsed_max = _range_numbers_from_text(segment)
            return _apply_range_policy(
                item,
                item.get("min", parsed_min if parsed_min is not None else 0),
                item.get("max", parsed_max if parsed_max is not None else ""),
                sub_field,
            )
    parsed_min, parsed_max = _range_numbers_from_text(text)
    return _apply_range_policy(
        item,
        item.get("min", parsed_min if parsed_min is not None else 0),
        item.get("max", parsed_max if parsed_max is not None else ""),
        sub_field,
    )


def _ratio_from_text(value: str) -> float | None:
    match = re.search(r"([0-9.]+)\s*%", value or "")
    if not match:
        return None
    try:
        return float(match.group(1)) / 100
    except ValueError:
        return None


def _label_value(lines: list[str], label: str) -> str:
    try:
        index = lines.index(label)
    except ValueError:
        prefix = next((line for line in lines if line.startswith(label)), "")
        return prefix[len(label) :].strip() if prefix else ""
    return lines[index + 1] if index + 1 < len(lines) else ""


def _metric_value_after(lines: list[str], label: str) -> float | None:
    return _number_from_text(_label_value(lines, label))


def _ratio_after_label(lines: list[str], label: str) -> float | None:
    try:
        index = lines.index(label)
    except ValueError:
        return None
    window = " ".join(lines[index + 1 : index + 4])
    return _ratio_from_text(window)


def _ratio_near_label(lines: list[str], label: str) -> float | None:
    for index, line in enumerate(lines):
        if label not in line:
            continue
        direct = _ratio_from_text(line)
        if direct is not None:
            return direct
        window = " ".join(lines[index + 1 : index + 4])
        value = _ratio_from_text(window)
        if value is not None:
            return value
    return None


def _line_between(lines: list[str], start: str, end: str) -> str:
    try:
        start_index = lines.index(start)
    except ValueError:
        return ""
    try:
        end_index = lines.index(end, start_index + 1)
    except ValueError:
        return ""
    return "".join(lines[start_index + 1 : end_index]).strip()


def _gender_ratio(lines: list[str], gender: str) -> float | None:
    for index, line in enumerate(lines):
        if gender not in line:
            continue
        window = " ".join(lines[max(0, index - 2) : index + 4])
        value = _ratio_from_text(window)
        if value is not None:
            return value
    return None


def _parse_region_distribution_line(value: Any) -> dict[str, Any]:
    text = _clean_text(value)
    if not text:
        return {}
    top_regions = [
        {"label": match.group(1), "ratio": float(match.group(2)) / 100}
        for match in re.finditer(r"([\u4e00-\u9fffA-Za-z .]+?)（([0-9]+(?:\.[0-9]+)?)%）", text)
    ]
    summary: dict[str, Any] = {"raw_text": text, "source": "visible_text"}
    if top_regions:
        summary["top_regions"] = top_regions
    dominant = top_regions[0] if top_regions else None
    if dominant:
        summary["dominant"] = dominant
    if "按省份" in text:
        summary["scope"] = "province"
    elif "按城市" in text:
        summary["scope"] = "city"
    elif "省份" in text:
        summary["scope"] = "province"
    elif "城市" in text:
        summary["scope"] = "city"
    return summary


def _parse_device_distribution_line(value: Any) -> dict[str, Any]:
    text = _clean_text(value)
    if not text:
        return {}
    match = re.search(r"([^，,]+?)用户占比\s*([0-9]+(?:\.[0-9]+)?)%", text)
    summary_text = text
    insight = ""
    if "消费力较强" in text:
        insight = "消费力较强"
    elif "消费力" in text and "，" in text:
        insight = text.split("，", 1)[-1].strip()
    if match:
        summary_text = f"{match.group(1).strip()}用户占比{match.group(2)}%"
        if insight:
            summary_text = f"{summary_text}，{insight}"
    summary: dict[str, Any] = {"raw_text": summary_text, "source": "visible_text"}
    if match:
        summary["dominant"] = {
            "label": match.group(1).strip(),
            "ratio": float(match.group(2)) / 100,
        }
    if insight:
        summary["insight"] = insight
    return summary


def _build_detail_collection_summary(detail: dict[str, Any]) -> dict[str, Any]:
    raw = detail.get("raw_payload") if isinstance(detail.get("raw_payload"), dict) else {}
    fan_analysis = raw.get("fan_analysis") if isinstance(raw.get("fan_analysis"), dict) else {}
    note_performance = raw.get("note_performance") if isinstance(raw.get("note_performance"), dict) else {}
    service_performance = raw.get("service_performance") if isinstance(raw.get("service_performance"), dict) else {}
    note_cases = []
    for key in ("cooperation_note_cases", "recent_notes", "note_cases", "recent_note_briefs"):
        value = raw.get(key)
        if isinstance(value, list):
            note_cases.extend(item for item in value if isinstance(item, dict) or isinstance(item, str))
    modules: list[str] = []
    if detail.get("nickname") and detail.get("followers_count") is not None:
        modules.append("basic_profile")
    if note_performance or detail.get("daily_read_median") is not None or detail.get("cooperation_read_median") is not None:
        modules.append("note_performance")
    if fan_analysis or detail.get("fans_35_plus_ratio") is not None:
        modules.append("fan_analysis")
    if detail.get("audience_age_distribution") or detail.get("audience_gender_distribution"):
        modules.append("audience_chart")
    if detail.get("audience_region_distribution"):
        modules.append("region_distribution")
    if detail.get("audience_device_distribution"):
        modules.append("device_distribution")
    if service_performance or detail.get("reply_rate_48h") is not None:
        modules.append("service_performance")
    if note_cases:
        modules.append("note_cases")
    if detail.get("audience_profile_screenshot"):
        modules.append("audience_profile_screenshot")
    return {
        "modules": modules,
        "module_count": len(modules),
        "note_case_count": len(note_cases),
        "has_basic_profile": "basic_profile" in modules,
        "has_note_performance": "note_performance" in modules,
        "has_fan_analysis": "fan_analysis" in modules,
        "has_audience_chart": "audience_chart" in modules,
        "has_region_distribution": "region_distribution" in modules,
        "has_device_distribution": "device_distribution" in modules,
        "has_service_performance": "service_performance" in modules,
        "has_note_cases": "note_cases" in modules,
        "has_audience_profile_screenshot": "audience_profile_screenshot" in modules,
    }


def _merge_ratio_metric(target: dict[str, Any], key: str, value: Any) -> None:
    parsed = _ratio_from_text(str(value)) if isinstance(value, str) else None
    if parsed is None:
        try:
            number = float(value)
            parsed = number / 100 if number > 1 else number
        except (TypeError, ValueError):
            parsed = None
    if parsed is not None:
        target[key] = min(max(parsed, 0), 1)


def _collect_audience_profile_chart_metrics(page: Any) -> dict[str, Any]:
    try:
        payload = page.evaluate(
            """
            () => {
              const normalize = value => String(value || '').replace(/\\s+/g, ' ').trim();
              const ratio = value => {
                if (value === null || value === undefined || value === '') return null;
                if (typeof value === 'number') return value > 1 ? value / 100 : value;
                const match = String(value).match(/([0-9]+(?:\\.[0-9]+)?)\\s*%/);
                return match ? Number(match[1]) / 100 : null;
              };
              const keyFor = name => {
                const text = normalize(name);
                if (/(<\\s*18|18\\s*岁?以下)/.test(text)) return 'fans_under_18_ratio';
                if (/18\\s*[-~～至到]\\s*24/.test(text)) return 'fans_18_24_ratio';
                if (/25\\s*[-~～至到]\\s*34/.test(text)) return 'fans_25_34_ratio';
                if (/35\\s*[-~～至到]\\s*44/.test(text)) return 'fans_35_44_ratio';
                if (/(>\\s*44|44\\s*岁?以上|44\\+)/.test(text)) return 'fans_44_plus_ratio';
                if (/女性/.test(text)) return 'female_fans_ratio';
                if (/男性/.test(text)) return 'male_fans_ratio';
                return '';
              };
              const put = (out, name, value, source) => {
                const key = keyFor(name);
                const val = ratio(value);
                if (key && val !== null && out[key] === undefined) {
                  out[key] = val;
                  out.sources.push({ key, name: normalize(name), value, source });
                }
              };
              const walk = (node, visit, seen = new Set()) => {
                if (node === null || node === undefined || seen.has(node)) return;
                seen.add(node);
                visit(node);
                if (Array.isArray(node)) {
                  node.forEach(item => walk(item, visit, seen));
                } else if (typeof node === 'object') {
                  Object.values(node).forEach(item => walk(item, visit, seen));
                }
              };
              const collectOption = (out, option, source) => {
                walk(option, item => {
                  if (!item || typeof item !== 'object' || Array.isArray(item)) return;
                  if (item.name !== undefined && item.value !== undefined) put(out, item.name, item.value, source);
                  if (item.label !== undefined && item.value !== undefined) put(out, item.label, item.value, source);
                });
                const series = Array.isArray(option?.series) ? option.series : [];
                for (const item of series) {
                  const data = Array.isArray(item?.data) ? item.data : [];
                  const xData = Array.isArray(option?.xAxis) ? option.xAxis.flatMap(axis => axis?.data || []) : (option?.xAxis?.data || []);
                  const yData = Array.isArray(option?.yAxis) ? option.yAxis.flatMap(axis => axis?.data || []) : (option?.yAxis?.data || []);
                  const names = [...xData, ...yData];
                  data.forEach((entry, index) => {
                    if (typeof entry === 'number') put(out, names[index], entry, source);
                    if (entry && typeof entry === 'object') put(out, entry.name ?? names[index], entry.value, source);
                  });
                }
              };
              const out = { sources: [] };
              const wrappers = Array.from(document.querySelectorAll('.age-chart__wrapper, .sex-chart__wrapper'))
                .filter(el => normalize(el.innerText).includes('分布'));
              for (const wrapper of wrappers) {
                const text = normalize(wrapper.innerText);
                const isAgeWrapper = wrapper.matches('.age-chart__wrapper') || text.includes('年龄分布');
                const isGenderWrapper = wrapper.matches('.sex-chart__wrapper') || text.includes('性别分布');
                const dominant = text.match(/(<\\s*18|18\\s*岁?以下|18\\s*[-~～至到]\\s*24|25\\s*[-~～至到]\\s*34|35\\s*[-~～至到]\\s*44|>\\s*44|44\\s*岁?以上|女性|男性)\\s*居多[^0-9%]{0,12}占比\\s*([0-9]+(?:\\.[0-9]+)?)\\s*%/);
                if (dominant && isAgeWrapper) out.audience_age_dominant = { label: normalize(dominant[1]), ratio: Number(dominant[2]) / 100 };
                if (dominant && isGenderWrapper) out.audience_gender_dominant = { label: normalize(dominant[1]), ratio: Number(dominant[2]) / 100 };
                const descMatches = text.matchAll(/(<\\s*18|18\\s*岁?以下|18\\s*[-~～至到]\\s*24|25\\s*[-~～至到]\\s*34|35\\s*[-~～至到]\\s*44|>\\s*44|44\\s*岁?以上|女性|男性)[^0-9%]{0,12}占比\\s*([0-9]+(?:\\.[0-9]+)?)\\s*%/g);
                for (const match of descMatches) put(out, match[1], `${match[2]}%`, 'visible_desc');
                const chartNodes = Array.from(wrapper.querySelectorAll('[_echarts_instance_], canvas, [aria-label], [title]'));
                for (const node of chartNodes) {
                  for (const attr of ['aria-label', 'title']) {
                    const value = node.getAttribute?.(attr);
                    if (!value) continue;
                    for (const match of normalize(value).matchAll(/(<\\s*18|18\\s*岁?以下|18\\s*[-~～至到]\\s*24|25\\s*[-~～至到]\\s*34|35\\s*[-~～至到]\\s*44|>\\s*44|44\\s*岁?以上|女性|男性)[^0-9%]{0,16}([0-9]+(?:\\.[0-9]+)?)\\s*%/g)) {
                      put(out, match[1], `${match[2]}%`, attr);
                    }
                  }
                  const chartRoot = node.closest?.('[_echarts_instance_]') || (node.hasAttribute?.('_echarts_instance_') ? node : null);
                  if (chartRoot && window.echarts?.getInstanceByDom) {
                    const instance = window.echarts.getInstanceByDom(chartRoot);
                    if (instance?.getOption) collectOption(out, instance.getOption(), 'echarts_option');
                  }
                }
              }
              return out;
            }
            """
        )
    except Exception:
        return {}
    if not isinstance(payload, dict):
        return {}
    result: dict[str, Any] = {
        "audience_age_distribution": {
            "segments": [
                {"label": "<18", "key": "fans_under_18_ratio", "ratio": payload.get("fans_under_18_ratio")},
                {"label": "18-24", "key": "fans_18_24_ratio", "ratio": payload.get("fans_18_24_ratio")},
                {"label": "25-34", "key": "fans_25_34_ratio", "ratio": payload.get("fans_25_34_ratio")},
                {"label": "35-44", "key": "fans_35_44_ratio", "ratio": payload.get("fans_35_44_ratio")},
                {"label": ">44", "key": "fans_44_plus_ratio", "ratio": payload.get("fans_44_plus_ratio")},
            ],
            "dominant": payload.get("audience_age_dominant") or "",
            "source": "dom_echarts",
        },
        "audience_gender_distribution": {
            "segments": [
                {"label": "女性", "key": "female_fans_ratio", "ratio": payload.get("female_fans_ratio")},
                {"label": "男性", "key": "male_fans_ratio", "ratio": payload.get("male_fans_ratio")},
            ],
            "dominant": payload.get("audience_gender_dominant") or "",
            "source": "dom_echarts",
        },
        "raw_payload": {
            "audience_profile_chart_metrics": {
                "source": "dom_echarts",
                "metrics": {key: value for key, value in payload.items() if key != "sources"},
                "sources": payload.get("sources") or [],
            }
        }
    }
    for key in [
        "female_fans_ratio",
        "male_fans_ratio",
        "fans_under_18_ratio",
        "fans_18_24_ratio",
        "fans_25_34_ratio",
        "fans_35_44_ratio",
        "fans_44_plus_ratio",
    ]:
        if key in payload:
            _merge_ratio_metric(result, key, payload.get(key))
    for distribution_key in ["audience_age_distribution", "audience_gender_distribution"]:
        segments = []
        for item in result[distribution_key]["segments"]:
            parsed = _ratio_from_text(str(item.get("ratio"))) if isinstance(item.get("ratio"), str) else item.get("ratio")
            try:
                parsed = float(parsed)
            except (TypeError, ValueError):
                parsed = None
            if parsed is not None:
                parsed = parsed / 100 if parsed > 1 else parsed
                segments.append({**item, "ratio": min(max(parsed, 0), 1)})
        result[distribution_key]["segments"] = segments
    return result


def _extract_note_cases(lines: list[str]) -> list[dict[str, Any]]:
    if "笔记案例" not in lines:
        return []
    try:
        start = lines.index("仅展示跨域合作笔记") + 1
    except ValueError:
        start = lines.index("笔记案例") + 1
    try:
        end = lines.index("前往TA的小红书APP主页", start)
    except ValueError:
        end = min(len(lines), start + 120)
    cases: list[dict[str, Any]] = []
    index = start
    while index < end and len(cases) < 12:
        if index + 1 >= end:
            break
        brand = lines[index]
        title = lines[index + 1]
        if brand in {"1", "2", "3", "4"} or title in {"阅读", "点赞", "收藏", "发布时间"}:
            index += 1
            continue
        case: dict[str, Any] = {"brand": brand, "title": title}
        cursor = index + 2
        if cursor < end and lines[cursor] == "含推广流量":
            case["has_promoted_traffic"] = True
            cursor += 1
        fields = {"阅读": "read_count", "点赞": "like_count", "收藏": "save_count"}
        while cursor < end:
            current = lines[cursor]
            if current in fields and cursor + 1 < end:
                case[fields[current]] = _number_from_text(lines[cursor + 1])
                cursor += 2
                continue
            if current == "发布时间" and cursor + 1 < end:
                case["published_at"] = lines[cursor + 1]
                cursor += 2
                break
            cursor += 1
        if case.get("published_at"):
            cases.append(case)
            index = cursor
        else:
            index += 1
    return cases


def _compact_ratio(value: float) -> str:
    if value >= 10:
        return f"{value:.0f}"
    if value >= 1:
        return f"{value:.1f}".rstrip("0").rstrip(".")
    return f"{value:.2f}".rstrip("0").rstrip(".")


def _note_interaction_count(case: dict[str, Any]) -> float | None:
    total = 0.0
    found = False
    for key in ["like_count", "save_count", "comment_count", "share_count"]:
        value = _number_from_text(str(case.get(key) or ""))
        if value is None:
            continue
        total += value
        found = True
    return total if found else None


def _case_title_key(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip().lower()


def _extract_note_case_dom_assets(page: Any) -> list[dict[str, Any]]:
    try:
        assets = page.evaluate(
            """
            () => {
              const normalize = value => String(value || '').replace(/\\s+/g, '\\n').trim();
              const inline = value => String(value || '').replace(/\\s+/g, ' ').trim();
              const visible = el => {
                const rect = el.getBoundingClientRect();
                const style = window.getComputedStyle(el);
                return rect.width > 0 && rect.height > 0 && style.display !== 'none' && style.visibility !== 'hidden';
              };
              const absolutize = value => {
                if (!value) return '';
                try { return new URL(value, location.href).href; } catch { return value; }
              };
              const imageFor = node => {
                const img = node.querySelector('img');
                if (img?.currentSrc || img?.src) return absolutize(img.currentSrc || img.src);
                for (const el of [node, ...Array.from(node.querySelectorAll('*'))]) {
                  const style = window.getComputedStyle(el);
                  const bg = style.backgroundImage || '';
                  const match = bg.match(/url\\(["']?([^"')]+)["']?\\)/);
                  if (match?.[1]) return absolutize(match[1]);
                }
                return '';
              };
              const linkFor = node => {
                const anchors = [
                  node.closest?.('a'),
                  ...Array.from(node.querySelectorAll('a')),
                ].filter(Boolean);
                const preferred = anchors.find(anchor => /xiaohongshu\\.com\\/(?:explore|discovery|search_result|user\\/profile)/.test(anchor.href || ''))
                  || anchors.find(anchor => anchor.href && !/javascript:void|^#$/.test(anchor.getAttribute('href') || ''));
                return preferred ? absolutize(preferred.href) : '';
              };
              const wrappers = Array.from(document.querySelectorAll('.note-case-wrapper, [class*="note-case"], [class*="case-wrapper"], section, div'))
                .filter(visible)
                .filter(el => (el.innerText || '').includes('笔记案例'));
              const wrapper = wrappers.sort((a, b) => (a.innerText || '').length - (b.innerText || '').length)[0];
              if (!wrapper) return [];
              let candidates = Array.from(wrapper.querySelectorAll('article, li, a, [class*="card"], [class*="item"], [class*="note"]'))
                .filter(visible)
                .map(node => ({ node, text: normalize(node.innerText || node.textContent || '') }))
                .filter(item => item.text.includes('阅读') && item.text.includes('点赞') && item.text.includes('收藏') && item.text.includes('发布时间'))
                .filter(item => item.text.length >= 20 && item.text.length <= 900);
              candidates = candidates.filter(item => !candidates.some(other => other.node !== item.node && item.node.contains(other.node) && other.text.length < item.text.length));
              const seen = new Set();
              return candidates.map((item, index) => {
                const lines = item.text.split('\\n').map(line => line.trim()).filter(Boolean);
                const title = lines.find((line, i) => i > 0 && !['阅读', '点赞', '收藏', '发布时间', '含推广流量'].includes(line) && !/^\\d+$/.test(line)) || '';
                const payload = {
                  index,
                  title,
                  raw_text: inline(item.text),
                  cover_url: imageFor(item.node),
                  note_url: linkFor(item.node),
                  source_url: location.href,
                };
                const key = `${payload.title}|${payload.cover_url}|${payload.note_url}|${payload.raw_text.slice(0, 80)}`;
                if (seen.has(key)) return null;
                seen.add(key);
                return payload;
              }).filter(Boolean);
            }
            """
        )
    except Exception:
        return []
    return [item for item in assets if isinstance(item, dict)]


def _merge_note_case_assets(cases: list[dict[str, Any]], assets: list[dict[str, Any]], source_url: str = "") -> list[dict[str, Any]]:
    if not cases:
        return []
    unused = list(enumerate(assets or []))
    merged: list[dict[str, Any]] = []
    for index, case in enumerate(cases):
        title_key = _case_title_key(case.get("title"))
        match_pos = next(
            (
                pos
                for pos, (_, asset) in enumerate(unused)
                if title_key
                and (
                    title_key in _case_title_key(asset.get("title") or asset.get("raw_text"))
                    or _case_title_key(asset.get("title") or asset.get("raw_text")) in title_key
                )
            ),
            None,
        )
        if match_pos is None:
            match_pos = next((pos for pos, (asset_index, _) in enumerate(unused) if asset_index == index), None)
        asset = unused.pop(match_pos)[1] if match_pos is not None else {}
        next_case = {**case}
        for source_key, target_key in [
            ("cover_url", "cover_url"),
            ("note_url", "note_url"),
            ("source_url", "source_url"),
        ]:
            if asset.get(source_key) and not next_case.get(target_key):
                next_case[target_key] = asset[source_key]
        if source_url and not next_case.get("source_url"):
            next_case["source_url"] = source_url
        merged.append(next_case)
    return merged


def _comment_summary(comments: Any, limit: int = 3, max_chars: int = 240) -> str:
    if not isinstance(comments, list):
        return ""
    texts: list[str] = []
    for item in comments:
        text = str(item.get("content") if isinstance(item, dict) else item or "").strip()
        if text:
            texts.append(text)
        if len(texts) >= limit:
            break
    summary = "；".join(texts)
    return summary[:max_chars]


def _note_title_key(value: Any) -> str:
    return re.sub(r"[\s，,。.!！?？:：;；《》\"'“”‘’（）()【】\\[\\]-]+", "", str(value or "")).lower()


def _note_detail_match_score(case: dict[str, Any], note: dict[str, Any], index: int, note_index: int) -> int:
    score = 0
    if case.get("note_id") and note.get("note_id") and str(case.get("note_id")) == str(note.get("note_id")):
        score += 100
    if case.get("note_url") and note.get("note_url") and str(case.get("note_url")) == str(note.get("note_url")):
        score += 80
    case_title = _note_title_key(case.get("title"))
    note_title = _note_title_key(note.get("title"))
    if case_title and note_title:
        if case_title == note_title:
            score += 60
        elif case_title in note_title or note_title in case_title:
            score += 35
    if case.get("published_at") and note.get("published_at") and str(case.get("published_at"))[:10] == str(note.get("published_at"))[:10]:
        score += 18
    if _number_from_text(str(case.get("read_count") or "")) is not None and _number_from_text(str(note.get("read_count") or "")) is not None:
        if _number_from_text(str(case.get("read_count") or "")) == _number_from_text(str(note.get("read_count") or "")):
            score += 22
    if index == note_index:
        score += 8
    return score


def _merge_note_detail_into_case(case: dict[str, Any], note: dict[str, Any]) -> dict[str, Any]:
    merged = dict(case)
    for key in [
        "note_id",
        "note_url",
        "content",
        "comments",
        "cover_url",
        "comment_count",
        "share_count",
        "follow_count",
        "note_type",
        "component_click_data",
    ]:
        value = note.get(key)
        if value not in ("", None, []) and not merged.get(key):
            merged[key] = value
    if not merged.get("comment_summary"):
        summary = _comment_summary(merged.get("comments"))
        if summary:
            merged["comment_summary"] = summary
    if note.get("source"):
        merged["note_detail_source"] = note["source"]
    return merged


def _enrich_note_cases_with_details(cases: list[dict[str, Any]], notes: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if not cases or not notes:
        return cases
    usable_notes = [note for note in notes if isinstance(note, dict) and (note.get("note_id") or note.get("note_url") or note.get("title"))]
    if not usable_notes:
        return cases
    enriched: list[dict[str, Any]] = []
    used: set[int] = set()
    for index, case in enumerate(cases):
        if not isinstance(case, dict):
            enriched.append(case)
            continue
        scored = [
            (_note_detail_match_score(case, note, index, note_index), note_index, note)
            for note_index, note in enumerate(usable_notes)
            if note_index not in used
        ]
        scored = [item for item in scored if item[0] >= 8]
        if not scored:
            enriched.append(case)
            continue
        scored.sort(key=lambda item: item[0], reverse=True)
        _, note_index, note = scored[0]
        used.add(note_index)
        enriched.append(_merge_note_detail_into_case(case, note))
    return enriched


def _merge_note_details_into_payload(detail: dict[str, Any]) -> dict[str, Any]:
    raw = detail.get("raw_payload")
    if not isinstance(raw, dict):
        return detail
    notes = raw.get("recent_notes") if isinstance(raw.get("recent_notes"), list) else []
    if not notes:
        return detail
    for key in ["note_cases", "cooperation_note_cases"]:
        if isinstance(raw.get(key), list):
            raw[key] = _enrich_note_cases_with_details(raw[key], notes)
    pages = raw.get("cooperation_note_case_pages")
    if isinstance(pages, list):
        for page_item in pages:
            if isinstance(page_item, dict) and isinstance(page_item.get("cases"), list):
                page_item["cases"] = _enrich_note_cases_with_details(page_item["cases"], notes)
    detail["raw_payload"] = raw
    return detail


def _annotate_note_cases_with_traffic_reference(detail: dict[str, Any]) -> dict[str, Any]:
    raw = detail.get("raw_payload")
    if not isinstance(raw, dict):
        return detail
    read_median = detail.get("cooperation_read_median") or detail.get("daily_read_median")
    interaction_median = detail.get("cooperation_interaction_median") or detail.get("daily_interaction_median")
    reference_source = "cooperation" if detail.get("cooperation_read_median") or detail.get("cooperation_interaction_median") else "daily"
    read_median = _number_from_text(str(read_median or ""))
    interaction_median = _number_from_text(str(interaction_median or ""))
    if not read_median and not interaction_median:
        return detail

    def annotate(case: dict[str, Any]) -> dict[str, Any]:
        if not isinstance(case, dict):
            return case
        next_case = {**case}
        comparisons: list[str] = []
        is_clear = False
        read_count = _number_from_text(str(next_case.get("read_count") or next_case.get("readCount") or ""))
        if read_count is not None and read_median:
            ratio = read_count / read_median
            next_case["read_vs_median"] = round(ratio, 3)
            comparisons.append(f"阅读为{_compact_ratio(ratio)}倍")
            is_clear = is_clear or ratio >= 1.5 or ratio <= 0.67
        interaction_count = _note_interaction_count(next_case)
        if interaction_count is not None and interaction_median:
            ratio = interaction_count / interaction_median
            next_case["interaction_count"] = interaction_count
            next_case["interaction_vs_median"] = round(ratio, 3)
            comparisons.append(f"互动为{_compact_ratio(ratio)}倍")
            is_clear = is_clear or ratio >= 1.5 or ratio <= 0.67
        if comparisons:
            next_case["traffic_median_reference"] = {
                "source": reference_source,
                "read_median": read_median,
                "interaction_median": interaction_median,
            }
            next_case["traffic_comparison"] = "，".join(comparisons)
            next_case["has_clear_median_contrast"] = is_clear
        return next_case

    for key in ["note_cases", "recent_note_cases", "recent_notes", "cooperation_note_cases", "notes"]:
        if isinstance(raw.get(key), list):
            raw[key] = [annotate(case) for case in raw[key]]
    pages = raw.get("cooperation_note_case_pages")
    if isinstance(pages, list):
        for page in pages:
            if isinstance(page, dict) and isinstance(page.get("cases"), list):
                page["cases"] = [annotate(case) for case in page["cases"]]
    detail["raw_payload"] = raw
    return detail


def _value_after_label(lines: list[str], label: str) -> str:
    try:
        index = lines.index(label)
    except ValueError:
        return ""
    stop_labels = {
        "曝光中位数",
        "阅读中位数",
        "互动中位数",
        "外溢进店中位数",
        "预估CPM",
        "预估CPM(图文)",
        "预估CPM(视频)",
        "预估阅读单价",
        "预估阅读单价(图文)",
        "预估阅读单价(视频)",
        "预估互动单价",
        "预估互动单价(图文)",
        "预估互动单价(视频)",
        "预估外溢进店单价(图文)",
        "预估外溢进店单价(视频)",
        "外溢进店单价",
        "中位点赞量",
        "中位收藏量",
        "中位评论量",
        "中位分享量",
        "中位关注量",
        "互动率",
        "视频完播率",
        "图文3秒阅读率",
        "千赞笔记比例",
        "百赞笔记比例",
    }
    unit_lines = {"元/千次曝光", "元/阅读", "元/互动", "元/进店"}
    parts: list[str] = []
    for item in lines[index + 1 : index + 4]:
        if item in stop_labels or item in {"优于", "同行", "按规模", "按成本", "日常笔记", "合作笔记", "其他指标", "核心指标"}:
            break
        if item.startswith("优于"):
            break
        parts.append(item)
        if _number_from_text(item) is not None or "%" in item or item == "-":
            if len(parts) >= 2 and item in unit_lines:
                break
            if len(parts) == 1:
                next_index = index + 2
                if next_index < len(lines) and lines[next_index] in unit_lines:
                    continue
            break
    return "".join(parts).strip()


def _metrics_from_lines(lines: list[str], labels: list[str]) -> dict[str, str]:
    metrics: dict[str, str] = {}
    for label in labels:
        value = _value_after_label(lines, label)
        if value:
            metrics[label] = value
    return metrics


def _extract_overview_note_state(text: str) -> dict[str, Any]:
    lines = _visible_text_lines(text)
    try:
        service_index = lines.index("服务表现")
    except ValueError:
        service_index = len(lines)
    note_indexes = [index for index, line in enumerate(lines[:service_index]) if line == "笔记数据"]
    start = note_indexes[-1] if note_indexes else max(0, service_index - 40)
    block = lines[start:service_index]
    labels = [
        "曝光中位数",
        "阅读中位数",
        "互动中位数",
        "外溢进店中位数",
        "预估CPM",
        "预估CPM(图文)",
        "预估CPM(视频)",
        "预估阅读单价",
        "预估阅读单价(图文)",
        "预估阅读单价(视频)",
        "预估互动单价",
        "预估互动单价(图文)",
        "预估互动单价(视频)",
        "预估外溢进店单价(图文)",
        "预估外溢进店单价(视频)",
        "外溢进店单价",
    ]
    return {"text": "\n".join(block), "metrics": _metrics_from_lines(block, labels)}


def _extract_performance_state(text: str) -> dict[str, Any]:
    lines = _visible_text_lines(text)
    try:
        start = lines.index("数据表现")
    except ValueError:
        start = 0
    try:
        end = lines.index("粉丝分析", start + 1)
    except ValueError:
        end = min(len(lines), start + 160)
    block = lines[start:end]
    labels = [
        "曝光中位数",
        "阅读中位数",
        "互动中位数",
        "外溢进店中位数",
        "预估CPM",
        "预估阅读单价",
        "预估互动单价",
        "预估外溢进店单价(图文)",
        "预估外溢进店单价(视频)",
        "外溢进店单价",
        "中位点赞量",
        "中位收藏量",
        "中位评论量",
        "中位分享量",
        "中位关注量",
        "互动率",
        "视频完播率",
        "图文3秒阅读率",
        "千赞笔记比例",
        "百赞笔记比例",
    ]
    return {"text": "\n".join(block), "metrics": _metrics_from_lines(block, labels)}


def _body_text(page: Any, timeout: int = 5000) -> str:
    try:
        return page.locator("body").inner_text(timeout=timeout)
    except Exception:
        return ""


def _container_text(page: Any, selector: str, contains: list[str] | None = None, occurrence: str = "first") -> str:
    try:
        return page.evaluate(
            """
            ({ selector, contains, occurrence }) => {
              const normalized = value => String(value || '').replace(/\\s+/g, '\\n').trim();
              const isVisibleEnough = el => {
                const rect = el.getBoundingClientRect();
                const style = window.getComputedStyle(el);
                return rect.width > 0 && rect.height > 0 && style.display !== 'none' && style.visibility !== 'hidden';
              };
              const items = Array.from(document.querySelectorAll(selector))
                .filter(isVisibleEnough)
                .filter(el => {
                  const text = el.innerText || '';
                  return (contains || []).every(item => text.includes(item));
                })
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  return ar.top - br.top || ar.left - br.left;
                });
              const target = occurrence === 'last' ? items[items.length - 1] : items[0];
              return target ? normalized(target.innerText || target.textContent || '') : '';
            }
            """,
            {"selector": selector, "contains": contains or [], "occurrence": occurrence},
        )
    except Exception:
        return ""


def _scroll_container(page: Any, selector: str, contains: list[str] | None = None, occurrence: str = "first") -> bool:
    try:
        scrolled = page.evaluate(
            """
            ({ selector, contains, occurrence }) => {
              const items = Array.from(document.querySelectorAll(selector))
                .filter(el => {
                  const rect = el.getBoundingClientRect();
                  const style = window.getComputedStyle(el);
                  const text = el.innerText || '';
                  return rect.width > 0 && rect.height > 0
                    && style.display !== 'none'
                    && style.visibility !== 'hidden'
                    && (contains || []).every(item => text.includes(item));
                })
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  return ar.top - br.top || ar.left - br.left;
                });
              const target = occurrence === 'last' ? items[items.length - 1] : items[0];
              if (!target) return false;
              target.scrollIntoView({ block: 'center', inline: 'nearest' });
              return true;
            }
            """,
            {"selector": selector, "contains": contains or [], "occurrence": occurrence},
        )
        if scrolled:
            page.wait_for_timeout(700)
            return True
    except Exception:
        pass
    return False


def _click_in_container(
    page: Any,
    container_selector: str,
    label: str,
    contains: list[str] | None = None,
    container_occurrence: str = "first",
    target_occurrence: str = "first",
) -> bool:
    _scroll_container(page, container_selector, contains, container_occurrence)
    try:
        box = page.evaluate(
            """
            ({ containerSelector, label, contains, containerOccurrence, targetOccurrence }) => {
              const normalized = value => String(value || '').replace(/\\s+/g, ' ').trim();
              const visible = el => {
                const rect = el.getBoundingClientRect();
                const style = window.getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden'
                  && rect.bottom >= 0
                  && rect.top <= window.innerHeight;
              };
              const containers = Array.from(document.querySelectorAll(containerSelector))
                .filter(el => {
                  const rect = el.getBoundingClientRect();
                  const style = window.getComputedStyle(el);
                  const text = el.innerText || '';
                  return rect.width > 0 && rect.height > 0
                    && style.display !== 'none'
                    && style.visibility !== 'hidden'
                    && (contains || []).every(item => text.includes(item));
                })
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  return ar.top - br.top || ar.left - br.left;
                });
              const container = containerOccurrence === 'last' ? containers[containers.length - 1] : containers[0];
              if (!container) return false;
              const selector = [
                'button',
                '[role="button"]',
                '[role="tab"]',
                '.d-tabs-header',
                '.d-radio-button',
                '.d-segment-item',
                '.d-segmented-item',
                '.d-pagination-page',
                'label',
                'span',
                'div'
              ].join(',');
              const targets = Array.from(container.querySelectorAll(selector))
                .filter(visible)
                .filter(el => normalized(el.innerText || el.textContent) === label)
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  return ar.top - br.top || ar.left - br.left;
                });
              let target = targetOccurrence === 'last' ? targets[targets.length - 1] : targets[0];
              if (!target) return false;
              target = target.closest('button,[role="button"],[role="tab"],.d-tabs-header,.d-radio-button,.d-segment-item,.d-segmented-item,.d-pagination-page,label') || target;
              const rect = target.getBoundingClientRect();
              return { x: rect.x + rect.width / 2, y: rect.y + rect.height / 2 };
            }
            """,
            {
                "containerSelector": container_selector,
                "label": label,
                "contains": contains or [],
                "containerOccurrence": container_occurrence,
                "targetOccurrence": target_occurrence,
            },
        )
        if box and isinstance(box, dict):
            page.mouse.click(float(box["x"]), float(box["y"]))
            page.wait_for_timeout(1000)
            return True
    except Exception:
        pass
    return False


def _scroll_to_visible_text(page: Any, label: str, occurrence: str = "first") -> bool:
    try:
        locator = page.locator(f"text={label}")
        count = locator.count()
        if not count:
            return False
        target = locator.last if occurrence == "last" else locator.first
        target.scroll_into_view_if_needed(timeout=3000)
        page.wait_for_timeout(600)
        return True
    except Exception:
        return False


def _click_visible_text(page: Any, label: str, occurrence: str = "first", viewport_only: bool = True) -> bool:
    try:
        clicked = page.evaluate(
            """
            ({ label, occurrence, viewportOnly }) => {
              const normalized = value => String(value || '').replace(/\\s+/g, ' ').trim();
              const isVisible = el => {
                const rect = el.getBoundingClientRect();
                const style = window.getComputedStyle(el);
                if (!rect.width || !rect.height) return false;
                if (style.display === 'none' || style.visibility === 'hidden' || style.opacity === '0') return false;
                if (!viewportOnly) return true;
                return rect.bottom >= 0 && rect.top <= window.innerHeight && rect.right >= 0 && rect.left <= window.innerWidth;
              };
              const selector = [
                'button',
                '[role="button"]',
                '[role="tab"]',
                '.d-tabs-tab',
                '.d-radio-button',
                '.d-segmented-item',
                '.d-dropdown-menu-item',
                '.d-select-option',
                '.d-pagination-page',
                'label',
                'span',
                'div'
              ].join(',');
              const candidates = Array.from(document.querySelectorAll(selector))
                .filter(isVisible)
                .filter(el => normalized(el.innerText || el.textContent) === label)
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  return ar.top - br.top || ar.left - br.left;
                });
              if (!candidates.length) return false;
              const target = occurrence === 'last' ? candidates[candidates.length - 1] : candidates[0];
              target.click();
              return true;
            }
            """,
            {"label": label, "occurrence": occurrence, "viewportOnly": viewport_only},
        )
        if clicked:
            page.wait_for_timeout(900)
            return True
    except Exception:
        pass
    try:
        locator = page.locator(f"text={label}")
        target = locator.last if occurrence == "last" else locator.first
        target.click(timeout=2000)
        page.wait_for_timeout(900)
        return True
    except Exception:
        return False


def _click_detail_tab(page: Any, section_label: str, tab_label: str, section_occurrence: str = "first", tab_occurrence: str = "first") -> bool:
    _scroll_to_visible_text(page, section_label, section_occurrence)
    return _click_visible_text(page, tab_label, tab_occurrence, viewport_only=True)


def _click_note_case_page(page: Any, page_no: int) -> bool:
    _scroll_container(page, ".note-case-wrapper", ["笔记案例"], "first")
    try:
        clicked = page.evaluate(
            """
            pageNo => {
              const normalized = value => String(value || '').replace(/\\s+/g, ' ').trim();
              const isVisible = el => {
                const rect = el.getBoundingClientRect();
                const style = window.getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden'
                  && rect.bottom >= 0
                  && rect.top <= window.innerHeight;
              };
              const pages = Array.from(document.querySelectorAll('.d-pagination-page'))
                .filter(isVisible)
                .filter(el => normalized(el.innerText || el.textContent).split(' ')[0] === String(pageNo))
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  return ar.top - br.top || ar.left - br.left;
                });
              const target = pages[pages.length - 1];
              if (!target) return false;
              target.click();
              return true;
            }
            """,
            page_no,
        )
        if clicked:
            page.wait_for_timeout(1400)
            return True
    except Exception:
        pass
    return False


def _case_key(case: dict[str, Any]) -> tuple[str, str, str]:
    return (
        str(case.get("brand") or "").strip(),
        str(case.get("title") or "").strip(),
        str(case.get("published_at") or "").strip(),
    )


def _collect_note_case_pages(page: Any, max_cases: int = 24) -> dict[str, Any]:
    _ensure_note_detail_api_pages(page, max_pages=3)
    api_result = _api_note_case_pages_from_cache(page, note_type=3, max_cases=max_cases) or _api_note_case_pages_from_cache(page, note_type=4, max_cases=max_cases)
    if api_result:
        return api_result
    _click_in_container(page, ".note-case-wrapper", "合作笔记", contains=["笔记案例"], target_occurrence="first")
    collected: list[dict[str, Any]] = []
    pages: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str]] = set()
    for page_no in range(1, 4):
        if page_no > 1:
            if not _click_note_case_page(page, page_no):
                break
        text = _container_text(page, ".note-case-wrapper", ["笔记案例"], "first") or _body_text(page)
        lines = _visible_text_lines(text)
        cases = _merge_note_case_assets(_extract_note_cases(lines), _extract_note_case_dom_assets(page), page.url)
        pages.append({"page": page_no, "count": len(cases), "cases": cases})
        before = len(collected)
        for case in cases:
            key = _case_key(case)
            if key in seen:
                continue
            seen.add(key)
            collected.append(case)
            if len(collected) >= max_cases:
                break
        if len(collected) >= max_cases:
            break
        if page_no > 1 and len(collected) == before:
            break
    return {"cooperation_note_cases": collected, "cooperation_note_case_pages": pages}


def _collect_overview_note_states(page: Any) -> dict[str, Any]:
    _ensure_detail_summary_api_cache(page)
    api_states = _overview_from_api_cache(page)
    if api_states:
        return api_states
    states: dict[str, Any] = {}
    selector = ".detail-item"
    contains = ["笔记数据", "按规模", "按成本"]
    for note_type in ["日常笔记", "合作笔记"]:
        _click_in_container(page, selector, "按规模", contains=contains, target_occurrence="first")
        if not _click_in_container(page, selector, note_type, contains=contains, target_occurrence="first"):
            _scroll_container(page, selector, contains, "first")
            _click_visible_text(page, note_type, "first", viewport_only=True)
        for mode in ["按规模", "按成本"]:
            if not _click_in_container(page, selector, mode, contains=contains, target_occurrence="first"):
                _scroll_container(page, selector, contains, "first")
                _click_visible_text(page, mode, "first", viewport_only=True)
            text = _container_text(page, selector, contains, "first") or _body_text(page)
            key = "daily" if note_type == "日常笔记" else "cooperation"
            mode_key = "scale" if mode == "按规模" else "cost"
            states.setdefault(key, {})[mode_key] = _extract_overview_note_state(text)
    return states


def _collect_performance_states(page: Any) -> dict[str, Any]:
    _ensure_detail_summary_api_cache(page)
    api_states = _performance_from_api_cache(page)
    if api_states:
        return api_states
    states: dict[str, Any] = {}
    selector = ".trans-data-wrapper"
    contains = ["数据表现", "日常笔记", "合作笔记"]
    _scroll_to_visible_text(page, "数据表现", "last")
    _click_visible_text(page, "合作笔记", "last", viewport_only=True)
    for mode in ["按规模", "按成本"]:
        _scroll_to_visible_text(page, "数据表现", "last")
        _click_visible_text(page, mode, "last", viewport_only=True)
        text = _container_text(page, selector, contains, "first") or _body_text(page)
        states.setdefault("cooperation", {})["scale" if mode == "按规模" else "cost"] = _extract_performance_state(text)
    return states


def _collect_detail_interaction_states(page: Any) -> dict[str, Any]:
    raw: dict[str, Any] = {}
    try:
        raw["overview_note_data"] = _collect_overview_note_states(page)
    except Exception as exc:
        raw["overview_note_data_error"] = str(exc)
    try:
        raw.update(_collect_note_case_pages(page))
    except Exception as exc:
        raw["cooperation_note_cases_error"] = str(exc)
    try:
        raw["data_performance"] = _collect_performance_states(page)
    except Exception as exc:
        raw["data_performance_error"] = str(exc)
    result: dict[str, Any] = {"raw_payload": raw}
    scale_metrics = (
        (
            raw.get("data_performance", {}).get("cooperation")
            or {}
        )
        .get("scale", {})
        .get("metrics", {})
        if isinstance(raw.get("data_performance"), dict)
        else {}
    )
    cost_metrics = (
        (
            raw.get("data_performance", {}).get("cooperation")
            or {}
        )
        .get("cost", {})
        .get("metrics", {})
        if isinstance(raw.get("data_performance"), dict)
        else {}
    )
    field_map = {
        "cooperation_exposure_median": scale_metrics.get("曝光中位数"),
        "cooperation_read_median": scale_metrics.get("阅读中位数"),
        "cooperation_interaction_median": scale_metrics.get("互动中位数"),
        "overflow_store_median": scale_metrics.get("外溢进店中位数"),
        "overflow_store_unit_price": cost_metrics.get("外溢进店单价"),
    }
    result.update({key: value for key, value in field_map.items() if value not in ("", None)})
    return result


def _safe_filename(value: str) -> str:
    text = re.sub(r"[^\w\u4e00-\u9fff.-]+", "_", str(value or "").strip(), flags=re.U)
    return text.strip("._")[:80] or "unknown"


def _normalize_export_header(value: str) -> str:
    return re.sub(r"\s+", "", str(value or "").replace("\ufeff", "")).strip()


def _export_value(row: dict[str, Any], aliases: list[str]) -> Any:
    normalized = {_normalize_export_header(key): value for key, value in row.items()}
    for alias in aliases:
        value = normalized.get(_normalize_export_header(alias))
        if value not in (None, ""):
            return value
    return None


def _read_export_rows(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        raw = path.read_bytes()
        for encoding in ["utf-8-sig", "gb18030", "utf-16"]:
            try:
                text = raw.decode(encoding)
                return [dict(row) for row in csv.DictReader(text.splitlines())]
            except Exception:
                continue
        return []
    if suffix in {".xlsx", ".xls"}:
        try:
            from openpyxl import load_workbook
        except ImportError:
            return []
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value or "").strip() for value in rows[0]]
        return [
            {headers[index]: value for index, value in enumerate(row) if index < len(headers)}
            for row in rows[1:]
            if any(value not in (None, "") for value in row)
        ]
    return []


def _normalize_export_row(row: dict[str, Any]) -> dict[str, Any]:
    normalized = {"source": "pgy", "raw_payload": {"export_row": row}}
    for field, aliases in PGY_EXPORT_FIELD_ALIASES.items():
        value = _export_value(row, aliases)
        if value not in (None, ""):
            normalized[field] = value
    if not normalized.get("creator_id"):
        seed = "|".join(
            str(normalized.get(key) or "")
            for key in ["xiaohongshu_id", "pgy_url", "nickname", "followers_count", "quote_price"]
        )
        if seed.strip("|"):
            normalized["creator_id"] = f"pgy-export:{seed}"
    return normalized


def parse_export_file(path: str | Path) -> dict[str, Any]:
    export_path = Path(path)
    if not export_path.exists():
        return {"status": "failed", "message": "导出文件不存在", "path": str(export_path), "creators": []}
    rows = _read_export_rows(export_path)
    creators = [_normalize_export_row(row) for row in rows]
    creators = [creator for creator in creators if creator.get("nickname") or creator.get("pgy_url") or creator.get("xiaohongshu_id")]
    return {
        "status": "parsed",
        "message": f"已解析导出文件 {len(creators)} 行",
        "path": str(export_path),
        "row_count": len(rows),
        "creators": creators,
    }


def _merge_filters(*groups: list[dict[str, Any]]) -> list[dict[str, Any]]:
    merged: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str, str]] = set()
    for group in groups:
        for item in group or []:
            field = str(item.get("field") or "")
            value = str(item.get("value") or "")
            sub_field = str(item.get("sub_field") or item.get("subField") or "")
            sub_value = str(item.get("sub_value") or item.get("subValue") or "")
            if not field or not value:
                continue
            key = (field, value, sub_field, sub_value)
            if key in seen:
                continue
            seen.add(key)
            merged.append({**item, "field": field, "value": value, "reason": str(item.get("reason") or "")})
    return merged


def _brief_emphasizes_region(text: str) -> bool:
    normalized = _clean_text(text).lower()
    if not any(keyword in normalized for keyword in ["地域", "地区", "城市", "ip", "北京", "上海", "一线", "省份"]):
        return False
    emphasis_keywords = [
        "地域要求",
        "地区要求",
        "城市要求",
        "ip要求",
        "地域优先",
        "地区优先",
        "城市优先",
        "ip优先",
        "优先北京",
        "优先上海",
        "北京优先",
        "上海优先",
        "重点城市",
        "核心城市",
        "指定城市",
        "必须",
        "限定",
        "限制",
        "只要",
        "仅限",
        "重点覆盖",
        "地域强调",
        "地区强调",
        "本地",
        "同城",
    ]
    return any(keyword in normalized for keyword in emphasis_keywords)


def _hard_filters_to_pgy_filters(hard_filters: list[dict[str, Any]]) -> list[dict[str, str]]:
    filters: list[dict[str, str]] = []

    def add(field: str, value: str, reason: str, **extra: Any) -> None:
        filters.append({"field": field, "value": value, "reason": reason, **{key: val for key, val in extra.items() if val not in (None, "", [])}})

    def add_marketing_goal(value: str, reason: str) -> None:
        parent = str(value or "").strip()
        metric = PGY_MARKETING_GOAL_DEFAULT_METRIC.get(parent) or parent
        parent = PGY_MARKETING_GOAL_METRIC_PARENT.get(metric) or parent
        if not parent or not metric:
            return
        add(
            "营销目标",
            metric,
            reason,
            control_type="marketing_goal_metric",
            goal=parent,
            parent_value=parent,
            priority="low",
        )

    for item in hard_filters or []:
        if item.get("required") is False:
            continue
        field = str(item.get("field") or item.get("standard") or "")
        value = str(item.get("value") or "")
        pgy_field = str(item.get("pgyField") or "")
        text = f"{field} {pgy_field} {value}".lower()
        reason = f"硬性条件：{field}{item.get('condition') or ''}{value}"
        value_control = str(item.get("valueControl") or "")
        sub_field = str(item.get("subField") or "")
        if pgy_field == "营销目标" and value:
            for part in [part.strip() for part in re.split(r"[、,，/]+", value) if part.strip()]:
                add_marketing_goal(part, reason)
            continue
        if pgy_field in {"博主类目", "粉丝量", "粉丝年龄", "家庭身份", "职业身份", "特色背景", "母婴阶段", "地域", "粉丝地域"} and value:
            if pgy_field == "粉丝年龄":
                age_values = []
                if any(keyword in value for keyword in ["35", "34", "40", "家长", "父母", ">44", "44岁以上"]):
                    age_values.append("35～44 占比高")
                    if ">44" in value or "44岁以上" in value or "35岁以上" in value:
                        age_values.append(">44 占比高")
                for age_value in age_values or [value]:
                    add(pgy_field, age_value, reason, control_type="dropdown")
                continue
            if pgy_field == "粉丝量" and "、" in value:
                for part in [part.strip() for part in re.split(r"[、,，/]+", value) if part.strip()]:
                    add(pgy_field, part, reason, control_type="preset_or_number_range")
                continue
            if pgy_field in {"地域", "粉丝地域"}:
                region_items = _standard_region_filter_items({"field": pgy_field, "value": value, "reason": reason, "control_type": "three_level_cascade_checkbox_popover"})
                if region_items:
                    filters.extend(region_items)
                else:
                    add(pgy_field, value, reason, control_type="three_level_cascade_checkbox_popover")
                continue
            add(pgy_field, value, reason, control_type=value_control or "checkbox_popover")
            continue
        if pgy_field in {"合作报价", "预估阅读单价", "预估互动单价", "阅读中位数", "互动中位数", "曝光中位数"} and value:
            if pgy_field == "合作报价":
                parsed_min, parsed_max = _range_numbers_from_text(value)
                min_quote = item.get("min", parsed_min if parsed_min is not None else 1000)
                max_quote = item.get("max", parsed_max if parsed_max is not None else (_threshold_from_text(value, "quote", 20000) or 20000))
                next_sub_field = sub_field or ("视频笔记" if "视频笔记" in value else "图文笔记")
                display_value = value if next_sub_field in value else f"{next_sub_field}：{value}"
                add(pgy_field, display_value, reason, control_type="subfield_preset_or_number_range", sub_field=next_sub_field, min=min_quote, max=max_quote)
                continue
            if pgy_field == "预估阅读单价":
                cpc_max = _threshold_from_text(value, "cpc", 2) or _threshold_from_text(value, "generic", 2) or 2
                add(pgy_field, f"{sub_field or '图文笔记阅读单价'}≤{cpc_max:g}", reason, control_type="subfield_preset_or_number_range", sub_field=sub_field or "图文笔记阅读单价", max=cpc_max)
                continue
            if pgy_field == "预估互动单价":
                cpe_max = _threshold_from_text(value, "cpe", 20) or _threshold_from_text(value, "generic", 20) or 20
                add(pgy_field, f"{sub_field or '图文笔记互动单价'}≤{cpe_max:g}", reason, control_type="subfield_preset_or_number_range", sub_field=sub_field or "图文笔记互动单价", max=cpe_max)
                continue
            add(pgy_field, value, reason, control_type="subfield_preset_or_number_range", sub_field=sub_field or "图文笔记")
            continue
        if any(keyword in text for keyword in ["报价", "预算", "合作价格", "平台价格"]):
            max_quote = _threshold_from_text(value, "quote", 20000) or 20000
            add("合作报价", f"图文笔记：0.1万～{max_quote / 10000:g}万", reason, control_type="subfield_preset_or_number_range", sub_field="图文笔记", min=1000, max=max_quote)
        if any(keyword in text for keyword in ["35", "34", "粉丝年龄", "宝妈", "家长"]):
            add("粉丝年龄", "35～44 占比高", reason, control_type="dropdown")
        if any(keyword in text for keyword in ["cpc", "cpe", "阅读单价", "互动单价"]):
            cpc_max = _threshold_from_text(value, "cpc", 2) or 2
            cpe_max = _threshold_from_text(value, "cpe", 20) or 20
            add("预估阅读单价", f"图文笔记阅读单价≤{cpc_max:g}", reason, control_type="subfield_preset_or_number_range", sub_field="图文笔记阅读单价", max=cpc_max)
            add("预估互动单价", f"图文笔记互动单价≤{cpe_max:g}", reason, control_type="subfield_preset_or_number_range", sub_field="图文笔记互动单价", max=cpe_max)
        if any(keyword in text for keyword in ["孩子年级", "小升初", "初中", "高中", "大孩"]):
            add("内容场景", "小升初/初中/高中", reason)
        if _brief_emphasizes_region(f"{field} {value}"):
            region_items = _standard_region_filter_items({"field": "地域", "value": value or "北京/上海优先", "reason": reason, "control_type": "three_level_cascade_checkbox_popover"})
            if region_items:
                filters.extend(region_items)
            else:
                add("地域", value or "北京/上海优先", reason, control_type="three_level_cascade_checkbox_popover")
        if any(keyword in text for keyword in ["限流", "违规", "流量稳定", "异常"]):
            add("常规剔除", "剔除低活博主", reason, control_type="checkbox")
            add("常规剔除", "剔除掉粉博主", reason, control_type="checkbox")
    return _merge_filters(filters)


def _normalize_pgy_filter_item(item: dict[str, Any]) -> dict[str, Any]:
    field = str(item.get("field") or "")
    value = str(item.get("value") or "")
    normalized = {**item, "field": field, "value": value, "reason": str(item.get("reason") or "")}
    if field == "博主类目":
        main_value = value.strip()
        sub_value = str(item.get("sub_value") or item.get("subValue") or "").strip()
        normalized.pop("sub_value", None)
        normalized.pop("subValue", None)
        if main_value not in PGY_BLOGGER_CATEGORY_TAXONOMY:
            for category, subcategories in PGY_BLOGGER_CATEGORY_TAXONOMY.items():
                if main_value in subcategories:
                    main_value, sub_value = category, main_value
                    break
        valid_subcategories = PGY_BLOGGER_CATEGORY_TAXONOMY.get(main_value) or []
        if sub_value and sub_value not in valid_subcategories:
            sub_value = ""
        return {
            **normalized,
            "value": main_value,
            "control_type": "tag_select_with_hover_subcategory",
            **({"sub_value": sub_value} if sub_value else {}),
        }
    if field == "营销目标":
        parent = str(item.get("goal") or item.get("parent_value") or item.get("parentValue") or "").strip()
        metric = value.strip()
        if metric in PGY_MARKETING_GOAL_DEFAULT_METRIC:
            parent = metric
            metric = PGY_MARKETING_GOAL_DEFAULT_METRIC.get(parent) or metric
        if not parent:
            parent = PGY_MARKETING_GOAL_METRIC_PARENT.get(metric) or ""
        return {
            **normalized,
            "value": metric,
            "control_type": "marketing_goal_metric",
            "goal": parent,
            "parent_value": parent,
            "priority": normalized.get("priority") or "low",
        }
    if field == "博主人设":
        mapping = {
            "家庭身份": {"field": "家庭身份", "value": "妈妈", "control_type": "checkbox_popover"},
            "职业身份": {"field": "职业身份", "value": "学生", "control_type": "checkbox_popover"},
            "特色背景": {"field": "特色背景", "value": "留学背景", "control_type": "checkbox_popover"},
        }
        return {**normalized, **mapping.get(value, {})}
    if field == "数据表现" and value in {"预估阅读/互动单价", "CPC<2/CPE<20"}:
        return {
            **normalized,
            "field": "预估阅读单价",
            "value": "图文笔记阅读单价≤2",
            "control_type": "subfield_preset_or_number_range",
            "sub_field": "图文笔记阅读单价",
            "max": 2,
            "pending_detail": "预估互动单价需作为独立条件补充",
        }
    if field == "报价":
        return {**normalized, "field": "合作报价", "value": "0.1万～2万", "control_type": "subfield_preset_or_number_range", "sub_field": "", "sub_fields": ["图文笔记", "视频笔记"], "min": 1000, "max": 20000}
    if field == "合作报价":
        parsed_min, parsed_max = _range_numbers_from_text(value)
        sub_names = _subfield_names_from_item(normalized, default=["图文笔记", "视频笔记"])
        sub_field = sub_names[0] if len(sub_names) == 1 else ""
        has_explicit_min = normalized.get("min") not in (None, "")
        has_explicit_max = normalized.get("max") not in (None, "")
        min_value = normalized.get("min") if has_explicit_min else parsed_min
        max_value = normalized.get("max") if has_explicit_max else parsed_max
        if min_value is None and max_value is None:
            min_value, max_value = 1000, 20000
        elif min_value is None:
            min_value = ""
        elif max_value is None:
            max_value = ""
        return {
            **normalized,
            "control_type": normalized.get("control_type") or "subfield_preset_or_number_range",
            "sub_field": sub_field,
            "sub_fields": sub_names,
            "min": min_value,
            "max": max_value,
        }
    if field in PGY_MIN_ONLY_RANGE_FIELDS:
        parsed_min, parsed_max = _range_numbers_from_text(value)
        min_value = normalized.get("min", parsed_min if parsed_min is not None else parsed_max)
        return {
            **normalized,
            "control_type": normalized.get("control_type") or ("number_range" if field == "合作订单数" else "preset_or_number_range"),
            "min": min_value if min_value not in (None, "") else "",
            "max": "",
            "range_policy": "min_only",
        }
    if field in PGY_MIN_ONLY_SUBFIELD_RANGE_FIELDS:
        return {
            **normalized,
            "control_type": normalized.get("control_type") or ("subfield_preset_or_percent_range" if field == "合作信用度" else "multi_subfield_preset_or_number_range"),
            "range_policy": "min_only",
        }
    if field in PGY_MAX_ONLY_RANGE_FIELDS:
        parsed_min, parsed_max = _range_numbers_from_text(value)
        max_value = normalized.get("max", parsed_max if parsed_max is not None else parsed_min)
        return {
            **normalized,
            "control_type": normalized.get("control_type") or "preset_or_number_range",
            "min": "",
            "max": max_value if max_value not in (None, "") else "",
            "range_policy": "max_only",
        }
    if field in PGY_MAX_ONLY_SUBFIELD_RANGE_FIELDS:
        parsed_min, parsed_max = _range_numbers_from_text(value)
        max_value = normalized.get("max", parsed_max if parsed_max is not None else parsed_min)
        return {
            **normalized,
            "control_type": normalized.get("control_type") or "subfield_preset_or_number_range",
            "min": "",
            "max": max_value if max_value not in (None, "") else "",
            "range_policy": "max_only",
        }
    if field == "粉丝年龄" and value in {"35岁以上优先", "35岁以上≥40%"}:
        return {**normalized, "value": "35～44 占比高", "control_type": "dropdown"}
    if field == "粉丝年龄":
        return {**normalized, "control_type": normalized.get("control_type") or "dropdown"}
    if field in {"地域", "粉丝地域"}:
        return {**normalized, "control_type": "three_level_cascade_checkbox_popover"}
    if field == "常规剔除" and value in {"低风险/流量稳定", "规避限流异常", "流量稳定"}:
        return {**normalized, "value": "剔除低活博主", "control_type": "checkbox"}
    return normalized


def _standard_region_filter_items(item: dict[str, Any]) -> list[dict[str, Any]]:
    field = str(item.get("field") or "")
    if field not in {"地域", "粉丝地域"}:
        return []
    raw_value = str(item.get("value") or "").strip()
    targets: list[str] = []
    if raw_value in PGY_REGION_ALIASES:
        targets.extend(PGY_REGION_ALIASES[raw_value])
    elif raw_value.replace(" ", "") in PGY_REGION_ALIASES:
        targets.extend(PGY_REGION_ALIASES[raw_value.replace(" ", "")])
    elif raw_value.startswith(("中国：", "中国:", "中国-")):
        targets.append(re.sub(r"^中国[：:-]", "", raw_value).strip())
    elif any(sep in raw_value for sep in ["、", ",", "，", "/", "|", "｜"]):
        for part in re.split(r"[、,，/|｜\s]+", raw_value):
            clean = part.strip(" ：:;；。优先重点必须地域要求地区城市IPip")
            if clean in PGY_REGION_ALIASES:
                targets.extend(PGY_REGION_ALIASES[clean])
            elif clean in PGY_FOREIGN_REGION_OPTIONS or clean in PGY_CHINA_REGION_OPTIONS or clean in {"广州", "深圳"}:
                targets.append(clean)
    elif raw_value and raw_value not in {"中国", "国内", "全国", "不限"}:
        targets.append(raw_value)
    for key in ("city", "province", "country"):
        value = str(item.get(key) or "").strip()
        if value and value not in {"中国", "国内", "全国", "不限"}:
            targets.append(value)
    result: list[dict[str, Any]] = []
    for target in dict.fromkeys([value for value in targets if value]):
        country = "中国"
        province = str(item.get("province") or "")
        city = str(item.get("city") or "")
        level = str(item.get("level") or "province")
        if target in PGY_FOREIGN_REGION_OPTIONS:
            country, province, city, level = target, "", "", "country"
        else:
            province = province or {"广州": "广东", "深圳": "广东"}.get(target, target)
        result.append(
            {
                **item,
                "field": field,
                "value": target,
                "country": country,
                "province": province,
                **({"city": city} if city else {}),
                "level": level,
                "control_type": "three_level_cascade_checkbox_popover",
                "label": "",
            }
        )
    return result


def _normalize_pgy_filters(filters: list[dict[str, Any]]) -> list[dict[str, Any]]:
    normalized: list[dict[str, Any]] = []
    for item in filters or []:
        region_items = _standard_region_filter_items(item)
        if region_items:
            normalized.extend(region_items)
            continue
        current = _normalize_pgy_filter_item(item)
        normalized.append(current)
        original_field = str(item.get("field") or "")
        original_value = str(item.get("value") or "")
        if original_field == "数据表现" and original_value in {"预估阅读/互动单价", "CPC<2/CPE<20"}:
            normalized.append(
                {
                    "field": "预估互动单价",
                    "value": "图文笔记互动单价≤20",
                    "reason": str(item.get("reason") or ""),
                    "control_type": "subfield_preset_or_number_range",
                    "sub_field": "图文笔记互动单价",
                    "max": 20,
                }
            )
        if original_field == "常规剔除" and original_value in {"低风险/流量稳定", "规避限流异常", "流量稳定"}:
            normalized.append(
                {
                    "field": "常规剔除",
                    "value": "剔除掉粉博主",
                    "reason": str(item.get("reason") or ""),
                    "control_type": "checkbox",
                }
            )
    return _merge_filters(normalized)


def build_collection_plan(brief: str = "", screening_plan: dict[str, Any] | None = None) -> dict[str, Any]:
    if isinstance(screening_plan, dict) and isinstance(screening_plan.get("pgyCollectionPlan"), dict):
        plan = screening_plan["pgyCollectionPlan"]
        if "collectionHardFilters" in screening_plan:
            hard_filters = screening_plan.get("collectionHardFilters") or []
        elif "hard_filters" in plan or "hardFilters" in plan:
            hard_filters = plan.get("hard_filters") or plan.get("hardFilters") or []
        else:
            hard_filters = screening_plan.get("hardFilters") or []
        schemes = plan.get("schemes") or []
        default_filters = plan.get("filters") or []
        if not default_filters and schemes and isinstance(schemes[0], dict):
            default_filters = schemes[0].get("filters") or []
        active_scheme = bool(plan.get("active_scheme_id"))
        derived_filters = [] if active_scheme else _hard_filters_to_pgy_filters(hard_filters)
        metrics = plan.get("display_metrics") or PGY_DISPLAY_METRICS
        if PGY_ALL_NON_LIVE_METRICS not in metrics:
            metrics = [PGY_ALL_NON_LIVE_METRICS]
        return {
            "strategy": plan.get("strategy") or "",
            "target_count_range": plan.get("target_count_range") or "",
            "schemes": schemes,
            "filters": _normalize_pgy_filters(_merge_filters(derived_filters, default_filters)),
            "hard_filters": hard_filters,
            "display_metrics": metrics,
            "detail_fields": plan.get("detail_fields") or ["基础画像", "粉丝画像", "报价", "合作表现", "内容表现"],
            "filter_catalog": plan.get("filter_catalog") or PGY_FILTER_CATALOG,
        }

    text = f"{brief} {screening_plan or ''}".lower()
    filters: list[dict[str, str]] = []

    def add(field: str, value: str, reason: str, **extra: Any) -> None:
        sub_value = str(extra.get("sub_value") or extra.get("subValue") or "")
        if not any(item["field"] == field and item["value"] == value and str(item.get("sub_value") or item.get("subValue") or "") == sub_value for item in filters):
            filters.append({"field": field, "value": value, "reason": reason, **{key: val for key, val in extra.items() if val not in (None, "", [])}})

    if any(keyword in text for keyword in ["曝光", "声量", "阅读", "播放"]):
        add("营销目标", "曝光表现", "Brief 提到曝光/声量目标", control_type="marketing_goal_metric", goal="曝光", parent_value="曝光", priority="low")
    if any(keyword in text for keyword in ["种草", "口碑", "测评", "内容"]):
        add("营销目标", "互动表现", "Brief 提到种草或内容测评", control_type="marketing_goal_metric", goal="种草", parent_value="种草", priority="low")
    if any(keyword in text for keyword in ["转化", "销售", "进店", "下单"]):
        add("营销目标", "外溢进店表现", "Brief 提到转化目标", control_type="marketing_goal_metric", goal="转化", parent_value="转化", priority="low")

    category_map = [
        ("教育", ["教育", "学习", "升学", "初中", "高中", "答疑", "教辅", "老师"]),
        ("母婴", ["母婴", "亲子", "妈妈", "孩子", "儿童", "大孩"]),
        ("家居家装", ["家居", "家装", "装修"]),
        ("科技数码", ["数码", "智能", "硬件", "电子"]),
    ]
    for category, keywords in category_map:
        if any(keyword in text for keyword in keywords):
            sub_values: list[str] = []
            if category == "教育":
                if any(keyword in text for keyword in ["家庭教育", "家长", "父母", "亲子", "大孩", "小升初", "初中", "高中"]):
                    sub_values.append("家庭教育")
                if any(keyword in text for keyword in ["k12", "K12", "小升初", "初中", "高中", "小学", "教辅", "答疑"]):
                    sub_values.append("k12教育")
                if any(keyword in text for keyword in ["学习日常", "学习博主", "学霸", "学习效率", "学习工具"]):
                    sub_values.append("学习日常")
            if category == "母婴":
                if any(keyword in text for keyword in ["育儿", "陪伴", "家长", "父母", "大孩", "小升初", "初中", "高中"]):
                    sub_values.append("育儿经验")
                if any(keyword in text for keyword in ["早教", "启蒙"]):
                    sub_values.append("早教")
                if any(keyword in text for keyword in ["日常", "家庭"]):
                    sub_values.append("母婴日常")
            if sub_values:
                for sub_value in sub_values:
                    add("博主类目", category, f"Brief 命中 {category}-{sub_value} 场景", sub_value=sub_value)
            else:
                add("博主类目", category, f"Brief 命中 {category} 场景")

    persona_map = [
        ("家庭身份", ["家庭", "亲子", "妈妈", "爸爸", "孩子"]),
        ("职业身份", ["老师", "教师", "专家", "医生", "博士"]),
        ("特色背景", ["高知", "中产", "精英", "升学"]),
    ]
    for field, keywords in persona_map:
        if any(keyword in text for keyword in keywords):
            add("博主人设", field, f"Brief 命中 {field} 画像")
    if any(keyword in text for keyword in ["妈妈", "宝妈", "母亲", "亲子", "萌娃", "爸爸", "父母"]):
        add("家庭身份", "妈妈", "Brief 命中家庭/亲子身份", control_type="checkbox_popover")
    if any(keyword in text for keyword in ["学生", "大学生", "留学生", "校园", "高校"]):
        add("职业身份", "学生", "Brief 命中学生身份", control_type="checkbox_popover")
    if any(keyword in text for keyword in ["留学", "海外留学", "留学生", "出国", "雅思", "托福"]):
        add("特色背景", "留学背景", "Brief 命中留学/海外背景", control_type="checkbox_popover")
    elif any(keyword in text for keyword in ["高知", "升学", "备考", "考研"]):
        add("特色背景", "备考经验", "Brief 命中特色背景", control_type="checkbox_popover")

    if _brief_emphasizes_region(brief):
        add("地域", "北京", "Brief 明确强调地域/IP/城市要求", control_type="three_level_cascade_checkbox_popover", country="中国", province="北京", level="province")
        add("地域", "上海", "Brief 明确强调地域/IP/城市要求", control_type="three_level_cascade_checkbox_popover", country="中国", province="上海", level="province")
    if any(keyword in text for keyword in ["35岁", "35 岁", "34岁", "家长", "父母"]):
        add("粉丝年龄", "35～44 占比高", "Brief 要求家长/35岁以上粉丝", control_type="dropdown")
    if any(keyword in text for keyword in ["cpe", "cpc", "阅读单价", "互动单价"]):
        add("预估阅读单价", "图文笔记阅读单价≤2", "Brief 要求控制 CPC", control_type="subfield_preset_or_number_range", sub_field="图文笔记阅读单价", max=2)
        add("预估互动单价", "图文笔记互动单价≤20", "Brief 要求控制 CPE", control_type="subfield_preset_or_number_range", sub_field="图文笔记互动单价", max=20)
    cooperation_brands = _extract_brands_after_labels(brief, ["近期合作品牌", "合作品牌", "已合作品牌"])
    competitor_brands = _extract_brands_after_labels(brief, ["竞品", "竞品品牌", "对标品牌", "类似品牌", "参考品牌"])
    recommendation_brands = _extract_brands_after_labels(brief, ["按博主粉丝推荐", "人群目标品牌", "目标品牌", "合作品牌", "竞品", "对标品牌", "类似品牌", "参考品牌"])
    if cooperation_brands or competitor_brands:
        brands = list(dict.fromkeys([*cooperation_brands, *competitor_brands]))
        add(
            "近期合作品牌",
            "、".join(brands[:6]) or "待补品牌",
            "Brief 提到近期合作品牌/竞品，需要记录品牌并可剔除已合作博主",
            control_type="searchable_multi_select_with_exclude",
            input_values=brands[:6],
            min_items=3,
            pending_detail="至少补足3个品牌",
        )
    if recommendation_brands or any(keyword in text for keyword in ["按博主粉丝推荐", "人群目标", "竞品粉丝", "品牌粉丝", "对标品牌"]):
        brands = recommendation_brands[:6]
        add(
            "按博主粉丝推荐",
            "、".join(brands) or "待选择合作品牌/竞品",
            "Brief 提到品牌/竞品粉丝人群，可通过右上角智能推荐博主入口搜索品牌或竞品",
            control_type="brand_search_recommendation",
            input_values=brands,
            competitor_values=competitor_brands[:6],
            pending_detail="在右上角合作品牌搜索框输入品牌名或ID；竞品也可作为搜索品牌输入",
        )
    if any(keyword in text for keyword in ["行业推荐", "推荐博主", "意向行业"]):
        add("行业推荐博主", "我的行业", "Brief 提到行业推荐/行业匹配", control_type="nested_select_popover", pending_detail="打开后继续选择我的行业")

    source_plan = screening_plan or {}
    hard_filters = (
        source_plan.get("collectionHardFilters") or []
        if "collectionHardFilters" in source_plan
        else source_plan.get("hardFilters") or []
    )
    return {
        "filters": _normalize_pgy_filters(_merge_filters(_hard_filters_to_pgy_filters(hard_filters), filters)),
        "hard_filters": hard_filters,
        "display_metrics": PGY_DISPLAY_METRICS,
        "detail_fields": ["基础画像", "粉丝画像", "报价", "合作表现", "内容表现"],
        "filter_catalog": PGY_FILTER_CATALOG,
    }


PGY_EMPTY_RESULT_HINTS = [
    "暂未找到相关博主",
    "暂未发现相关博主",
    "试试按行业找博主",
    "放宽条件才能找到更多的博主",
    "放宽条件才能找到更多博主",
    "未发现相关博主",
    "没有找到相关博主",
    "暂无相关博主",
    "暂无数据",
]

PGY_BROAD_KEEP_FIELDS = {"营销目标", "博主类目"}
PGY_INVALID_ROW_NAMES = {"暂无数据", "暂无相关博主", "暂未找到相关博主", "暂未发现相关博主", "没有找到相关博主"}
PGY_NO_ORDER_PERMISSION_HINTS = (
    "无接单权限",
    "暂无接单权限",
    "不可接单",
    "不能接单",
    "暂不接单",
    "未开通接单",
)


def _visible_text_lines(text: str) -> list[str]:
    return [line.strip() for line in text.splitlines() if line.strip()]


def _clean_text(value: str) -> str:
    return re.sub(r"\s+", "", value or "")


def _pgy_no_order_permission_hint(*values: Any) -> str:
    for value in values:
        if value in (None, ""):
            continue
        if isinstance(value, (dict, list)):
            text = json.dumps(value, ensure_ascii=False)
        else:
            text = str(value)
        cleaned = _clean_text(text)
        if any(hint in cleaned for hint in PGY_NO_ORDER_PERMISSION_HINTS):
            return "无接单权限"
    return ""


def _page_empty_result_hint(page: Any) -> str:
    try:
        text = page.locator("body").inner_text(timeout=2500)
    except Exception:
        return ""
    normalized = _clean_text(text)
    for hint in PGY_EMPTY_RESULT_HINTS:
        if hint in normalized:
            return hint
    return ""


def _relaxed_collection_plan(plan: dict[str, Any], stage: str) -> dict[str, Any] | None:
    filters = [item for item in plan.get("filters") or [] if isinstance(item, dict)]
    if not filters:
        return None
    if stage == "broad":
        kept = [item for item in filters if str(item.get("field") or "") in PGY_BROAD_KEEP_FIELDS]
        if len(kept) == len(filters):
            return None
        strategy = "页面无结果，自动放宽为基础类目/目标条件"
    elif stage == "unfiltered":
        kept = []
        strategy = "基础类目/目标仍无结果，清空蒲公英页面筛选后采集，再由本地筛选工作台标记"
    else:
        return None
    removed = [
        item
        for item in filters
        if not any(
            str(item.get("field") or "") == str(next_item.get("field") or "")
            and str(item.get("value") or "") == str(next_item.get("value") or "")
            for next_item in kept
        )
    ]
    return {
        **plan,
        "filters": kept,
        "auto_relaxed": True,
        "relaxation": {
            "stage": stage,
            "strategy": strategy,
            "original_filter_count": len(filters),
            "relaxed_filter_count": len(kept),
            "removed_filters": removed,
            "kept_filters": kept,
        },
    }


def _merge_plan_results(base: dict[str, Any], retry: dict[str, Any], relaxation: dict[str, Any]) -> dict[str, Any]:
    removed_skips = [
        {**item, "message": relaxation.get("strategy") or "页面无结果，自动放宽该筛选项"}
        for item in relaxation.get("removed_filters") or []
    ]
    return {
        "applied_filters": retry.get("applied_filters") or [],
        "skipped_filters": [
            *(base.get("skipped_filters") or []),
            *removed_skips,
            *(retry.get("skipped_filters") or []),
        ],
        "selected_metrics": retry.get("selected_metrics") or base.get("selected_metrics") or [],
        "skipped_metrics": [
            *(base.get("skipped_metrics") or []),
            *(retry.get("skipped_metrics") or []),
        ],
    }


def _is_visible(locator: Any) -> bool:
    try:
        box = locator.bounding_box(timeout=700)
        return bool(box and box.get("width", 0) > 0 and box.get("height", 0) > 0)
    except Exception:
        return False


def _click_first_visible(locator: Any, timeout: int = 1500) -> bool:
    try:
        count = min(locator.count(), 30)
    except Exception:
        return False
    for index in range(count):
        candidate = locator.nth(index)
        if not _is_visible(candidate):
            continue
        try:
            candidate.scroll_into_view_if_needed(timeout=timeout)
            candidate.click(timeout=timeout)
            return True
        except Exception:
            continue
    return False


def _click_locator(page: Any, locator: Any, timeout: int = 1500) -> bool:
    try:
        locator.scroll_into_view_if_needed(timeout=timeout)
    except Exception:
        pass
    try:
        locator.click(timeout=timeout)
        return True
    except Exception:
        pass
    try:
        locator.evaluate("node => node.click()")
        return True
    except Exception:
        pass
    try:
        box = locator.bounding_box(timeout=800)
        if not box:
            return False
        page.mouse.click(box["x"] + box["width"] / 2, box["y"] + box["height"] / 2)
        return True
    except Exception:
        return False


def _selected_filter_text(page: Any) -> str:
    for selector in [
        ".selected-box",
        ".filter-selected",
        ".selected-filter",
        "[class*='selected'][class*='box']",
        "[class*='filter'][class*='selected']",
    ]:
        try:
            selected = page.locator(selector).first
            if selected.count():
                text = selected.inner_text(timeout=1200)
                if text:
                    return text
        except Exception:
            continue
    try:
        text = page.evaluate(
            """
            () => {
              const visible = el => {
                const style = window.getComputedStyle(el);
                const box = el.getBoundingClientRect();
                return style.visibility !== 'hidden' && style.display !== 'none' && box.width > 0 && box.height > 0;
              };
              const nodes = Array.from(document.querySelectorAll('*'))
                .filter(el => visible(el))
                .map(el => ({ el, text: (el.innerText || el.textContent || '').trim() }))
                .filter(item => item.text.includes('重置') && (item.text.includes('存为常用筛选') || item.text.includes('不限') || item.text.includes('已选')));
              nodes.sort((a, b) => a.text.length - b.text.length);
              return nodes[0]?.text || '';
            }
            """
        )
        return str(text or "")
    except Exception:
        return ""


def _item_selection_value(item: dict[str, Any]) -> str:
    value = str(item.get("value") or "")
    sub_field = str(item.get("sub_field") or item.get("subField") or "")
    if sub_field:
        for sep in ("：", ":"):
            prefix = f"{sub_field}{sep}"
            if value.startswith(prefix):
                return value[len(prefix):].strip()
    return value.strip()


def _filter_value_candidates(item: dict[str, Any]) -> list[str]:
    raw_value = str(item.get("value") or "").strip()
    selection_value = _item_selection_value(item)
    candidates = [selection_value, raw_value]
    sub_fields = item.get("sub_fields") or item.get("subFields") or []
    if not isinstance(sub_fields, list):
        sub_fields = [sub_fields] if sub_fields else []
    max_value = item.get("max")
    min_value = item.get("min")
    for sub_field in [str(part).strip() for part in sub_fields if str(part).strip()]:
        candidates.append(sub_field)
        if max_value not in (None, ""):
            candidates.extend([f"{sub_field}≤{max_value:g}" if isinstance(max_value, (int, float)) else f"{sub_field}≤{max_value}", f"{sub_field}{max_value}以下"])
        if min_value not in (None, ""):
            candidates.extend([f"{sub_field}≥{min_value:g}" if isinstance(min_value, (int, float)) else f"{sub_field}≥{min_value}", f"{sub_field}{min_value}以上"])
    sub_value = str(item.get("sub_value") or item.get("subValue") or "").strip()
    if str(item.get("field") or "") == "博主类目" and sub_value:
        candidates = [sub_value, f"{selection_value}-{sub_value}", f"{selection_value}/{sub_value}", f"{selection_value}：{sub_value}", selection_value, raw_value]
    if str(item.get("field") or "") == "营销目标":
        parent = str(item.get("goal") or item.get("parent_value") or item.get("parentValue") or "").strip()
        if selection_value in PGY_MARKETING_GOAL_DEFAULT_METRIC:
            candidates.append(PGY_MARKETING_GOAL_DEFAULT_METRIC[selection_value])
        if parent:
            candidates.append(parent)
    for value in [selection_value, raw_value]:
        candidates.extend(PGY_FILTER_ALIASES.get(value) or [])
    return list(dict.fromkeys([candidate for candidate in candidates if candidate]))


def _filter_threshold_candidates(item: dict[str, Any]) -> list[str]:
    candidates: list[str] = []
    for key, suffixes in [("max", ["以下", "以内"]), ("min", ["以上", "不低于"])]:
        value = item.get(key)
        if value in (None, ""):
            continue
        number = _number_from_text(str(value))
        if number is None:
            text = str(value)
        else:
            text = f"{number:g}"
        candidates.append(text)
        if key == "max":
            candidates.extend([f"≤{text}", f"<={text}", f"不限-{text}", f"不限～{text}"])
        else:
            candidates.extend([f"≥{text}", f">={text}", f"{text}-不限", f"{text}～不限"])
        candidates.extend([f"{text}{suffix}" for suffix in suffixes])
    return list(dict.fromkeys([_clean_text(candidate) for candidate in candidates if candidate]))


def _subfield_selected_in_text(selected_text: str, field: str, sub_field: str) -> bool:
    aliases = []
    try:
        aliases = _subfield_aliases(field, sub_field)
    except Exception:
        aliases = [sub_field]
    aliases = [_clean_text(item) for item in aliases if item]
    return any(alias and alias in selected_text for alias in aliases)


def _catalog_option_groups(field: str) -> list[dict[str, Any]]:
    for catalog_item in PGY_FILTER_CATALOG:
        if catalog_item.get("field") == field:
            groups = catalog_item.get("option_groups") or catalog_item.get("optionGroups") or []
            return groups if isinstance(groups, list) else []
    return []


def _infer_option_group(field: str, value: str) -> str:
    for group in _catalog_option_groups(field):
        options = group.get("options") or []
        if value in options:
            return str(group.get("label") or "")
    return ""


def _filter_already_selected(page: Any, item: dict[str, str]) -> bool:
    selected_text = _clean_text(_selected_filter_text(page))
    field = _clean_text(item.get("field") or "")
    sub_value = _clean_text(str(item.get("sub_value") or item.get("subValue") or ""))
    if field == "博主类目" and sub_value:
        return bool(selected_text and sub_value in selected_text)
    sub_fields = item.get("sub_fields") or item.get("subFields") or []
    if isinstance(sub_fields, list) and len(sub_fields) > 1:
        normalized_sub_fields = [_clean_text(str(part)) for part in sub_fields if str(part)]
        has_all_subfields = all(_subfield_selected_in_text(selected_text, field, part) for part in normalized_sub_fields)
        threshold_candidates = _filter_threshold_candidates(item)
        has_threshold = not threshold_candidates or any(candidate in selected_text for candidate in threshold_candidates)
        return bool(selected_text and field in selected_text and has_all_subfields and has_threshold)
    single_sub_field = str(item.get("sub_field") or item.get("subField") or "").strip()
    if not single_sub_field and isinstance(sub_fields, list) and len(sub_fields) == 1:
        single_sub_field = str(sub_fields[0] or "").strip()
    threshold_candidates = _filter_threshold_candidates(item)
    if single_sub_field:
        has_subfield = _subfield_selected_in_text(selected_text, field, single_sub_field)
        has_threshold = not threshold_candidates or any(candidate in selected_text for candidate in threshold_candidates)
        return bool(selected_text and field in selected_text and has_subfield and has_threshold)
    if selected_text and field in selected_text and threshold_candidates and any(candidate in selected_text for candidate in threshold_candidates):
        return True
    values = [_clean_text(candidate) for candidate in _filter_value_candidates(item)]
    values = [value for value in values if value]
    if not selected_text or not values:
        return False
    if any(value in selected_text for value in values):
        return True
    return bool(field and field in selected_text and any(value and value in selected_text for value in values))


def _blogger_category_parent_all_selected(page: Any, main_value: str) -> bool:
    selected_text = _clean_text(_selected_filter_text(page))
    main_text = _clean_text(main_value)
    if not selected_text or not main_text or "博主类目" not in selected_text:
        return False
    return any(
        marker in selected_text
        for marker in [
            f"{main_text}-全部",
            f"{main_text}/全部",
            f"{main_text}：全部",
            f"{main_text}:全部",
        ]
    )


def _remove_selected_filter_chip(page: Any, required_parts: list[str]) -> bool:
    parts = [_clean_text(part) for part in required_parts if _clean_text(part)]
    if not parts:
        return False
    try:
        removed = page.evaluate(
            """
            parts => {
              const clean = text => String(text || '').replace(/\\s+/g, '');
              const visible = el => {
                const style = window.getComputedStyle(el);
                const box = el.getBoundingClientRect();
                return style.visibility !== 'hidden' && style.display !== 'none' && box.width > 0 && box.height > 0;
              };
              const closeSelector = [
                '[class*="close"]',
                '[class*="Close"]',
                '[class*="icon-close"]',
                '[aria-label*="关闭"]',
                '[aria-label*="删除"]',
                'svg',
                'i',
                'button'
              ].join(',');
              const candidates = Array.from(document.querySelectorAll('*'))
                .filter(visible)
                .map(el => ({ el, text: clean(el.innerText || el.textContent || '') }))
                .filter(item => item.text && parts.every(part => item.text.includes(part)))
                .filter(item => item.text.length <= 80 || /tag|chip|selected|filter/i.test(String(item.el.className || '')))
                .sort((a, b) => a.text.length - b.text.length);
              for (const { el } of candidates) {
                const close = Array.from(el.querySelectorAll(closeSelector))
                  .filter(visible)
                  .sort((a, b) => b.getBoundingClientRect().left - a.getBoundingClientRect().left)[0];
                if (close) {
                  close.click();
                  return true;
                }
                if (/×|x/i.test(el.textContent || '')) {
                  el.click();
                  return true;
                }
              }
              return false;
            }
            """,
            parts,
        )
    except Exception:
        removed = False
    if removed:
        try:
            page.wait_for_timeout(500)
        except Exception:
            pass
    return bool(removed)


def _verify_filter_selected(page: Any, item: dict[str, Any]) -> bool:
    control_type = str(item.get("control_type") or "")
    if control_type in {"brand_search_recommendation"}:
        return True
    page.wait_for_timeout(500)
    return _filter_already_selected(page, item)


def _filter_acceptance_required(item: dict[str, Any]) -> bool:
    field = str(item.get("field") or "")
    control_type = str(item.get("control_type") or "")
    if field in {"博主类目", "粉丝量", "粉丝年龄", "合作报价", "互动中位数", "阅读中位数", "预估阅读单价", "预估互动单价"}:
        return True
    return control_type in {
        "tag",
        "tag_select_with_hover_subcategory",
        "dropdown",
        "select_popover",
        "single_select_popover",
        "dropdown_single",
        "checkbox_popover",
        "preset_or_number_range",
        "preset_or_percent_range",
        "subfield_preset_or_number_range",
        "subfield_preset_or_percent_range",
        "multi_subfield_preset_or_number_range",
    }


def _parse_row_text(text: str, page_url: str, table_payload: dict[str, Any] | None = None) -> dict[str, Any] | None:
    lines = _visible_text_lines(text)
    if not lines or any("skeleton-block" in line for line in lines):
        return None
    cleaned_text = _clean_text(text)
    if any(hint in cleaned_text for hint in PGY_EMPTY_RESULT_HINTS):
        return None
    stop_words = {"添加合作", "发起邀约", "更多操作"}
    lines = [line for line in lines if line not in {"合作", "\t"}]
    if not lines:
        return None
    nickname = lines[0]
    if nickname in stop_words or nickname in PGY_INVALID_ROW_NAMES or len(nickname) > 60:
        return None

    quote_index = next((i for i, line in enumerate(lines) if line in {"¥", "￥"}), -1)
    quote_price = _number_from_text(lines[quote_index + 1]) if quote_index >= 0 and quote_index + 1 < len(lines) else None
    before_quote = lines[:quote_index] if quote_index >= 0 else lines
    row_metrics: dict[str, str] = {}
    metric_label_values: dict[str, str] = {}
    metric_start = -1
    for index in range(2, len(before_quote)):
        value = before_quote[index]
        if re.fullmatch(r"\d+(?:\.\d+)?\s*(?:w|万)", value, flags=re.I):
            metric_start = index
            break
    if metric_start >= 0:
        remaining_metric_count = len(before_quote) - metric_start
        metric_fields = (
            PGY_ROW_METRIC_FIELDS_WITH_LIVE
            if remaining_metric_count >= len(PGY_ROW_METRIC_FIELDS_WITH_LIVE)
            else PGY_NON_LIVE_ROW_METRIC_FIELDS
        )
        metric_values = before_quote[metric_start : metric_start + len(metric_fields)]
        profile_lines = before_quote[:metric_start]
        for (field, label), value in zip(metric_fields, metric_values):
            row_metrics[field] = value
            metric_label_values[label] = value
    else:
        metrics = [line for line in before_quote if _number_from_text(line) is not None or line == "--"]
        followers = metrics[-3] if len(metrics) >= 3 else None
        read_median = metrics[-2] if len(metrics) >= 2 else None
        interaction_median = metrics[-1] if len(metrics) >= 1 else None
        profile_lines = before_quote[: max(1, len(before_quote) - len(metrics))]
        if followers is not None:
            row_metrics["followers_count"] = followers
            metric_label_values["粉丝数"] = followers
        if read_median is not None:
            row_metrics["daily_read_median"] = read_median
            metric_label_values["阅读中位数（日常）"] = read_median
        if interaction_median is not None:
            row_metrics["daily_interaction_median"] = interaction_median
            metric_label_values["互动中位数（日常）"] = interaction_median
    followers = _number_from_text(row_metrics.get("followers_count") or "")
    read_median = _number_from_text(row_metrics.get("daily_read_median") or "")
    interaction_median = _number_from_text(row_metrics.get("daily_interaction_median") or "")
    location = profile_lines[1] if len(profile_lines) > 1 else ""
    tags = [line for line in profile_lines[2:] if not line.startswith("期待与") and line not in stop_words]
    cooperation_hint = next((line for line in profile_lines if line.startswith("期待与")), "")
    table_payload = table_payload or {}
    if not any([followers, read_median, interaction_median, table_payload.get("pgy_url"), table_payload.get("xiaohongshu_id")]) and len(lines) <= 3:
        return None
    raw_payload = {
        "text": text,
        "lines": lines,
        "metrics": metric_label_values,
        "read_median": read_median,
        "interaction_median": interaction_median,
        "cooperation_hint": cooperation_hint,
    }
    if table_payload.get("raw_table"):
        raw_payload["raw_table"] = table_payload["raw_table"]
    no_order_permission = _pgy_no_order_permission_hint(text, table_payload)
    if no_order_permission:
        raw_payload["order_permission_status"] = no_order_permission
    creator_id_seed = f"pgy:list:{nickname}:{location}"
    creator = {
        **row_metrics,
        "creator_id": creator_id_seed,
        "source": "pgy",
        "nickname": nickname,
        "pgy_url": "",
        "profile_url": "" if "/solar/pre-trade/note/kol" in page_url else page_url,
        "followers_count": followers,
        "quote_price": quote_price,
        "ip_city": location,
        "persona_tags": "/".join(tags),
        "raw_payload": raw_payload,
    }
    if no_order_permission:
        creator["order_permission_status"] = no_order_permission
    for key, value in table_payload.items():
        if key == "raw_table":
            continue
        if key in {"quote_price", "video_quote_price"} and not _looks_like_quote_text(value):
            continue
        if value not in (None, "") and creator.get(key) in (None, ""):
            creator[key] = value
    if table_payload.get("followers_count") not in (None, "") and _looks_like_count_text(table_payload["followers_count"]):
        creator["followers_count"] = _number_from_text(str(table_payload["followers_count"]))
    if creator.get("quote_price") in (None, "") and table_payload.get("quote_price") not in (None, "") and _looks_like_quote_text(table_payload["quote_price"]):
        creator["quote_price"] = _number_from_text(str(table_payload["quote_price"]))
    if creator.get("video_quote_price") in (None, "") and table_payload.get("video_quote_price") not in (None, "") and _looks_like_quote_text(table_payload["video_quote_price"]):
        creator["video_quote_price"] = _number_from_text(str(table_payload["video_quote_price"]))
    return creator


def _extract_table_headers(page: Any) -> list[str]:
    selectors = [
        ".blogger-list_list .d-new-table thead th",
        ".blogger-list_list table thead th",
    ]
    for selector in selectors:
        try:
            headers = [
                _clean_text(item)
                for item in page.locator(selector).evaluate_all(
                    "(nodes) => nodes.map((node) => node.innerText || node.textContent || '')"
                )
            ]
        except Exception:
            headers = []
        if any(headers):
            return headers
    return []


def _extract_row_table_payload(row: Any, headers: list[str]) -> dict[str, Any]:
    if not headers:
        return {}
    try:
        values = [
            _clean_text(item)
            for item in row.locator("td").evaluate_all(
                "(nodes) => nodes.map((node) => node.innerText || node.textContent || '')"
            )
        ]
    except Exception:
        values = []
    if not values:
        return {}
    pairs = dict(zip(headers, values))
    payload: dict[str, Any] = {"raw_table": pairs}
    for header, value in pairs.items():
        if not header or not value:
            continue
        payload[header] = value
        normalized_header = _normalize_export_header(header)
        field = PGY_TABLE_TO_PAYLOAD_FIELDS.get(header) or PGY_TABLE_TO_PAYLOAD_FIELDS.get(normalized_header)
        if not field:
            field = next(
                (
                    target
                    for target, aliases in PGY_EXPORT_FIELD_ALIASES.items()
                    if normalized_header in {_normalize_export_header(alias) for alias in aliases}
                ),
                "",
            )
        if field and (field not in {"quote_price", "video_quote_price"} or _looks_like_quote_text(value)):
            payload[field] = value
    return payload


def _extract_detail_fields(text: str, url: str) -> dict[str, Any]:
    lines = _visible_text_lines(text)

    def after(label: str) -> str:
        return _label_value(lines, label)

    nickname = ""
    if "笔记主页" in lines:
        index = lines.index("笔记主页")
        nickname = lines[index + 2] if index + 2 < len(lines) and lines[index + 1] == "直播主页" else ""
        if nickname in {"", "小红书号：", "小红书号", "数据概览", "笔记数据", "粉丝分析"}:
            nickname = ""
    xhs_id = after("小红书号：")
    pgy_blogger_id = ""
    blogger_match = re.search(r"/blogger-detail/([^?/#]+)", url)
    if blogger_match:
        pgy_blogger_id = blogger_match.group(1)
    fans_35_plus_ratio = None
    age_line = next((line for line in lines if "35-44" in line and "占比" in line), "")
    match = re.search(r"占比\s*([0-9.]+)%", age_line)
    if match:
        fans_35_plus_ratio = float(match.group(1)) / 100
    if fans_35_plus_ratio is None:
        fans_35_plus_ratio = _ratio_after_label(lines, "35-44")
    topic_point = ""
    interest_line = next((line for line in lines if "用户最感兴趣的内容类型为" in line), "")
    if interest_line:
        topic_point = interest_line.replace("用户最感兴趣的内容类型为", "").strip()
    personal_intro = ""
    for label in ("个人简介：", "个人简介", "简介：", "简介"):
        personal_intro = after(label)
        if personal_intro:
            break
    quote_price = _number_from_text(after("图文笔记一口价"))
    video_quote = _number_from_text(after("视频笔记一口价"))
    followers = _number_from_text(after("粉丝数"))
    liked_collected = _number_from_text(after("获赞与收藏"))
    persona_tags = ""
    ip_city = ""
    mcn_status = ""
    home_city = ""
    primary_categories: list[str] = []
    if nickname and nickname in lines:
        name_index = lines.index(nickname)
        persona_tags = lines[name_index + 3] if name_index + 3 < len(lines) and lines[name_index + 1] == "小红书号：" else ""
        ip_city = lines[name_index + 4] if name_index + 4 < len(lines) and lines[name_index + 1] == "小红书号：" else ""
        mcn_status = lines[name_index + 5] if name_index + 5 < len(lines) and lines[name_index + 1] == "小红书号：" else ""
        home_city = lines[name_index + 6] if name_index + 6 < len(lines) and lines[name_index + 1] == "小红书号：" else ""
        categories_start = name_index + 7 if name_index + 1 < len(lines) and lines[name_index + 1] == "小红书号：" else name_index + 1
        for line in lines[categories_start:]:
            if line in PGY_DETAIL_CATEGORY_STOP_LINES or line.startswith("与") or "相似的博主" in line:
                break
            category = sanitize_creator_type(line)
            if category:
                primary_categories.append(category)
    note_case_items = _extract_note_cases(lines)
    note_performance = {
        "exposure_median": _metric_value_after(lines, "曝光中位数"),
        "read_median": _metric_value_after(lines, "阅读中位数"),
        "interaction_median": _metric_value_after(lines, "互动中位数"),
        "median_like_count": _metric_value_after(lines, "中位点赞量"),
        "median_save_count": _metric_value_after(lines, "中位收藏量"),
        "median_comment_count": _metric_value_after(lines, "中位评论量"),
        "median_share_count": _metric_value_after(lines, "中位分享量"),
        "median_follow_count": _metric_value_after(lines, "中位关注量"),
        "interaction_rate": _ratio_after_label(lines, "互动率"),
        "video_completion_rate": _ratio_after_label(lines, "视频完播率"),
        "thousand_like_note_ratio": _ratio_after_label(lines, "千赞笔记比例"),
        "hundred_like_note_ratio": _ratio_after_label(lines, "百赞笔记比例"),
    }
    note_performance = {key: value for key, value in note_performance.items() if value is not None}
    fan_analysis = {
        "fan_growth": _metric_value_after(lines, "粉丝增量"),
        "fan_growth_ratio": _ratio_after_label(lines, "粉丝量变化幅度"),
        "active_fans_ratio": _ratio_after_label(lines, "活跃粉丝占比"),
        "read_fans_ratio": _ratio_after_label(lines, "阅读粉丝占比"),
        "interaction_fans_ratio": _ratio_after_label(lines, "互动粉丝占比"),
        "order_fans_ratio": _ratio_after_label(lines, "下单粉丝占比"),
        "female_fans_ratio": _gender_ratio(lines, "女性"),
        "male_fans_ratio": _gender_ratio(lines, "男性"),
        "fans_18_24_ratio": _ratio_near_label(lines, "18-24"),
        "fans_25_34_ratio": _ratio_near_label(lines, "25-34"),
        "fans_35_44_ratio": _ratio_near_label(lines, "35-44"),
        "fans_44_plus_ratio": _ratio_near_label(lines, ">44") or _ratio_near_label(lines, "44岁以上"),
        "gender_distribution": _line_between(lines, "性别分布", "年龄分布"),
        "age_distribution": _line_between(lines, "年龄分布", "地域分布"),
        "region_distribution": _line_between(lines, "地域分布", "按省份"),
        "device_distribution": _line_between(lines, "用户设备分布", "用户兴趣"),
    }
    fan_analysis = {key: value for key, value in fan_analysis.items() if value not in ("", None)}
    region_distribution = _parse_region_distribution_line(fan_analysis.get("region_distribution"))
    device_distribution = _parse_device_distribution_line(fan_analysis.get("device_distribution"))
    services = {
        "active_days_7d": _metric_value_after(lines, "近7天活跃天数"),
        "reply_rate_48h": _ratio_after_label(lines, "邀约48小时回复率"),
    }
    services = {key: value for key, value in services.items() if value is not None}
    overview_match = re.search(r"博主优势(.*?)笔记数据", text, flags=re.S)
    blogger_advantage = ""
    if overview_match:
        blogger_advantage = re.sub(r"\s+", "", overview_match.group(1))
    raw_detail = {
        "detail_url": url,
        "detail_text": text[:12000],
        "lines": lines,
        "data_updated_to": after("数据更新至："),
        "personal_intro": personal_intro,
        "blogger_advantage": blogger_advantage,
        "video_quote_price": video_quote,
        "note_performance": note_performance,
        "fan_analysis": fan_analysis,
        "service_performance": services,
        "note_cases": note_case_items,
        "fans_age_line": age_line,
        "interest_line": interest_line,
        "active_fans_ratio": after("活跃粉丝占比"),
        "read_fans_ratio": after("阅读粉丝占比"),
        "interaction_fans_ratio": after("互动粉丝占比"),
    }
    no_order_permission = _pgy_no_order_permission_hint(text)
    if no_order_permission:
        raw_detail["order_permission_status"] = no_order_permission
    result = {
        "pgy_url": url,
        "profile_url": f"https://www.xiaohongshu.com/user/profile/{pgy_blogger_id}" if pgy_blogger_id else "",
        "xiaohongshu_id": xhs_id,
        "pgy_blogger_id": pgy_blogger_id,
        "raw_payload": raw_detail,
    }
    if no_order_permission:
        result["order_permission_status"] = no_order_permission
    if nickname:
        result["nickname"] = nickname
    if followers is not None:
        result["followers_count"] = followers
    if liked_collected is not None:
        result["liked_collected_count"] = liked_collected
    if quote_price is not None:
        result["quote_price"] = quote_price
    if video_quote is not None:
        result["video_quote_price"] = video_quote
    if persona_tags:
        result["persona_tags"] = persona_tags
    if ip_city:
        result["ip_city"] = ip_city
    if primary_categories:
        result["creator_type"] = sanitize_creator_type("/".join(primary_categories))
    if topic_point:
        result["topic_point"] = topic_point
    if personal_intro:
        result["personal_intro"] = personal_intro
    if region_distribution:
        result["audience_region_distribution"] = region_distribution
    if device_distribution:
        result["audience_device_distribution"] = device_distribution
    metric_mapping = {
        "daily_exposure_median": note_performance.get("exposure_median"),
        "daily_read_median": note_performance.get("read_median"),
        "daily_interaction_median": note_performance.get("interaction_median"),
        "daily_thousand_like_note_ratio": note_performance.get("thousand_like_note_ratio"),
        "daily_hundred_like_note_ratio": note_performance.get("hundred_like_note_ratio"),
        "video_completion_rate": note_performance.get("video_completion_rate"),
        "active_fans_ratio": fan_analysis.get("active_fans_ratio"),
        "female_fans_ratio": fan_analysis.get("female_fans_ratio"),
        "male_fans_ratio": fan_analysis.get("male_fans_ratio"),
        "fans_18_24_ratio": fan_analysis.get("fans_18_24_ratio"),
        "fans_25_34_ratio": fan_analysis.get("fans_25_34_ratio"),
        "fans_35_44_ratio": fan_analysis.get("fans_35_44_ratio"),
        "fans_44_plus_ratio": fan_analysis.get("fans_44_plus_ratio"),
        "fans_growth_ratio": fan_analysis.get("fan_growth_ratio"),
        "read_fans_ratio": fan_analysis.get("read_fans_ratio"),
        "interaction_fans_ratio": fan_analysis.get("interaction_fans_ratio"),
        "order_fans_ratio": fan_analysis.get("order_fans_ratio"),
        "reply_rate_48h": services.get("reply_rate_48h"),
        "active_days_7d": services.get("active_days_7d"),
    }
    derived_fans_35_plus = None
    if fan_analysis.get("fans_35_44_ratio") is not None or fan_analysis.get("fans_44_plus_ratio") is not None:
        derived_fans_35_plus = (fan_analysis.get("fans_35_44_ratio") or 0) + (fan_analysis.get("fans_44_plus_ratio") or 0)
    if derived_fans_35_plus is not None:
        metric_mapping["fans_35_plus_ratio"] = min(derived_fans_35_plus, 1)
    elif fans_35_plus_ratio is not None:
        metric_mapping["fans_35_plus_ratio"] = fans_35_plus_ratio
    result.update({key: value for key, value in metric_mapping.items() if value is not None})
    return result


def _extract_row_link_fields(row: Any) -> dict[str, Any]:
    try:
        urls = row.evaluate(
            """
            node => Array.from(node.querySelectorAll('*'))
              .flatMap(el => {
                const attrs = Array.from(el.attributes || []).map(attr => attr.value);
                return [el.href, el.getAttribute('href'), el.getAttribute('data-href'), el.getAttribute('data-url'), ...attrs];
              })
              .filter(Boolean)
            """
        )
    except Exception:
        urls = []
    pgy_url = next((url for url in urls if "pgy.xiaohongshu.com" in url and "/blogger-detail/" in url), "")
    if not pgy_url:
        pgy_url = next((url for url in urls if "/blogger-detail/" in url), "")
    if not pgy_url:
        return {}
    return _detail_url_fields(pgy_url, source="row_dom") or {"pgy_url": pgy_url}


def _detail_url_fields(url: str, source: str = "") -> dict[str, Any]:
    if not url or "/blogger-detail/" not in url:
        return {}
    blogger_match = re.search(r"/blogger-detail/([^?/#]+)", url)
    if not blogger_match:
        return {}
    blogger_id = blogger_match.group(1)
    canonical_url = f"https://pgy.xiaohongshu.com/solar/pre-trade/blogger-detail/{blogger_id}"
    return {
        "pgy_url": canonical_url,
        "profile_url": f"https://www.xiaohongshu.com/user/profile/{blogger_id}",
        "pgy_blogger_id": blogger_id,
        "pgy_url_source": source,
    }


def _kol_api_key(kol: dict[str, Any]) -> str:
    for field in ["userId", "user_id", "bloggerId", "blogger_id", "kolId", "kol_id", "redId"]:
        value = str(kol.get(field) or "").strip()
        if value:
            return f"{field}:{value}"
    return "|".join(
        str(kol.get(field) or "").strip()
        for field in ["name", "nickName", "nickname", "location", "city", "fansCount", "picturePrice"]
    )


def _remember_kol_api_kols(page: Any, kols: list[dict[str, Any]]) -> None:
    if not kols:
        return
    pool = getattr(page, "_pgy_api_kol_pool", []) or []
    keys = getattr(page, "_pgy_api_kol_keys", set()) or set()
    if not isinstance(pool, list):
        pool = []
    if not isinstance(keys, set):
        keys = set(keys) if isinstance(keys, (list, tuple)) else set()
    for kol in kols:
        if not isinstance(kol, dict):
            continue
        key = _kol_api_key(kol)
        if not key or key in keys:
            continue
        keys.add(key)
        pool.append(kol)
    setattr(page, "_pgy_api_kol_pool", pool)
    setattr(page, "_pgy_api_kol_keys", keys)
    setattr(page, "_pgy_latest_api_kols", [kol for kol in kols if isinstance(kol, dict)])


def _kol_nested_value(payload: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        value = payload.get(key)
        if value not in (None, ""):
            return value
    for value in payload.values():
        if isinstance(value, dict):
            nested = _kol_nested_value(value, *keys)
            if nested not in (None, ""):
                return nested
    return None


def _creator_from_api_kol(kol: dict[str, Any], page_number: int = 0, row_index: int = 0) -> dict[str, Any]:
    user_id = str(_kol_nested_value(kol, "userId", "user_id", "bloggerId", "blogger_id", "kolId", "kol_id") or "").strip()
    nickname = str(_kol_nested_value(kol, "name", "nickName", "nickname") or "").strip()
    creator_id = f"pgy-api:{user_id}" if user_id else f"pgy-api:{_kol_api_key(kol)}"
    raw_payload = {
        "list_api_kol": kol,
        "collection_source": "list_api",
        "collection_page": page_number,
        "collection_row_index": row_index,
    }
    no_order_permission = _pgy_no_order_permission_hint(kol)
    if no_order_permission:
        raw_payload["order_permission_status"] = no_order_permission
    recent_note_briefs = _recent_note_briefs_from_kol(kol, max_notes=2)
    if recent_note_briefs:
        raw_payload["recent_note_briefs"] = recent_note_briefs
    creator: dict[str, Any] = {
        "creator_id": creator_id,
        "source": "pgy",
        "nickname": nickname,
        "ip_city": str(_kol_nested_value(kol, "location", "city") or "").strip(),
        "raw_payload": raw_payload,
    }
    if no_order_permission:
        creator["order_permission_status"] = no_order_permission
    if user_id:
        creator.update(_detail_url_fields(f"https://pgy.xiaohongshu.com/solar/pre-trade/blogger-detail/{user_id}", source="list_api"))
    red_id = _kol_nested_value(kol, "redId", "red_id", "xiaohongshuId", "xiaohongshu_id")
    if red_id not in (None, ""):
        creator["xiaohongshu_id"] = str(red_id).strip()
    avatar = _kol_nested_value(kol, "headPhoto", "avatar", "avatarUrl", "imageUrl")
    if avatar not in (None, ""):
        creator["avatar_url"] = str(avatar).strip()
    metric_fields = {
        "followers_count": ("fansNum", "fansCount", "fans_count", "followerCount", "followersCount"),
        "liked_collected_count": ("likeCollectCountInfo", "likedCollectedCount", "likeCollectCount"),
        "quote_price": ("picturePrice", "quotePrice", "imageQuotePrice", "picPrice"),
        "video_quote_price": ("videoPrice", "videoQuotePrice"),
        "daily_exposure_median": ("accumCommonImpMedinNum30d", "impMedian", "mAccumImpNum", "exposureMedian"),
        "daily_read_median": ("clickMidNum", "readMedian", "readMedianNum"),
        "daily_interaction_median": ("mEngagementNum", "mengagementNum", "interactionMedian"),
        "image_daily_exposure_median": ("accumPicCommonImpMedinNum30d",),
        "image_daily_read_median": ("pictureClickMidNum",),
        "image_daily_interaction_median": ("pictureInterMidNum",),
        "video_daily_exposure_median": ("accumVideoCommonImpMedinNum30d",),
        "video_daily_read_median": ("videoClickMidNum",),
        "video_daily_interaction_median": ("videoInterMidNum",),
        "cooperation_exposure_median": ("accumCoopImpMedinNum30d",),
        "cooperation_read_median": ("readMidCoop30",),
        "cooperation_interaction_median": ("interMidCoop30",),
        "overflow_store_median": ("mCpuvNum30d", "mcpuvNum30d"),
        "overflow_store_unit_price": ("estimateCpuv30d",),
        "image_cpm": ("estimatePictureCpm", "pictureCpm", "picCpm"),
        "video_cpm": ("estimateVideoCpm", "videoCpm"),
        "image_read_unit_price": ("pictureReadCost", "pictureReadUnitPrice", "imageReadUnitPrice"),
        "image_interaction_unit_price": ("estimatePictureEngageCost", "pictureInteractionUnitPrice", "imageInteractionUnitPrice"),
        "video_read_unit_price": ("videoReadCost", "videoReadCostV2", "videoReadUnitPrice"),
        "video_interaction_unit_price": ("estimateVideoEngageCost", "videoInteractionUnitPrice"),
    }
    for field, keys in metric_fields.items():
        value = _kol_nested_value(kol, *keys)
        if value not in (None, ""):
            creator[field] = _number_from_text(str(value))
    for ratio_field, keys in {
        "fans_25_34_ratio": ("fans25To34Rate", "fans_25_34_ratio"),
        "fans_35_44_ratio": ("fans35To44Rate", "fans_35_44_ratio"),
        "fans_44_plus_ratio": ("fans44PlusRate", "fans_44_plus_ratio"),
        "active_fans_ratio": ("fansActiveIn28dLv", "activeFansRate"),
        "fans_growth_ratio": ("fans30GrowthRate", "fansGrowthRate"),
        "read_fans_ratio": ("readFansRate",),
        "interaction_fans_ratio": ("fansEngageNum30dLv", "engageFansRate", "interactionFansRate"),
        "order_fans_ratio": ("payFansUserRate30d",),
        "reply_rate_48h": ("inviteReply48hNumRatio", "responseRate", "replyRate48h"),
    }.items():
        value = _kol_nested_value(kol, *keys)
        if value not in (None, ""):
            creator[ratio_field] = _ratio_from_percent_value(value)
    if creator.get("fans_35_44_ratio") is not None or creator.get("fans_44_plus_ratio") is not None:
        creator["fans_35_plus_ratio"] = min((creator.get("fans_35_44_ratio") or 0) + (creator.get("fans_44_plus_ratio") or 0), 1)
    content_tags = []
    for item in kol.get("contentTags") or []:
        if isinstance(item, dict):
            if item.get("taxonomy1Tag"):
                content_tags.append(str(item["taxonomy1Tag"]))
            content_tags.extend(str(tag) for tag in item.get("taxonomy2Tags") or [] if str(tag).strip())
        elif str(item).strip():
            content_tags.append(str(item).strip())
    if content_tags:
        creator["creator_type"] = sanitize_creator_type("/".join(dict.fromkeys(content_tags)))
    personal_tags = [str(item).strip() for item in kol.get("personalTags") or [] if str(item).strip()]
    if personal_tags:
        creator["persona_tags"] = "、".join(personal_tags)
    return creator


def _install_kol_response_capture(page: Any) -> None:
    if getattr(page, "_pgy_kol_response_capture_installed", False):
        return
    _reset_kol_response_capture_state(page)

    def handle_response(response: Any) -> None:
        if "/api/solar/cooperator/blogger/v2" not in getattr(response, "url", ""):
            return
        try:
            payload = response.json()
        except Exception:
            return
        data = payload.get("data") if isinstance(payload, dict) else {}
        kols = data.get("kols") if isinstance(data, dict) else []
        if not isinstance(kols, list) or not kols:
            return
        try:
            total = data.get("total") or data.get("totalCount") or data.get("count")
            if total not in (None, ""):
                setattr(page, "_pgy_latest_api_total", int(float(str(total).replace(",", "").replace("，", ""))))
        except Exception:
            pass
        _remember_kol_api_kols(page, kols)
        try:
            request = response.request
            setattr(page, "_pgy_latest_kol_request", {
                "url": getattr(response, "url", ""),
                "method": getattr(request, "method", "GET"),
                "post_data": getattr(request, "post_data", "") or "",
            })
        except Exception:
            pass

    try:
        page.on("response", handle_response)
        setattr(page, "_pgy_kol_response_capture_installed", True)
    except Exception:
        pass


def _reset_kol_response_capture_state(page: Any) -> None:
    setattr(page, "_pgy_latest_api_kols", [])
    setattr(page, "_pgy_api_kol_pool", [])
    setattr(page, "_pgy_api_kol_keys", set())
    setattr(page, "_pgy_latest_kol_request", {})
    setattr(page, "_pgy_latest_api_total", None)


def _kol_request_body(request: dict[str, Any] | None) -> dict[str, Any]:
    if not isinstance(request, dict):
        return {}
    post_data = str(request.get("post_data") or request.get("body") or "")
    if post_data:
        try:
            payload = json.loads(post_data)
            return payload if isinstance(payload, dict) else {}
        except Exception:
            return {}
    url = str(request.get("url") or "")
    try:
        return dict(parse_qsl(urlsplit(url).query, keep_blank_values=True))
    except Exception:
        return {}


def _kol_request_from_snapshot(snapshot: dict[str, Any] | None) -> dict[str, Any]:
    if not isinstance(snapshot, dict):
        return {}
    request = snapshot.get("request") if isinstance(snapshot.get("request"), dict) else snapshot
    url = str(request.get("url") or "").strip()
    if not url or "/api/solar/cooperator/blogger/v2" not in url:
        return {}
    return {
        "url": url,
        "method": str(request.get("method") or "GET").upper(),
        "post_data": str(request.get("post_data") or request.get("body") or ""),
    }


def _current_kol_request_snapshot(page: Any) -> dict[str, Any]:
    request = _kol_request_from_snapshot(getattr(page, "_pgy_latest_kol_request", {}) or {})
    latest = getattr(page, "_pgy_latest_api_kols", []) or []
    pool = getattr(page, "_pgy_api_kol_pool", []) or []
    if not request:
        return {}
    return {
        "request": request,
        "latest_count": len(latest) if isinstance(latest, list) else 0,
        "pool_count": len(pool) if isinstance(pool, list) else 0,
        "total_count": getattr(page, "_pgy_latest_api_total", None),
        "captured_at": _now_text(),
    }


def _request_payload_values(payload: dict[str, Any], key: str) -> list[Any]:
    value = payload.get(key)
    if value in (None, "", [], {}):
        return []
    if isinstance(value, list):
        return value
    return [value]


def _request_range_min(value: Any) -> float | None:
    if not isinstance(value, list) or not value:
        return None
    try:
        number = float(value[0])
    except (TypeError, ValueError):
        return None
    return number if number >= 0 else None


def _request_range_max(value: Any) -> float | None:
    if not isinstance(value, list) or len(value) < 2:
        return None
    try:
        number = float(value[1])
    except (TypeError, ValueError):
        return None
    return number if number >= 0 else None


def _numbers_equal(left: Any, right: Any) -> bool:
    try:
        return abs(float(left) - float(right)) < 0.0001
    except (TypeError, ValueError):
        return False


def expected_kol_request_constraints_from_plan(plan: dict[str, Any] | None) -> dict[str, Any]:
    constraints: dict[str, Any] = {}
    filters = plan.get("filters") if isinstance(plan, dict) else []
    for item in filters or []:
        if not isinstance(item, dict):
            continue
        field = str(item.get("field") or "")
        if field == "博主类目":
            sub_value = str(item.get("sub_value") or item.get("subValue") or "").strip()
            value = sub_value or str(item.get("value") or "").strip()
            if value:
                constraints.setdefault("contentTag", [])
                if value not in constraints["contentTag"]:
                    constraints["contentTag"].append(value)
        elif field in {"家庭身份", "职业身份", "特色背景"}:
            value = str(item.get("value") or "").strip()
            if value:
                constraints.setdefault("personalTags", [])
                if value not in constraints["personalTags"]:
                    constraints["personalTags"].append(value)
        elif field == "笔记类型":
            value = str(item.get("value") or "").strip()
            if "视频" in value:
                constraints["noteType"] = 2
            elif "图文" in value:
                constraints["noteType"] = 1
        elif field == "合作报价":
            _, max_value = _apply_range_policy(item, item.get("min"), item.get("max"))
            if max_value in (None, ""):
                _, max_value = _range_numbers_from_text(str(item.get("value") or ""))
            for sub_field in _subfield_names_from_item(item, default=["图文笔记", "视频笔记"]):
                if max_value in (None, ""):
                    continue
                if "视频" in sub_field:
                    constraints["videoPriceUpper"] = max_value
                elif "图文" in sub_field or "笔记" in sub_field:
                    constraints["notePriceUpper"] = max_value
        elif field == "阅读中位数":
            min_value, max_value = _apply_range_policy(item, item.get("min"), item.get("max"))
            if min_value in (None, ""):
                min_value, _ = _range_numbers_from_text(str(item.get("value") or ""))
            if min_value not in (None, ""):
                constraints["readMidNor30_min"] = min_value
        elif field == "互动中位数":
            min_value, max_value = _apply_range_policy(item, item.get("min"), item.get("max"))
            if min_value in (None, ""):
                min_value, _ = _range_numbers_from_text(str(item.get("value") or ""))
            if min_value not in (None, ""):
                constraints["interMidNor30_min"] = min_value
        elif field == "粉丝量":
            min_value, max_value = _apply_range_policy(item, item.get("min"), item.get("max"))
            if min_value in (None, ""):
                min_value, _ = _range_numbers_from_text(str(item.get("value") or ""))
            if min_value not in (None, ""):
                constraints["fansNumberLower"] = min_value
    return constraints


def validate_kol_request_snapshot_against_plan(
    snapshot: dict[str, Any] | None,
    plan: dict[str, Any] | None,
    *,
    require_clean_pool: bool = True,
) -> dict[str, Any]:
    expected = expected_kol_request_constraints_from_plan(plan)
    request = _kol_request_from_snapshot(snapshot)
    body = _kol_request_body(request)
    issues: list[str] = []
    warnings: list[str] = []
    latest_count = int((snapshot or {}).get("latest_count") or 0) if isinstance(snapshot, dict) else 0
    pool_count = int((snapshot or {}).get("pool_count") or 0) if isinstance(snapshot, dict) else 0
    if not request:
        issues.append("未捕获蒲公英达人列表 API 请求")
    if body.get("similarUserId"):
        issues.append("当前请求仍包含 similarUserId，相似达人搜索会污染采前筛选结果")
    if str(body.get("searchType") or "") == "0" and expected:
        warnings.append("当前请求 searchType=0，可能是泛推荐或空筛选请求")
    if require_clean_pool and latest_count and pool_count > latest_count:
        warnings.append(f"API 池包含 {pool_count} 条，最新响应 {latest_count} 条；正式采集前应只沿用当前请求分页扩展，避免旧响应混入")

    content_tags = [str(item) for item in _request_payload_values(body, "contentTag")]
    for value in expected.get("contentTag") or []:
        if value not in content_tags:
            issues.append(f"请求体缺少 contentTag={value}")
    personal_tags = [str(item) for item in _request_payload_values(body, "personalTags")]
    for value in expected.get("personalTags") or []:
        if value not in personal_tags:
            issues.append(f"请求体缺少 personalTags={value}")
    for key in ("notePriceUpper", "videoPriceUpper", "fansNumberLower"):
        if key in expected and not _numbers_equal(body.get(key), expected[key]):
            issues.append(f"请求体 {key}={body.get(key)!r} 与期望 {expected[key]!r} 不一致")
    if "noteType" in expected and not _numbers_equal(body.get("noteType"), expected["noteType"]):
        issues.append(f"请求体 noteType={body.get('noteType')!r} 与期望 {expected['noteType']!r} 不一致")
    if "readMidNor30_min" in expected:
        actual = _request_range_min(body.get("readMidNor30"))
        if actual is None or actual < float(expected["readMidNor30_min"]):
            issues.append(f"请求体 readMidNor30 下限 {actual!r} 未达到期望 {expected['readMidNor30_min']!r}")
    if "interMidNor30_min" in expected:
        actual = _request_range_min(body.get("interMidNor30"))
        if actual is None or actual < float(expected["interMidNor30_min"]):
            issues.append(f"请求体 interMidNor30 下限 {actual!r} 未达到期望 {expected['interMidNor30_min']!r}")
    return {
        "request_valid": not issues,
        "expected_constraints": expected,
        "actual_constraints": {
            "searchType": body.get("searchType"),
            "similarUserId": body.get("similarUserId"),
            "contentTag": body.get("contentTag"),
            "personalTags": body.get("personalTags"),
            "fansNumberLower": body.get("fansNumberLower"),
            "fansNumberUpper": body.get("fansNumberUpper"),
            "readMidNor30": body.get("readMidNor30"),
            "interMidNor30": body.get("interMidNor30"),
            "noteType": body.get("noteType"),
            "notePriceUpper": body.get("notePriceUpper"),
            "videoPriceUpper": body.get("videoPriceUpper"),
        },
        "issues": issues,
        "warnings": warnings,
    }


def _base_kol_request_for_repair(page: Any, snapshot: dict[str, Any] | None) -> dict[str, Any]:
    request = _kol_request_from_snapshot(snapshot)
    if request:
        return request
    request = _kol_request_from_snapshot(getattr(page, "_pgy_latest_kol_request", {}) or {})
    if request:
        return request
    template = _read_kol_api_template()
    request = _kol_request_from_snapshot(template)
    if request:
        return request
    return {
        "url": "https://pgy.xiaohongshu.com/api/solar/cooperator/blogger/v2",
        "method": "POST",
        "post_data": "{}",
    }


def _repair_kol_request_from_constraints(
    page: Any,
    snapshot: dict[str, Any] | None,
    plan: dict[str, Any] | None,
) -> dict[str, Any]:
    expected = expected_kol_request_constraints_from_plan(plan)
    if not expected:
        return {}
    base_request = _base_kol_request_for_repair(page, snapshot)
    body = _kol_request_body(base_request)
    if not body:
        body = {}
    for key in ("similarUserId", "similarWord", "filterList"):
        if key in body:
            if key == "filterList":
                body[key] = []
            else:
                body.pop(key, None)
    body["searchType"] = 1
    body["pageNum"] = 1
    body["pageSize"] = int(body.get("pageSize") or 20)
    if expected.get("contentTag"):
        body["contentTag"] = expected["contentTag"]
    if expected.get("personalTags"):
        current = [str(item) for item in body.get("personalTags") or [] if str(item)]
        for tag in expected["personalTags"]:
            if tag not in current:
                current.append(tag)
        body["personalTags"] = current
    if "noteType" in expected:
        body["noteType"] = expected["noteType"]
    if "fansNumberLower" in expected:
        body["fansNumberLower"] = expected["fansNumberLower"]
        body.setdefault("fansNumberUpper", None)
    if "readMidNor30_min" in expected:
        body["readMidNor30"] = [expected["readMidNor30_min"], -1]
    if "interMidNor30_min" in expected:
        body["interMidNor30"] = [expected["interMidNor30_min"], -1]
    if "notePriceUpper" in expected:
        body["notePriceLower"] = body.get("notePriceLower", -1)
        body["notePriceUpper"] = expected["notePriceUpper"]
    if "videoPriceUpper" in expected:
        body["videoPriceLower"] = body.get("videoPriceLower", -1)
        body["videoPriceUpper"] = expected["videoPriceUpper"]
    repaired = {
        "url": str(base_request.get("url") or "https://pgy.xiaohongshu.com/api/solar/cooperator/blogger/v2"),
        "method": "POST",
        "post_data": json.dumps(body, ensure_ascii=False),
    }
    return repaired


def _repair_kol_capture_for_plan(
    page: Any,
    snapshot: dict[str, Any] | None,
    plan: dict[str, Any] | None,
) -> dict[str, Any]:
    repaired_request = _repair_kol_request_from_constraints(page, snapshot, plan)
    if not repaired_request:
        return {"ok": False, "message": "没有可修复的期望请求约束"}
    _apply_kol_request_source(page, repaired_request, clear_pool=True)
    first_page = _fetch_kol_api_page(page, repaired_request, 1, int(_kol_request_body(repaired_request).get("pageSize") or 20))
    repaired_snapshot = _current_kol_request_snapshot(page)
    validation = validate_kol_request_snapshot_against_plan(repaired_snapshot, plan, require_clean_pool=False)
    if not first_page:
        return {
            "ok": False,
            "message": "已重组蒲公英 API 请求，但接口未返回达人",
            "kol_request_snapshot": repaired_snapshot,
            "kol_request_validation": validation,
        }
    if not validation.get("request_valid"):
        return {
            "ok": False,
            "message": "已重组蒲公英 API 请求，但修复后仍未通过校验：" + "；".join(validation.get("issues") or []),
            "kol_request_snapshot": repaired_snapshot,
            "kol_request_validation": validation,
        }
    return {
        "ok": True,
        "message": "已自动修复为直接蒲公英 API 采集请求",
        "kol_request_snapshot": repaired_snapshot,
        "kol_request_validation": validation,
        "repaired_request": repaired_request,
        "first_page_count": len(first_page),
    }


def _restore_kol_request_snapshot(page: Any, snapshot: dict[str, Any] | None, *, clear_pool: bool = True) -> bool:
    request = _kol_request_from_snapshot(snapshot)
    if not request:
        return False
    if clear_pool:
        setattr(page, "_pgy_latest_api_kols", [])
        setattr(page, "_pgy_api_kol_pool", [])
        setattr(page, "_pgy_api_kol_keys", set())
        setattr(page, "_pgy_latest_api_total", None)
    setattr(page, "_pgy_latest_kol_request", request)
    return True


def _read_kol_api_template() -> dict[str, Any]:
    try:
        if not PGY_KOL_API_TEMPLATE_PATH.exists():
            return {}
        with PGY_KOL_API_TEMPLATE_PATH.open("r", encoding="utf-8") as handle:
            payload = json.load(handle)
    except Exception:
        return {}
    request = _kol_request_from_snapshot(payload)
    if not request:
        return {}
    return {
        "request": request,
        "saved_at": payload.get("saved_at") or "",
        "source": payload.get("source") or "template",
    }


def _write_kol_api_template(snapshot: dict[str, Any] | None, *, source: str = "capture") -> None:
    request = _kol_request_from_snapshot(snapshot)
    if not request:
        return
    try:
        PGY_KOL_API_TEMPLATE_PATH.parent.mkdir(parents=True, exist_ok=True)
        with PGY_KOL_API_TEMPLATE_PATH.open("w", encoding="utf-8") as handle:
            json.dump(
                {
                    "request": request,
                    "source": source,
                    "saved_at": _now_text(),
                },
                handle,
                ensure_ascii=False,
                indent=2,
            )
    except Exception:
        pass


def _apply_kol_request_source(page: Any, request: dict[str, Any], *, clear_pool: bool = True) -> bool:
    return _restore_kol_request_snapshot(page, {"request": request}, clear_pool=clear_pool)


def _validate_current_kol_request(page: Any, *, page_size: int = 20) -> bool:
    request = _kol_request_from_snapshot(getattr(page, "_pgy_latest_kol_request", {}) or {})
    if not request:
        return False
    if _fetch_kol_api_page(page, request, 1, page_size):
        _write_kol_api_template({"request": request}, source="validated")
        return True
    return False


def _capture_kol_api_from_browser(page: Any, *, timeout_ms: int = 8000) -> bool:
    try:
        with page.expect_response(
            lambda response: "/api/solar/cooperator/blogger/v2" in getattr(response, "url", ""),
            timeout=timeout_ms,
        ) as response_info:
            try:
                refresh = page.locator("button").filter(has_text=re.compile("查询|搜索|刷新")).first
                if refresh.count():
                    refresh.click(timeout=1500)
                else:
                    page.evaluate("window.dispatchEvent(new Event('scroll'))")
            except Exception:
                page.evaluate("window.dispatchEvent(new Event('scroll'))")
        response = response_info.value
        try:
            request = response.request
            captured = {
                "url": getattr(response, "url", ""),
                "method": getattr(request, "method", "GET"),
                "post_data": getattr(request, "post_data", "") or "",
            }
        except Exception:
            captured = {}
        if not _apply_kol_request_source(page, captured, clear_pool=True):
            return False
        return _validate_current_kol_request(page)
    except Exception:
        return False


def _restore_kol_request_from_template(page: Any) -> bool:
    template = _read_kol_api_template()
    request = _kol_request_from_snapshot(template)
    if not request:
        return False
    if not _apply_kol_request_source(page, request, clear_pool=True):
        return False
    return _validate_current_kol_request(page)


def _emit_collection_progress(callback: Any, payload: dict[str, Any]) -> None:
    if not callable(callback):
        return
    try:
        callback(payload)
    except Exception:
        pass


def _stop_requested(callback: Any) -> bool:
    if not callable(callback):
        return False
    try:
        return bool(callback())
    except Exception:
        return False


def _prime_kol_api_capture(page: Any, reload_if_empty: bool = True, stop_callback: Any = None) -> bool:
    for _ in range(4):
        if _stop_requested(stop_callback):
            return False
        kols = getattr(page, "_pgy_latest_api_kols", []) or []
        if isinstance(kols, list) and kols:
            return True
        page.wait_for_timeout(250)
    if _stop_requested(stop_callback):
        return False
    if not reload_if_empty:
        return False
    try:
        try:
            with page.expect_response(
                lambda response: "/api/solar/cooperator/blogger/v2" in getattr(response, "url", ""),
                timeout=6000,
            ):
                page.reload(wait_until="domcontentloaded", timeout=12000)
        except Exception:
            page.reload(wait_until="domcontentloaded", timeout=12000)
        page.wait_for_timeout(1200)
    except Exception:
        return False
    for _ in range(4):
        if _stop_requested(stop_callback):
            return False
        kols = getattr(page, "_pgy_latest_api_kols", []) or []
        if isinstance(kols, list) and kols:
            return True
        page.wait_for_timeout(250)
    kols = getattr(page, "_pgy_latest_api_kols", []) or []
    return isinstance(kols, list) and bool(kols)


def _should_reload_for_api_prime(*, apply_filters: bool, preserve_existing_filters: bool, preflight_only: bool) -> bool:
    return bool(not apply_filters and not preserve_existing_filters and not preflight_only)


def _set_nested_pagination(payload: Any, page_number: int, page_size: int) -> bool:
    if not isinstance(payload, dict):
        return False
    changed = False
    page_keys = {"page", "pageNo", "pageNum", "pageNumber", "pageIndex", "current", "currentPage"}
    size_keys = {"pageSize", "page_size", "size", "limit"}
    for key, value in list(payload.items()):
        if key in page_keys:
            payload[key] = page_number
            changed = True
        elif key in size_keys:
            payload[key] = page_size
            changed = True
        elif isinstance(value, dict):
            changed = _set_nested_pagination(value, page_number, page_size) or changed
    return changed


def _paginated_kol_request(request_info: dict[str, Any], page_number: int, page_size: int) -> dict[str, Any]:
    method = str(request_info.get("method") or "GET").upper()
    url = str(request_info.get("url") or "")
    post_data = str(request_info.get("post_data") or "")
    body = post_data
    if method == "GET":
        parts = urlsplit(url)
        pairs = dict(parse_qsl(parts.query, keep_blank_values=True))
        had_page = False
        for key in ["page", "pageNo", "pageNum", "pageNumber", "pageIndex", "current", "currentPage"]:
            if key in pairs:
                pairs[key] = str(page_number)
                had_page = True
        if not had_page:
            pairs["pageNum"] = str(page_number)
        for key in ["pageSize", "page_size", "size", "limit"]:
            if key in pairs:
                pairs[key] = str(page_size)
        url = urlunsplit((parts.scheme, parts.netloc, parts.path, urlencode(pairs), parts.fragment))
    elif post_data:
        try:
            payload = json.loads(post_data)
            if isinstance(payload, dict) and not _set_nested_pagination(payload, page_number, page_size):
                payload["pageNum"] = page_number
                payload["pageSize"] = page_size
            body = json.dumps(payload, ensure_ascii=False)
        except Exception:
            body = post_data
    return {"url": url, "method": method, "body": body}


def _fetch_kol_api_page(page: Any, request_info: dict[str, Any], page_number: int, page_size: int) -> list[dict[str, Any]]:
    request_payload = _paginated_kol_request(request_info, page_number, page_size)
    try:
        payload = page.evaluate(
            """
            async ({ url, method, body }) => {
              const controller = new AbortController();
              const timer = setTimeout(() => controller.abort(), 8000);
              const init = {
                method,
                credentials: 'include',
                signal: controller.signal,
                headers: {
                  'accept': 'application/json, text/plain, */*',
                  'content-type': 'application/json;charset=UTF-8'
                }
              };
              if (method !== 'GET' && body) init.body = body;
              try {
                const response = await fetch(url, init);
                return await response.json();
              } finally {
                clearTimeout(timer);
              }
            }
            """,
            request_payload,
        )
    except Exception:
        return []
    data = payload.get("data") if isinstance(payload, dict) else {}
    kols = data.get("kols") if isinstance(data, dict) else []
    if not isinstance(kols, list):
        return []
    try:
        total = data.get("total") or data.get("totalCount") or data.get("count")
        if total not in (None, ""):
            setattr(page, "_pgy_latest_api_total", int(float(str(total).replace(",", "").replace("，", ""))))
    except Exception:
        pass
    valid_kols = [kol for kol in kols if isinstance(kol, dict)]
    _remember_kol_api_kols(page, valid_kols)
    return valid_kols


def _append_expanded_api_creators(page: Any, creators: list[dict[str, Any]], seen: set[str], limit: int, stop_callback: Any = None) -> int:
    api_added = _expand_kol_api_pool(page, limit, stop_callback=stop_callback)
    if not api_added:
        return 0
    return _append_api_creators(page, creators, seen, limit)


def _expand_kol_api_pool(page: Any, limit: int, progress_callback: Any = None, mode: str = "api", stop_callback: Any = None) -> int:
    request_info = getattr(page, "_pgy_latest_kol_request", {}) or {}
    if not isinstance(request_info, dict) or not request_info.get("url"):
        return 0
    latest = getattr(page, "_pgy_latest_api_kols", []) or []
    page_size = max(20, min(100, len(latest) or 20))
    before = len(getattr(page, "_pgy_api_kol_pool", []) or [])
    max_pages = max(2, min(80, int((max(limit, before) / page_size) + 6)))
    idle_rounds = 0
    for page_number in range(2, max_pages + 1):
        if _stop_requested(stop_callback):
            break
        current_count = len(getattr(page, "_pgy_api_kol_pool", []) or [])
        if current_count >= limit:
            break
        pool_before = len(getattr(page, "_pgy_api_kol_pool", []) or [])
        kols = _fetch_kol_api_page(page, request_info, page_number, page_size)
        pool_after = len(getattr(page, "_pgy_api_kol_pool", []) or [])
        if not kols or pool_after <= pool_before:
            idle_rounds += 1
        else:
            idle_rounds = 0
        if page_number == 2 or pool_after >= limit or page_number % 5 == 0 or idle_rounds:
            _emit_collection_progress(
                progress_callback,
                {
                    "mode": mode,
                    "page_number": page_number,
                    "collected_count": min(pool_after, limit),
                    "limit": limit,
                    "page_size": page_size,
                    "idle_rounds": idle_rounds,
                },
            )
        if idle_rounds >= 2:
            break
        page.wait_for_timeout(250)
    return max(0, len(getattr(page, "_pgy_api_kol_pool", []) or []) - before)


def _append_api_creators(
    page: Any,
    creators: list[dict[str, Any]],
    seen: set[str],
    limit: int,
) -> int:
    pool = getattr(page, "_pgy_api_kol_pool", []) or []
    if not isinstance(pool, list):
        return 0
    added = 0
    for index, kol in enumerate(pool, start=1):
        if len(creators) >= limit:
            break
        if not isinstance(kol, dict):
            continue
        creator = _creator_from_api_kol(kol, page_number=((index - 1) // 20) + 1, row_index=index)
        if not creator.get("nickname") and not creator.get("pgy_url") and not creator.get("xiaohongshu_id"):
            continue
        key = _creator_collection_key(creator)
        if key in seen:
            continue
        seen.add(key)
        creators.append(creator)
        added += 1
    return added


def _collect_creators_from_api_pool(
    page: Any,
    limit: int,
    *,
    progress_callback: Any = None,
    mode: str = "api",
    stop_callback: Any = None,
) -> list[dict[str, Any]]:
    if _stop_requested(stop_callback):
        return []
    request_info = getattr(page, "_pgy_latest_kol_request", {}) or {}
    if not isinstance(request_info, dict) or not request_info.get("url"):
        return []
    creators: list[dict[str, Any]] = []
    seen: set[str] = set()
    before = len(getattr(page, "_pgy_api_kol_pool", []) or [])
    if before == 0:
        first_page = _fetch_kol_api_page(page, request_info, 1, 20)
        before = len(getattr(page, "_pgy_api_kol_pool", []) or [])
        if first_page:
            _emit_collection_progress(
                progress_callback,
                {"mode": mode, "page_number": 1, "collected_count": min(before, limit), "limit": limit, "page_size": len(first_page)},
            )
    if _stop_requested(stop_callback):
        return creators
    _append_api_creators(page, creators, seen, limit)
    _expand_kol_api_pool(page, limit, progress_callback=progress_callback, mode=mode, stop_callback=stop_callback)
    _append_api_creators(page, creators, seen, limit)
    if creators:
        _write_kol_api_template({"request": request_info}, source=mode)
        _emit_collection_progress(
            progress_callback,
            {"mode": f"{mode}_done", "collected_count": len(creators), "limit": limit},
        )
    return creators


def _match_api_kol(page: Any, row_index: int, creator: dict[str, Any]) -> dict[str, Any]:
    kols = getattr(page, "_pgy_latest_api_kols", []) or []
    if not isinstance(kols, list):
        return {}
    nickname = _clean_text(str(creator.get("nickname") or ""))
    location = _clean_text(str(creator.get("ip_city") or ""))
    candidates: list[dict[str, Any]] = []
    if 0 <= row_index < len(kols) and isinstance(kols[row_index], dict):
        candidates.append(kols[row_index])
    candidates.extend(kol for kol in kols if isinstance(kol, dict))
    seen_ids: set[int] = set()
    for kol in candidates:
        marker = id(kol)
        if marker in seen_ids:
            continue
        seen_ids.add(marker)
        kol_name = _clean_text(str(_kol_nested_value(kol, "name", "nickName", "nickname") or ""))
        kol_location = _clean_text(str(_kol_nested_value(kol, "location", "city") or ""))
        if nickname and kol_name and kol_name != nickname:
            continue
        if location and kol_location and kol_location != location and row_index >= len(kols):
            continue
        return kol
    return {}


def _api_kol_link_fields(page: Any, row_index: int, creator: dict[str, Any]) -> dict[str, Any]:
    kol = _match_api_kol(page, row_index, creator)
    if not kol:
        return {}
    user_id = str(_kol_nested_value(kol, "userId", "user_id", "bloggerId", "blogger_id", "kolId", "kol_id") or "").strip()
    if not user_id:
        return {}
    fields = _detail_url_fields(
        f"https://pgy.xiaohongshu.com/solar/pre-trade/blogger-detail/{user_id}",
        source="list_api",
    )
    if not fields:
        return {}
    red_id = _kol_nested_value(kol, "redId", "red_id", "xiaohongshuId", "xiaohongshu_id")
    if red_id not in (None, ""):
        fields["xiaohongshu_id"] = str(red_id).strip()
    avatar = _kol_nested_value(kol, "headPhoto", "avatar", "avatarUrl", "imageUrl")
    if avatar not in (None, ""):
        fields["avatar_url"] = str(avatar).strip()
    return fields


def _api_pool_matches_current_rows(page: Any) -> bool:
    kols = getattr(page, "_pgy_latest_api_kols", []) or []
    if not isinstance(kols, list) or not kols:
        return False
    rows = _creator_list_rows(page)
    try:
        count = min(rows.count(), len(kols), 5)
    except Exception:
        return False
    if count <= 0:
        return False
    checked = 0
    matched = 0
    for index in range(count):
        try:
            creator = _parse_row_text(rows.nth(index).inner_text(timeout=1200), page.url)
        except Exception:
            continue
        nickname = _clean_text(str((creator or {}).get("nickname") or ""))
        kol_name = _clean_text(str(_kol_nested_value(kols[index], "name", "nickName", "nickname") or ""))
        if not nickname or not kol_name:
            continue
        checked += 1
        if nickname == kol_name:
            matched += 1
    return bool(checked and matched == checked)


def _recent_note_type_label(value: Any) -> str:
    try:
        note_type = int(value)
    except (TypeError, ValueError):
        return ""
    if note_type == 2:
        return "视频笔记"
    if note_type == 1:
        return "图文笔记"
    return ""


def _normalize_media_url(value: Any) -> str:
    url = str(value or "").strip()
    if not url:
        return ""
    if url.startswith("//"):
        return f"https:{url}"
    if url.startswith("http://"):
        return f"https://{url[7:]}"
    return url


def _recent_note_briefs_from_kol(kol: dict[str, Any], max_notes: int = 2) -> list[dict[str, Any]]:
    note_list = kol.get("noteList")
    if not isinstance(note_list, list):
        return []
    briefs: list[dict[str, Any]] = []
    seen: set[str] = set()
    for index, item in enumerate(note_list):
        if not isinstance(item, dict):
            continue
        note_id = str(item.get("noteId") or item.get("id") or "").strip()
        if not note_id or note_id in seen:
            continue
        seen.add(note_id)
        briefs.append(
            {
                "note_id": note_id,
                "cover_url": _normalize_media_url(item.get("imageUrl") or item.get("cover") or item.get("coverUrl")),
                "note_type": _recent_note_type_label(item.get("noteType")) or str(item.get("noteType") or ""),
                "content_category": str(item.get("contentTag") or "").strip(),
                "feature_tags": [str(tag).strip() for tag in (item.get("featureTags") or []) if str(tag).strip()],
                "industry_tags": [str(tag).strip() for tag in (item.get("industryTags") or []) if str(tag).strip()],
                "bind": bool(item.get("bind")),
                "source": "list_api",
                "index": index,
            }
        )
        if len(briefs) >= max_notes:
            break
    return briefs


def _detail_note_cover(data: dict[str, Any]) -> str:
    images = data.get("imagesList")
    if isinstance(images, list):
        for item in images:
            if isinstance(item, str) and item.strip():
                return _normalize_media_url(item)
            if isinstance(item, dict):
                for key in ("url", "imageUrl", "originUrl", "src", "traceId", "fileId"):
                    value = _normalize_media_url(item.get(key))
                    if value:
                        return value
    video = data.get("videoInfo")
    if isinstance(video, dict):
        for key in ("coverUrl", "imageUrl", "poster", "firstFrameUrl"):
            value = _normalize_media_url(video.get(key))
            if value:
                return value
    return ""


def _recent_note_comments_from_payload(payload: Any, limit: int = 6) -> list[str]:
    if not isinstance(payload, list):
        return []
    comments: list[str] = []
    for item in payload:
        if not isinstance(item, dict):
            continue
        comment = item.get("comment")
        if isinstance(comment, dict):
            text = str(comment.get("content") or "").strip()
            if text:
                comments.append(text)
        for reply in item.get("l1L2Comments") or []:
            if not isinstance(reply, dict):
                continue
            text = str(reply.get("content") or "").strip()
            if text:
                comments.append(text)
        if len(comments) >= limit:
            break
    return comments[:limit]


def _fetch_note_detail_from_api(page: Any, note_id: str, comment_limit: int = 5) -> dict[str, Any]:
    note_id = str(note_id or "").strip()
    if not note_id:
        return {}
    try:
        payload = page.evaluate(
            """
            async ({ noteId, commentLimit }) => {
              const detailUrl = `https://pgy.xiaohongshu.com/api/solar/note/${noteId}/detail?bizCode=`;
              const commentUrl = `https://pgy.xiaohongshu.com/api/solar/note/${noteId}/comments?pageSize=${commentLimit}&pageIndex=0`;
              const [detailRes, commentsRes] = await Promise.all([
                fetch(detailUrl, { credentials: 'include' }),
                fetch(commentUrl, { credentials: 'include' }),
              ]);
              const detail = await detailRes.json();
              const comments = await commentsRes.json();
              return { detail, comments };
            }
            """,
            {"noteId": note_id, "commentLimit": max(1, min(comment_limit, 20))},
        )
    except Exception:
        return {}
    detail_data = payload.get("detail", {}).get("data") if isinstance(payload, dict) else {}
    comments_data = payload.get("comments", {}).get("data") if isinstance(payload, dict) else {}
    if not isinstance(detail_data, dict):
        return {}
    component_click_data = detail_data.get("compClickData")
    return {
        "note_id": note_id,
        "title": str(detail_data.get("title") or "").strip(),
        "content": str(detail_data.get("content") or "").strip(),
        "note_url": str(detail_data.get("noteLink") or "").strip(),
        "cover_url": _detail_note_cover(detail_data),
        "published_at": str(detail_data.get("createTime") or detail_data.get("time") or "").strip(),
        "read_count": detail_data.get("readNum"),
        "like_count": detail_data.get("likeNum"),
        "save_count": detail_data.get("favNum"),
        "comment_count": detail_data.get("cmtNum"),
        "share_count": detail_data.get("shareNum"),
        "follow_count": detail_data.get("followCnt"),
        "comments": _recent_note_comments_from_payload(comments_data),
        "component_click_data": component_click_data if isinstance(component_click_data, dict) else {},
        "source": "note_detail_api",
    }


def _collect_recent_note_details_from_api(page: Any, kol: dict[str, Any], max_notes: int = 2) -> list[dict[str, Any]]:
    briefs = _recent_note_briefs_from_kol(kol, max_notes=max_notes)
    detailed: list[dict[str, Any]] = []
    for brief in briefs:
        detail = _fetch_note_detail_from_api(page, brief.get("note_id") or "")
        note = {**brief, **detail}
        if brief.get("cover_url") and not note.get("cover_url"):
            note["cover_url"] = brief["cover_url"]
        if brief.get("content_category") and not note.get("brand"):
            note["brand"] = brief["content_category"]
        detailed.append(note)
    return [item for item in detailed if item.get("note_id")]


def _normalize_detail_note_item(item: dict[str, Any], index: int = 0) -> dict[str, Any]:
    note_id = str(item.get("noteId") or item.get("note_id") or item.get("id") or "").strip()
    note_type = "视频笔记" if item.get("isVideo") else (_recent_note_type_label(item.get("noteType")) or str(item.get("noteType") or "图文笔记"))
    note = {
        "note_id": note_id,
        "title": str(item.get("title") or item.get("noteTitle") or "").strip(),
        "cover_url": _normalize_media_url(item.get("imgUrl") or item.get("imageUrl") or item.get("coverUrl") or item.get("cover_url")),
        "note_type": note_type,
        "brand": str(item.get("brandName") or item.get("contentTag") or "").strip(),
        "content_category": str(item.get("contentTag") or item.get("brandName") or "").strip(),
        "published_at": str(item.get("date") or item.get("createTime") or item.get("time") or "").strip(),
        "read_count": item.get("readNum") if item.get("readNum") is not None else item.get("read_count"),
        "like_count": item.get("likeNum") if item.get("likeNum") is not None else item.get("like_count"),
        "save_count": item.get("collectNum") if item.get("collectNum") is not None else item.get("save_count"),
        "comment_count": item.get("cmtNum") if item.get("cmtNum") is not None else item.get("comment_count"),
        "share_count": item.get("shareNum") if item.get("shareNum") is not None else item.get("share_count"),
        "third_read_user_num": item.get("thirdReadUserNum"),
        "has_promoted_traffic": bool(item.get("isAdvertise")),
        "source": "detail_notes_api",
        "index": index,
    }
    return {key: value for key, value in note.items() if value not in ("", None)}


def _install_detail_notes_capture(page: Any) -> None:
    if getattr(page, "_pgy_detail_notes_capture_installed", False):
        return
    setattr(page, "_pgy_latest_detail_notes", [])

    def handle_response(response: Any) -> None:
        if "/api/solar/kol/data_v2/notes_detail" not in getattr(response, "url", ""):
            return
        try:
            payload = response.json()
        except Exception:
            return
        data = payload.get("data") if isinstance(payload, dict) else {}
        items = data.get("list") if isinstance(data, dict) else []
        if not isinstance(items, list) or not items:
            return
        notes = [_normalize_detail_note_item(item, index) for index, item in enumerate(items) if isinstance(item, dict)]
        notes = [item for item in notes if item.get("note_id") or item.get("title")]
        if notes:
            setattr(page, "_pgy_latest_detail_notes", notes)

    try:
        page.on("response", handle_response)
        setattr(page, "_pgy_detail_notes_capture_installed", True)
    except Exception:
        pass


def _install_detail_api_capture(page: Any) -> None:
    if getattr(page, "_pgy_detail_api_capture_installed", False):
        return
    setattr(page, "_pgy_detail_api_cache", {})

    def handle_response(response: Any) -> None:
        url = getattr(response, "url", "")
        if "pgy.xiaohongshu.com/api/" not in url:
            return
        interesting = [
            "/api/solar/cooperator/user/blogger/",
            "/api/solar/kol/data_v3/fans_summary",
            "/api/solar/kol/data/",
            "/api/pgy/kol/data/data_summary",
            "/api/solar/kol/data_v3/notes_rate",
            "/api/solar/kol/data_v2/notes_detail",
        ]
        if not any(item in url for item in interesting):
            return
        try:
            payload = response.json()
        except Exception:
            return
        data = payload.get("data") if isinstance(payload, dict) else None
        cache = getattr(page, "_pgy_detail_api_cache", {}) or {}
        if "/api/solar/cooperator/user/blogger/" in url and isinstance(data, dict):
            cache["blogger_profile"] = data
        elif "/api/solar/kol/data_v3/fans_summary" in url and isinstance(data, dict):
            cache["fans_summary"] = data
        elif "/fans_profile" in url and isinstance(data, dict):
            cache["fans_profile"] = data
        elif "/api/pgy/kol/data/data_summary" in url and isinstance(data, dict):
            business = "cooperation" if "business=1" in url else "daily"
            summary_cache = cache.setdefault("data_summary", {})
            summary_cache[business] = data
            request_cache = cache.setdefault("requests", {})
            request_cache.setdefault("data_summary", {})[business] = _response_request_info(response)
        elif "/api/solar/kol/data_v3/notes_rate" in url and isinstance(data, dict):
            business = "cooperation" if "business=1" in url else "daily"
            rate_cache = cache.setdefault("notes_rate", {})
            rate_cache[business] = data
            request_cache = cache.setdefault("requests", {})
            request_cache.setdefault("notes_rate", {})[business] = _response_request_info(response)
        elif "/api/solar/kol/data_v2/notes_detail" in url and isinstance(data, dict):
            match_page = re.search(r"[?&]pageNumber=(\d+)", url)
            match_note_type = re.search(r"[?&]noteType=(\d+)", url)
            page_no = int(match_page.group(1)) if match_page else 1
            note_type = int(match_note_type.group(1)) if match_note_type else 0
            note_detail_cache = cache.setdefault("notes_detail", {})
            note_type_cache = note_detail_cache.setdefault(note_type, {})
            note_type_cache[page_no] = data
            request_cache = cache.setdefault("requests", {})
            request_cache.setdefault("notes_detail", {}).setdefault(note_type, {})[page_no] = _response_request_info(response)
        setattr(page, "_pgy_detail_api_cache", cache)

    try:
        page.on("response", handle_response)
        setattr(page, "_pgy_detail_api_capture_installed", True)
    except Exception:
        pass


def _detail_api_cache(page: Any) -> dict[str, Any]:
    cache = getattr(page, "_pgy_detail_api_cache", None)
    return cache if isinstance(cache, dict) else {}


def _response_request_info(response: Any) -> dict[str, Any]:
    try:
        request = response.request
        return {
            "url": getattr(response, "url", ""),
            "method": getattr(request, "method", "GET"),
            "post_data": getattr(request, "post_data", "") or "",
        }
    except Exception:
        return {"url": getattr(response, "url", ""), "method": "GET", "post_data": ""}


def _url_with_query_params(url: str, params: dict[str, Any]) -> str:
    if not url:
        return url
    parts = urlsplit(url)
    pairs = dict(parse_qsl(parts.query, keep_blank_values=True))
    for key, value in params.items():
        if value is None:
            continue
        pairs[key] = str(value)
    return urlunsplit((parts.scheme, parts.netloc, parts.path, urlencode(pairs), parts.fragment))


def _fetch_detail_api_data(page: Any, request_info: dict[str, Any], query_params: dict[str, Any] | None = None) -> dict[str, Any]:
    if not isinstance(request_info, dict) or not request_info.get("url"):
        return {}
    method = str(request_info.get("method") or "GET").upper()
    url = _url_with_query_params(str(request_info.get("url") or ""), query_params or {})
    body = str(request_info.get("post_data") or "")
    try:
        payload = page.evaluate(
            """
            async ({ url, method, body }) => {
              const init = {
                method,
                credentials: 'include',
                headers: {
                  'accept': 'application/json, text/plain, */*',
                  'content-type': 'application/json;charset=UTF-8'
                }
              };
              if (method !== 'GET' && body) init.body = body;
              const response = await fetch(url, init);
              return await response.json();
            }
            """,
            {"url": url, "method": method, "body": body},
        )
    except Exception:
        return {}
    data = payload.get("data") if isinstance(payload, dict) else None
    return data if isinstance(data, dict) else {}


def _detail_cached_request(page: Any, group: str, business: str = "") -> dict[str, Any]:
    requests = _detail_api_cache(page).get("requests")
    if not isinstance(requests, dict):
        return {}
    group_requests = requests.get(group)
    if isinstance(group_requests, dict) and business:
        request = group_requests.get(business)
        if isinstance(request, dict):
            return request
        for value in group_requests.values():
            if isinstance(value, dict):
                return value
    if isinstance(group_requests, dict):
        for value in group_requests.values():
            if isinstance(value, dict):
                return value
    return {}


def _ensure_detail_summary_api_cache(page: Any) -> None:
    cache = _detail_api_cache(page)
    for group, cache_key in [("data_summary", "data_summary"), ("notes_rate", "notes_rate")]:
        group_cache = cache.get(cache_key) if isinstance(cache.get(cache_key), dict) else {}
        for business, business_value in [("daily", 0), ("cooperation", 1)]:
            if isinstance(group_cache, dict) and isinstance(group_cache.get(business), dict):
                continue
            request = _detail_cached_request(page, group, business)
            data = _fetch_detail_api_data(page, request, {"business": business_value})
            if not data:
                continue
            cache = _detail_api_cache(page)
            next_group_cache = cache.setdefault(cache_key, {})
            next_group_cache[business] = data
            setattr(page, "_pgy_detail_api_cache", cache)


def _first_notes_detail_request(page: Any, note_type: int | None = None) -> dict[str, Any]:
    requests = _detail_api_cache(page).get("requests")
    if not isinstance(requests, dict):
        return {}
    detail_requests = requests.get("notes_detail")
    if not isinstance(detail_requests, dict):
        return {}
    if note_type is not None and isinstance(detail_requests.get(note_type), dict):
        for value in detail_requests[note_type].values():
            if isinstance(value, dict):
                return value
    for type_requests in detail_requests.values():
        if isinstance(type_requests, dict):
            for value in type_requests.values():
                if isinstance(value, dict):
                    return value
    return {}


def _ensure_note_detail_api_pages(page: Any, max_pages: int = 3) -> None:
    cache = _detail_api_cache(page)
    notes_detail = cache.get("notes_detail") if isinstance(cache.get("notes_detail"), dict) else {}
    note_types = list(notes_detail.keys()) if isinstance(notes_detail, dict) and notes_detail else [3, 4]
    for note_type in note_types:
        request = _first_notes_detail_request(page, int(note_type) if str(note_type).isdigit() else None)
        if not request:
            continue
        for page_no in range(1, max_pages + 1):
            cache = _detail_api_cache(page)
            current = cache.get("notes_detail") if isinstance(cache.get("notes_detail"), dict) else {}
            current_type = current.get(note_type) if isinstance(current, dict) else {}
            if isinstance(current_type, dict) and isinstance(current_type.get(page_no), dict):
                continue
            data = _fetch_detail_api_data(page, request, {"noteType": note_type, "pageNumber": page_no})
            items = data.get("list") if isinstance(data, dict) else []
            if not isinstance(items, list) or not items:
                break
            cache = _detail_api_cache(page)
            cache.setdefault("notes_detail", {}).setdefault(note_type, {})[page_no] = data
            setattr(page, "_pgy_detail_api_cache", cache)


def _ratio_from_percent_value(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return _ratio_from_text(str(value))
    return number / 100 if number > 1 else number


def _metric_string(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (int, float)):
        return f"{value:,}"
    return str(value).strip()


def _summary_note_types_text(note_types: Any) -> str:
    if not isinstance(note_types, list):
        return ""
    parts = []
    for item in note_types:
        if not isinstance(item, dict):
            continue
        tag = str(item.get("contentTag") or "").strip()
        percent = str(item.get("percent") or "").strip()
        if tag and percent:
            parts.append(f"{tag}(占比{percent}%)")
        elif tag:
            parts.append(tag)
    return "、".join(parts)


def _build_blogger_advantage_from_api(summary: dict[str, Any], business: str = "daily") -> str:
    if not isinstance(summary, dict):
        return ""
    parts: list[str] = []
    advantage = str(summary.get("kolAdvantage") or "").strip()
    if advantage:
        parts.append(f"{advantage}博主")
    note_number = summary.get("noteNumber")
    if note_number not in (None, ""):
        parts.append(f"发布笔记{note_number}篇")
    note_types = _summary_note_types_text(summary.get("noteType"))
    if note_types:
        parts.append(f"内容类目{note_types}")
    trade_names = [str(item).strip() for item in (summary.get("tradeNames") or []) if str(item).strip()]
    if trade_names:
        parts.append(f"合作行业{'、'.join(trade_names)}")
    return "".join(parts)


def _note_case_from_notes_detail_item(item: dict[str, Any]) -> dict[str, Any]:
    case = {
        "note_id": str(item.get("noteId") or "").strip(),
        "title": str(item.get("title") or "").strip(),
        "brand": str(item.get("brandName") or item.get("contentTag") or "").strip(),
        "cover_url": _normalize_media_url(item.get("imgUrl") or item.get("imageUrl") or item.get("coverUrl")),
        "published_at": str(item.get("date") or item.get("publishTime") or "").strip(),
        "read_count": item.get("readNum"),
        "like_count": item.get("likeNum"),
        "save_count": item.get("collectNum"),
        "comment_count": item.get("cmtNum"),
        "share_count": item.get("shareNum"),
        "has_promoted_traffic": bool(item.get("isAdvertise")),
        "note_type": "视频笔记" if item.get("isVideo") else "图文笔记",
        "source": "notes_detail_api",
    }
    return {key: value for key, value in case.items() if value not in ("", None, [])}


def _api_note_case_pages_from_cache(page: Any, note_type: int = 3, max_cases: int = 24) -> dict[str, Any]:
    cache = _detail_api_cache(page)
    notes_detail = cache.get("notes_detail") if isinstance(cache.get("notes_detail"), dict) else {}
    pages_map = notes_detail.get(note_type) if isinstance(notes_detail, dict) else {}
    if not isinstance(pages_map, dict):
        return {}
    pages: list[dict[str, Any]] = []
    collected: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str]] = set()
    for page_no in sorted(pages_map):
        payload = pages_map.get(page_no)
        items = payload.get("list") if isinstance(payload, dict) else []
        if not isinstance(items, list):
            continue
        cases = [_note_case_from_notes_detail_item(item) for item in items if isinstance(item, dict)]
        pages.append({"page": page_no, "count": len(cases), "cases": cases})
        for case in cases:
            key = _case_key(case)
            if key in seen:
                continue
            seen.add(key)
            collected.append(case)
            if len(collected) >= max_cases:
                break
        if len(collected) >= max_cases:
            break
    if not collected:
        return {}
    return {"cooperation_note_cases": collected, "cooperation_note_case_pages": pages}


def _overview_from_api_cache(page: Any) -> dict[str, Any]:
    cache = _detail_api_cache(page)
    summary_cache = cache.get("data_summary") if isinstance(cache.get("data_summary"), dict) else {}
    result: dict[str, Any] = {}
    for business in ["daily", "cooperation"]:
        summary = summary_cache.get(business)
        if not isinstance(summary, dict):
            continue
        rate_cache = cache.get("notes_rate") if isinstance(cache.get("notes_rate"), dict) else {}
        rate_data = rate_cache.get(business) if isinstance(rate_cache, dict) else {}
        scale_metrics = {
            "曝光中位数": _metric_string(summary.get("mAccumImpNum")),
            "阅读中位数": _metric_string(summary.get("readMedian")),
            "互动中位数": _metric_string(summary.get("mEngagementNum") or summary.get("interactionMedian")),
        }
        if business == "cooperation" and summary.get("estimateVideoCpuv") not in (None, ""):
            scale_metrics["预估阅读单价"] = _metric_string(summary.get("estimateVideoCpuv"))
        if business == "cooperation" and summary.get("estimateVideoEngageCost") not in (None, ""):
            scale_metrics["预估互动单价"] = _metric_string(summary.get("estimateVideoEngageCost"))
        cost_metrics = {
            "预估CPM": _metric_string(summary.get("estimateVideoCpm") or summary.get("estimatePictureCpm")),
            "预估阅读单价": _metric_string(summary.get("videoReadCostV2") or summary.get("estimateVideoCpuv") or summary.get("picReadCost")),
            "预估互动单价": _metric_string(summary.get("estimateVideoEngageCost") or summary.get("estimatePictureEngageCost")),
        }
        state = {
            "text": "",
            "metrics": {key: value for key, value in scale_metrics.items() if value not in ("", None)},
        }
        state_cost = {
            "text": "",
            "metrics": {key: value for key, value in cost_metrics.items() if value not in ("", None)},
        }
        if isinstance(rate_data, dict):
            if rate_data.get("interactionRate") not in (None, ""):
                state["metrics"]["互动率"] = f"{rate_data['interactionRate']}%"
                state_cost["metrics"]["互动率"] = f"{rate_data['interactionRate']}%"
            if rate_data.get("videoFullViewRate") not in (None, ""):
                state["metrics"]["视频完播率"] = f"{rate_data['videoFullViewRate']}%"
                state_cost["metrics"]["视频完播率"] = f"{rate_data['videoFullViewRate']}%"
            if rate_data.get("thousandLikePercent") not in (None, ""):
                state["metrics"]["千赞笔记比例"] = f"{rate_data['thousandLikePercent']}%"
                state_cost["metrics"]["千赞笔记比例"] = f"{rate_data['thousandLikePercent']}%"
            if rate_data.get("hundredLikePercent") not in (None, ""):
                state["metrics"]["百赞笔记比例"] = f"{rate_data['hundredLikePercent']}%"
                state_cost["metrics"]["百赞笔记比例"] = f"{rate_data['hundredLikePercent']}%"
        result[business] = {"scale": state, "cost": state_cost}
    return result


def _performance_from_api_cache(page: Any) -> dict[str, Any]:
    cache = _detail_api_cache(page)
    rate_cache = cache.get("notes_rate") if isinstance(cache.get("notes_rate"), dict) else {}
    result: dict[str, Any] = {}
    for business in ["daily", "cooperation"]:
        rate_data = rate_cache.get(business)
        if not isinstance(rate_data, dict):
            continue
        scale_metrics = {
            "曝光中位数": _metric_string(rate_data.get("impMedian")),
            "阅读中位数": _metric_string(rate_data.get("readMedian")),
            "互动中位数": _metric_string(rate_data.get("interactionMedian")),
            "中位点赞量": _metric_string(rate_data.get("likeMedian")),
            "中位收藏量": _metric_string(rate_data.get("collectMedian")),
            "中位评论量": _metric_string(rate_data.get("commentMedian")),
            "中位分享量": _metric_string(rate_data.get("shareMedian")),
            "中位关注量": _metric_string(rate_data.get("mfollowCnt") or rate_data.get("mFollowCnt")),
        }
        cost_metrics = dict(scale_metrics)
        if rate_data.get("interactionRate") not in (None, ""):
            scale_metrics["互动率"] = f"{rate_data['interactionRate']}%"
            cost_metrics["互动率"] = f"{rate_data['interactionRate']}%"
        if rate_data.get("videoFullViewRate") not in (None, ""):
            scale_metrics["视频完播率"] = f"{rate_data['videoFullViewRate']}%"
            cost_metrics["视频完播率"] = f"{rate_data['videoFullViewRate']}%"
        if rate_data.get("picture3sViewRate") not in (None, ""):
            scale_metrics["图文3秒阅读率"] = f"{rate_data['picture3sViewRate']}%"
            cost_metrics["图文3秒阅读率"] = f"{rate_data['picture3sViewRate']}%"
        if rate_data.get("thousandLikePercent") not in (None, ""):
            scale_metrics["千赞笔记比例"] = f"{rate_data['thousandLikePercent']}%"
            cost_metrics["千赞笔记比例"] = f"{rate_data['thousandLikePercent']}%"
        if rate_data.get("hundredLikePercent") not in (None, ""):
            scale_metrics["百赞笔记比例"] = f"{rate_data['hundredLikePercent']}%"
            cost_metrics["百赞笔记比例"] = f"{rate_data['hundredLikePercent']}%"
        result[business] = {
            "scale": {"text": "", "metrics": {key: value for key, value in scale_metrics.items() if value not in ("", None)}},
            "cost": {"text": "", "metrics": {key: value for key, value in cost_metrics.items() if value not in ("", None)}},
        }
    return result


def _detail_from_api_cache(page: Any, detail: dict[str, Any]) -> dict[str, Any]:
    cache = _detail_api_cache(page)
    if not cache:
        return detail
    result = dict(detail)
    raw = result.get("raw_payload") if isinstance(result.get("raw_payload"), dict) else {}
    profile = cache.get("blogger_profile") if isinstance(cache.get("blogger_profile"), dict) else {}
    fans_summary = cache.get("fans_summary") if isinstance(cache.get("fans_summary"), dict) else {}
    fans_profile = cache.get("fans_profile") if isinstance(cache.get("fans_profile"), dict) else {}
    daily_summary = (cache.get("data_summary") or {}).get("daily") if isinstance(cache.get("data_summary"), dict) else {}
    if isinstance(profile, dict):
        if profile.get("name") and not result.get("nickname"):
            result["nickname"] = str(profile.get("name") or "").strip()
        if profile.get("redId") and not result.get("xiaohongshu_id"):
            result["xiaohongshu_id"] = str(profile.get("redId") or "").strip()
        if profile.get("location") and not result.get("ip_city"):
            result["ip_city"] = str(profile.get("location") or "").strip()
        if profile.get("fansCount") not in (None, "") and result.get("followers_count") in (None, ""):
            result["followers_count"] = _number_from_text(str(profile.get("fansCount")))
        if profile.get("likeCollectCountInfo") not in (None, "") and result.get("liked_collected_count") in (None, ""):
            result["liked_collected_count"] = _number_from_text(str(profile.get("likeCollectCountInfo")))
        if profile.get("picturePrice") not in (None, "") and result.get("quote_price") in (None, ""):
            result["quote_price"] = _number_from_text(str(profile.get("picturePrice")))
        if profile.get("videoPrice") not in (None, "") and result.get("video_quote_price") in (None, ""):
            result["video_quote_price"] = _number_from_text(str(profile.get("videoPrice")))
        tags = [str(item).strip() for item in (profile.get("personalTags") or []) if str(item).strip()]
        if tags and not result.get("persona_tags"):
            result["persona_tags"] = "、".join(tags)
        content_tags = []
        for item in (profile.get("contentTags") or []):
            if not isinstance(item, dict):
                continue
            taxonomy1 = str(item.get("taxonomy1Tag") or "").strip()
            if taxonomy1:
                content_tags.append(taxonomy1)
            for tag in item.get("taxonomy2Tags") or []:
                text = str(tag).strip()
                if text:
                    content_tags.append(text)
        if content_tags and not result.get("creator_type"):
            result["creator_type"] = sanitize_creator_type("/".join(dict.fromkeys(content_tags)))
    if isinstance(daily_summary, dict):
        if daily_summary.get("dateKey"):
            raw["data_updated_to"] = str(daily_summary.get("dateKey") or "")
        if not raw.get("blogger_advantage"):
            raw["blogger_advantage"] = _build_blogger_advantage_from_api(daily_summary, "daily")
        if daily_summary.get("readMedian") not in (None, "") and result.get("daily_read_median") in (None, ""):
            result["daily_read_median"] = float(daily_summary["readMedian"])
        if daily_summary.get("mAccumImpNum") not in (None, "") and result.get("daily_exposure_median") in (None, ""):
            result["daily_exposure_median"] = float(daily_summary["mAccumImpNum"])
        if daily_summary.get("mEngagementNum") not in (None, "") and result.get("daily_interaction_median") in (None, ""):
            result["daily_interaction_median"] = float(daily_summary["mEngagementNum"])
        if daily_summary.get("responseRate") not in (None, "") and result.get("reply_rate_48h") in (None, ""):
            result["reply_rate_48h"] = _ratio_from_percent_value(daily_summary.get("responseRate"))
        if daily_summary.get("activeDayInLast7") not in (None, "") and result.get("active_days_7d") in (None, ""):
            result["active_days_7d"] = float(daily_summary["activeDayInLast7"])
    fan_analysis = raw.get("fan_analysis") if isinstance(raw.get("fan_analysis"), dict) else {}
    if isinstance(fans_summary, dict):
        mappings = {
            "fansIncreaseNum": "fan_growth",
            "fansGrowthRate": "fan_growth_ratio",
            "activeFansRate": "active_fans_ratio",
            "readFansRate": "read_fans_ratio",
            "engageFansRate": "interaction_fans_ratio",
            "payFansUserRate30d": "order_fans_ratio",
        }
        for source_key, target_key in mappings.items():
            value = fans_summary.get(source_key)
            if value in (None, "") or fan_analysis.get(target_key) not in (None, ""):
                continue
            fan_analysis[target_key] = _ratio_from_percent_value(value) if "Rate" in source_key else float(value)
        raw["fan_analysis"] = fan_analysis
        if result.get("active_fans_ratio") in (None, "") and fan_analysis.get("active_fans_ratio") is not None:
            result["active_fans_ratio"] = fan_analysis["active_fans_ratio"]
        if result.get("read_fans_ratio") in (None, "") and fan_analysis.get("read_fans_ratio") is not None:
            result["read_fans_ratio"] = fan_analysis["read_fans_ratio"]
        if result.get("interaction_fans_ratio") in (None, "") and fan_analysis.get("interaction_fans_ratio") is not None:
            result["interaction_fans_ratio"] = fan_analysis["interaction_fans_ratio"]
        if result.get("order_fans_ratio") in (None, "") and fan_analysis.get("order_fans_ratio") is not None:
            result["order_fans_ratio"] = fan_analysis["order_fans_ratio"]
        if result.get("fans_growth_ratio") in (None, "") and fan_analysis.get("fan_growth_ratio") is not None:
            result["fans_growth_ratio"] = fan_analysis["fan_growth_ratio"]
    if isinstance(fans_profile, dict):
        gender = fans_profile.get("gender") if isinstance(fans_profile.get("gender"), dict) else {}
        if gender:
            female = _ratio_from_percent_value(gender.get("female"))
            male = _ratio_from_percent_value(gender.get("male"))
            if female is not None and result.get("female_fans_ratio") in (None, ""):
                result["female_fans_ratio"] = female
            if male is not None and result.get("male_fans_ratio") in (None, ""):
                result["male_fans_ratio"] = male
        ages = fans_profile.get("ages") if isinstance(fans_profile.get("ages"), list) else []
        age_segments = []
        for item in ages:
            if not isinstance(item, dict):
                continue
            label = str(item.get("group") or "").strip()
            ratio = _ratio_from_percent_value(item.get("percent"))
            if not label or ratio is None:
                continue
            age_segments.append({"label": label, "ratio": ratio})
            if label == "25-34" and result.get("fans_25_34_ratio") in (None, ""):
                result["fans_25_34_ratio"] = ratio
            if label == "35-44" and result.get("fans_35_44_ratio") in (None, ""):
                result["fans_35_44_ratio"] = ratio
            if label in {">44", "44岁以上"} and result.get("fans_44_plus_ratio") in (None, ""):
                result["fans_44_plus_ratio"] = ratio
        if age_segments and not result.get("audience_age_distribution"):
            dominant = max(age_segments, key=lambda item: item.get("ratio") or 0)
            result["audience_age_distribution"] = {"segments": age_segments, "dominant": dominant, "source": "detail_api"}
        female = result.get("female_fans_ratio")
        male = result.get("male_fans_ratio")
        if female is not None or male is not None:
            segments = []
            if female is not None:
                segments.append({"label": "女性", "key": "female_fans_ratio", "ratio": female})
            if male is not None:
                segments.append({"label": "男性", "key": "male_fans_ratio", "ratio": male})
            dominant = max(segments, key=lambda item: item.get("ratio") or 0)
            result["audience_gender_distribution"] = {"segments": segments, "dominant": dominant, "source": "detail_api"}
        provinces = fans_profile.get("provinces") if isinstance(fans_profile.get("provinces"), list) else []
        if provinces and not result.get("audience_region_distribution"):
            top_regions = []
            for item in provinces[:10]:
                if not isinstance(item, dict):
                    continue
                label = str(item.get("name") or "").strip()
                ratio = _ratio_from_percent_value(item.get("percent"))
                if label and ratio is not None:
                    top_regions.append({"label": label, "ratio": ratio})
            if top_regions:
                result["audience_region_distribution"] = {
                    "raw_text": "、".join(f"{item['label']}（{round(item['ratio'] * 100, 1)}%）" for item in top_regions[:3]),
                    "source": "detail_api",
                    "top_regions": top_regions,
                    "dominant": top_regions[0],
                    "scope": "province",
                }
        devices = fans_profile.get("devices") if isinstance(fans_profile.get("devices"), list) else []
        if devices and not result.get("audience_device_distribution"):
            top_device = None
            parsed_devices = []
            for item in devices:
                if not isinstance(item, dict):
                    continue
                label = str(item.get("name") or "").strip()
                ratio = _ratio_from_percent_value(item.get("percent"))
                desc = str(item.get("desc") or "").strip()
                if label and ratio is not None:
                    parsed_devices.append({"label": label, "ratio": ratio, "desc": desc})
            if parsed_devices:
                top_device = parsed_devices[0]
                result["audience_device_distribution"] = {
                    "raw_text": f"{top_device['label']}用户占比{round(top_device['ratio'] * 100, 2)}%" + (f"，{top_device['desc']}" if top_device.get("desc") else ""),
                    "source": "detail_api",
                    "dominant": {"label": top_device["label"], "ratio": top_device["ratio"]},
                    "insight": top_device.get("desc") or "",
                }
        interests = fans_profile.get("interests") if isinstance(fans_profile.get("interests"), list) else []
        if interests and not result.get("topic_point"):
            topic_labels = [str(item.get("name") or "").strip() for item in interests[:3] if isinstance(item, dict) and str(item.get("name") or "").strip()]
            if topic_labels:
                result["topic_point"] = "、".join(topic_labels)
    result["raw_payload"] = raw
    return result


def _detail_notes_from_capture(page: Any, max_notes: int = 8, include_text: bool = True) -> list[dict[str, Any]]:
    notes = getattr(page, "_pgy_latest_detail_notes", []) or []
    if not isinstance(notes, list):
        return []
    results: list[dict[str, Any]] = []
    seen: set[str] = set()
    for note in notes:
        if not isinstance(note, dict):
            continue
        key = str(note.get("note_id") or note.get("title") or "").strip()
        if not key or key in seen:
            continue
        seen.add(key)
        merged = dict(note)
        if include_text and note.get("note_id"):
            detail = _fetch_note_detail_from_api(page, str(note.get("note_id")))
            merged = {**merged, **{key: value for key, value in detail.items() if value not in ("", None, [])}}
            if note.get("cover_url") and not merged.get("cover_url"):
                merged["cover_url"] = note["cover_url"]
        if merged.get("cover_url"):
            merged["cover_url"] = _normalize_media_url(merged.get("cover_url"))
        results.append(merged)
        if len(results) >= max_notes:
            break
    return results


def _merge_recent_notes_into_detail(detail: dict[str, Any], page: Any, max_notes: int = 8) -> dict[str, Any]:
    recent_notes = _detail_notes_from_capture(page, max_notes=max_notes)
    if not recent_notes:
        return detail
    raw = detail.get("raw_payload") if isinstance(detail.get("raw_payload"), dict) else {}
    existing = raw.get("recent_notes") if isinstance(raw.get("recent_notes"), list) else []
    by_key: dict[str, dict[str, Any]] = {}
    for item in [*existing, *recent_notes]:
        if not isinstance(item, dict):
            continue
        key = str(item.get("note_id") or item.get("title") or "").strip()
        if not key:
            continue
        by_key[key] = {**by_key.get(key, {}), **item}
    detail["raw_payload"] = {**raw, "recent_notes": list(by_key.values())}
    return detail


def _attach_recent_note_payload(page: Any, row_index: int, creator: dict[str, Any], include_details: bool = False, max_notes: int = 2) -> None:
    kol = _match_api_kol(page, row_index, creator)
    if not kol:
        return
    raw_payload = creator.get("raw_payload") if isinstance(creator.get("raw_payload"), dict) else {}
    note_briefs = _recent_note_briefs_from_kol(kol, max_notes=max_notes)
    if note_briefs:
        raw_payload["recent_note_briefs"] = note_briefs
    if include_details:
        recent_notes = _collect_recent_note_details_from_api(page, kol, max_notes=max_notes)
        if recent_notes:
            raw_payload["recent_notes"] = recent_notes
    if raw_payload:
        creator["raw_payload"] = raw_payload


def _audience_profile_clip(page: Any) -> dict[str, Any] | None:
    try:
        return page.evaluate(
            """
            () => {
              const normalized = value => String(value || '').replace(/\\s+/g, ' ').trim();
              const visible = el => {
                const rect = el.getBoundingClientRect();
                const style = window.getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              };
              const textNodes = label => Array.from(document.querySelectorAll('section, div, h1, h2, h3, h4, span, p'))
                .filter(visible)
                .filter(el => normalized(el.innerText || el.textContent) === label)
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  return ar.top - br.top || ar.left - br.left;
                });
              const heading = textNodes('粉丝画像')[0];
              if (!heading) return null;
              const titleBlock = heading.closest('.content__title') || heading;
              const titleRect = titleBlock.getBoundingClientRect();
              const fansPic = Array.from(document.querySelectorAll('.contentPic.fansPic'))
                .filter(visible)
                .filter(el => {
                  const text = el.innerText || '';
                  const rect = el.getBoundingClientRect();
                  return text.includes('性别分布') && text.includes('年龄分布') && rect.top >= titleRect.top;
                })
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  return ar.top - br.top || ar.left - br.left;
                })[0];
              if (fansPic) {
                const picRect = fansPic.getBoundingClientRect();
                const left = Math.min(titleRect.left, picRect.left);
                const top = Math.max(0, titleRect.top - 12) + window.scrollY;
                const topViewport = Math.max(0, titleRect.top - 12);
                const right = Math.max(titleRect.right, picRect.right);
                const bottom = picRect.bottom + 12;
                const width = right - left;
                const height = bottom - topViewport;
                if (width >= 240 && height >= 160) {
                  return {
                    x: left,
                    y: topViewport,
                    width,
                    height,
                    pageX: left + window.scrollX,
                    pageY: top,
                    source: 'audience_profile_fanspic_clip',
                    rootText: normalized(fansPic.innerText || '').slice(0, 240)
                  };
                }
              }
              const gender = textNodes('性别分布')[0];
              const age = textNodes('年龄分布')[0];
              if (!gender || !age) return null;
              const rootCandidates = Array.from(document.querySelectorAll('section, div'))
                .filter(visible)
                .filter(el => {
                  const text = el.innerText || '';
                  if (!text.includes('粉丝画像') || !text.includes('性别分布') || !text.includes('年龄分布')) return false;
                  const rect = el.getBoundingClientRect();
                  const headingRect = heading.getBoundingClientRect();
                  return rect.top <= headingRect.top && rect.bottom >= age.getBoundingClientRect().bottom;
                })
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  const areaA = ar.width * ar.height;
                  const areaB = br.width * br.height;
                  return areaA - areaB;
                });
              const root = rootCandidates[0];
              if (!root) return null;
              const rootRect = root.getBoundingClientRect();
              const headingRect = heading.getBoundingClientRect();
              const genderRect = gender.getBoundingClientRect();
              const ageRect = age.getBoundingClientRect();
              const allRects = [headingRect, genderRect, ageRect];
              const genderCard = Array.from(document.querySelectorAll('section, div'))
                .filter(visible)
                .filter(el => (el.innerText || '').includes('性别分布'))
                .filter(el => {
                  const rect = el.getBoundingClientRect();
                  return rect.top >= headingRect.top - 20 && rect.left >= rootRect.left - 20 && rect.right <= rootRect.right + 20;
                })
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  return (ar.width * ar.height) - (br.width * br.height);
                })[0];
              const ageCard = Array.from(document.querySelectorAll('section, div'))
                .filter(visible)
                .filter(el => (el.innerText || '').includes('年龄分布'))
                .filter(el => {
                  const rect = el.getBoundingClientRect();
                  return rect.top >= headingRect.top - 20 && rect.left >= rootRect.left - 20 && rect.right <= rootRect.right + 20;
                })
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  return (ar.width * ar.height) - (br.width * br.height);
                })[0];
              const nextSection = textNodes('地域分布').find(el => {
                const rect = el.getBoundingClientRect();
                return rect.top > headingRect.top && rect.left >= rootRect.left - 20 && rect.left <= rootRect.right + 20;
              }) || textNodes('用户设备分布').find(el => {
                const rect = el.getBoundingClientRect();
                return rect.top > headingRect.top && rect.left >= rootRect.left - 20 && rect.left <= rootRect.right + 20;
              });
              const nextTop = nextSection ? nextSection.getBoundingClientRect().top : rootRect.bottom;
              const cardRects = Array.from(root.querySelectorAll('section, div'))
                .filter(visible)
                .map(el => el.getBoundingClientRect())
                .filter(rect => rect.top >= headingRect.top - 8 && rect.top < nextTop && rect.bottom > Math.max(genderRect.bottom, ageRect.bottom))
                .filter(rect => rect.width > 120 && rect.height > 80);
              if (genderCard) allRects.push(genderCard.getBoundingClientRect());
              if (ageCard) allRects.push(ageCard.getBoundingClientRect());
              allRects.push(...cardRects);
              const leftViewport = Math.max(0, Math.min(...allRects.map(rect => rect.left), rootRect.left));
              const topViewport = Math.max(0, headingRect.top - 12);
              const rightViewport = Math.max(...allRects.map(rect => rect.right), rootRect.right);
              const cardsBottom = Math.max(...allRects.map(rect => rect.bottom));
              const bottomViewport = Math.max(cardsBottom + 12, nextTop ? Math.min(cardsBottom + 12, nextTop - 12) : cardsBottom + 12);
              const left = leftViewport;
              const top = topViewport;
              const right = rightViewport;
              const bottom = bottomViewport;
              const width = right - left;
              const height = bottom - top;
              if (width < 240 || height < 160) return null;
              return {
                x: left,
                y: top,
                width,
                height,
                pageX: leftViewport + window.scrollX,
                pageY: topViewport + window.scrollY,
                source: 'audience_profile_gender_age_clip',
                rootText: normalized(root.innerText || '').slice(0, 240)
              };
            }
            """
        )
    except Exception:
        return None


def _prepare_audience_profile_capture_target(page: Any) -> dict[str, Any]:
    try:
        return page.evaluate(
            """
            () => {
              const normalized = value => String(value || '').replace(/\\s+/g, ' ').trim();
              const visible = el => {
                const rect = el.getBoundingClientRect();
                const style = window.getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              };
              const exact = label => Array.from(document.querySelectorAll('section, div, h1, h2, h3, h4, span, p'))
                .filter(visible)
                .filter(el => normalized(el.innerText || el.textContent) === label)
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  return ar.top - br.top || ar.left - br.left;
                });
              const heading = exact('粉丝画像')[0];
              if (!heading) return { ok: false, message: 'missing heading' };
              const roots = Array.from(document.querySelectorAll('section, div'))
                .filter(visible)
                .filter(el => {
                  const text = el.innerText || '';
                  return text.includes('粉丝画像') && text.includes('性别分布') && text.includes('年龄分布');
                })
                .filter(el => {
                  const rect = el.getBoundingClientRect();
                  const headingRect = heading.getBoundingClientRect();
                  return rect.top <= headingRect.top && rect.bottom > headingRect.bottom;
                })
                .sort((a, b) => {
                  const ar = a.getBoundingClientRect();
                  const br = b.getBoundingClientRect();
                  return (ar.width * ar.height) - (br.width * br.height);
                });
              const root = roots[0];
              if (!root) return { ok: false, message: 'missing root' };
              const before = exact('地域分布').find(el => {
                const r = el.getBoundingClientRect();
                const h = heading.getBoundingClientRect();
                return r.top > h.top;
              }) || exact('用户设备分布').find(el => {
                const r = el.getBoundingClientRect();
                const h = heading.getBoundingClientRect();
                return r.top > h.top;
              });
              const previous = document.getElementById('__pgy_fans_profile_capture__');
              if (previous) previous.remove();
              const wrapper = document.createElement('div');
              wrapper.id = '__pgy_fans_profile_capture__';
              wrapper.style.background = '#fff';
              wrapper.style.boxSizing = 'border-box';
              wrapper.style.width = `${root.getBoundingClientRect().width}px`;
              wrapper.style.padding = '0';
              wrapper.style.margin = '0';
              wrapper.style.overflow = 'visible';
              const clone = root.cloneNode(true);
              if (before) {
                const beforeText = normalized(before.innerText || before.textContent);
                const walker = document.createTreeWalker(clone, NodeFilter.SHOW_ELEMENT);
                const toRemove = [];
                while (walker.nextNode()) {
                  const node = walker.currentNode;
                  if (normalized(node.innerText || node.textContent) === beforeText) {
                    let section = node;
                    while (section.parentElement && section.parentElement !== clone) {
                      const text = section.parentElement.innerText || '';
                      if (text.includes(beforeText) && !text.includes('粉丝画像')) {
                        section = section.parentElement;
                        break;
                      }
                      section = section.parentElement;
                    }
                    toRemove.push(section);
                    break;
                  }
                }
                for (const node of toRemove) {
                  let current = node;
                  while (current) {
                    const next = current.nextElementSibling;
                    current.remove();
                    current = next;
                  }
                }
              }
              wrapper.appendChild(clone);
              document.body.appendChild(wrapper);
              wrapper.scrollIntoView({ block: 'center', inline: 'nearest' });
              const rect = wrapper.getBoundingClientRect();
              return { ok: true, selector: '#__pgy_fans_profile_capture__', box: { x: rect.x, y: rect.y, width: rect.width, height: rect.height } };
            }
            """
        )
    except Exception as error:
        return {"ok": False, "message": str(error)}


def _remove_audience_profile_capture_target(page: Any) -> None:
    try:
        page.evaluate("() => document.getElementById('__pgy_fans_profile_capture__')?.remove()")
    except Exception:
        pass


def _scroll_audience_profile_to_viewport_top(page: Any) -> None:
    try:
        page.evaluate(
            """
            () => {
              const normalized = value => String(value || '').replace(/\\s+/g, ' ').trim();
              const visible = el => {
                const rect = el.getBoundingClientRect();
                const style = window.getComputedStyle(el);
                return rect.width > 0 && rect.height > 0
                  && style.display !== 'none'
                  && style.visibility !== 'hidden';
              };
              const heading = Array.from(document.querySelectorAll('section, div, h1, h2, h3, h4, span, p'))
                .filter(visible)
                .find(el => normalized(el.innerText || el.textContent) === '粉丝画像');
              const target = heading?.closest('.content') || heading?.closest('.content__title') || heading;
              if (!target) return false;
              target.scrollIntoView({ block: 'start', inline: 'nearest' });
              const scrollParents = [];
              let parent = target.parentElement;
              while (parent) {
                const style = window.getComputedStyle(parent);
                if (/(auto|scroll)/.test(`${style.overflowY} ${style.overflow}`)) scrollParents.push(parent);
                parent = parent.parentElement;
              }
              for (const parent of scrollParents) {
                const targetRect = target.getBoundingClientRect();
                const parentRect = parent.getBoundingClientRect();
                parent.scrollTop += targetRect.top - parentRect.top - 24;
              }
              window.scrollBy(0, -24);
              return true;
            }
            """
        )
        page.wait_for_timeout(700)
    except Exception:
        pass


def _scroll_to_fans_analysis(page: Any) -> None:
    for label in ["粉丝分析", "粉丝画像"]:
        try:
            target = page.locator(f"text={label}").last
            if target.count():
                target.scroll_into_view_if_needed(timeout=3000)
                page.wait_for_timeout(1200)
                return
        except Exception:
            continue
    try:
        page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
        page.wait_for_timeout(1200)
    except Exception:
        pass


def _capture_audience_profile_screenshot(page: Any, detail: dict[str, Any]) -> dict[str, Any]:
    _scroll_to_fans_analysis(page)
    _scroll_audience_profile_to_viewport_top(page)
    PGY_DETAIL_SCREENSHOT_DIR.mkdir(parents=True, exist_ok=True)
    blogger_id = str(detail.get("pgy_blogger_id") or "unknown")
    nickname = str(detail.get("nickname") or blogger_id)
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    path = PGY_DETAIL_SCREENSHOT_DIR / f"{_safe_filename(nickname)}-{_safe_filename(blogger_id)}-fans-profile-{timestamp}.png"
    clip = _audience_profile_clip(page)
    if not clip:
        return {}
    try:
        _scroll_audience_profile_to_viewport_top(page)
        page.wait_for_timeout(500)
        refreshed = _audience_profile_clip(page)
        if refreshed:
            clip = refreshed
    except Exception:
        pass
    box = {key: float(clip[key]) for key in ["x", "y", "width", "height"]}
    if box["y"] < 120:
        shift = 120 - box["y"]
        box["y"] = 120
        box["height"] = max(80, box["height"] - shift)
    if not box or box.get("width", 0) < 120 or box.get("height", 0) < 80:
        return {}
    original_viewport = None
    try:
        original_viewport = page.viewport_size
        if original_viewport:
            target_width = max(int(original_viewport.get("width") or 0), int(box["x"] + box["width"] + 40))
            target_height = max(int(original_viewport.get("height") or 0), int(box["y"] + box["height"] + 80))
            if target_width != original_viewport.get("width") or target_height != original_viewport.get("height"):
                page.set_viewport_size({"width": target_width, "height": target_height})
                page.wait_for_timeout(500)
                refreshed = _audience_profile_clip(page)
                if refreshed:
                    clip = refreshed
                    box = {key: float(clip[key]) for key in ["x", "y", "width", "height"]}
        page.screenshot(path=str(path), clip=box)
    except Exception:
        return {}
    finally:
        if original_viewport:
            try:
                page.set_viewport_size(original_viewport)
            except Exception:
                pass
    source = clip.get("source")
    return {
        "audience_profile_screenshot": str(path),
        "raw_payload": {
            "audience_profile_screenshot": {
                "path": str(path),
                "dom_only": True,
                "scope": "粉丝画像-性别分布-年龄分布",
                "source": source,
                "box": box,
            }
        },
    }


def _merge_detail_payload(detail: dict[str, Any], extra: dict[str, Any]) -> dict[str, Any]:
    if not extra:
        return detail
    merged = {**detail, **{key: value for key, value in extra.items() if key != "raw_payload"}}
    raw = {}
    if isinstance(detail.get("raw_payload"), dict):
        raw.update(detail["raw_payload"])
    if isinstance(extra.get("raw_payload"), dict):
        raw.update(extra["raw_payload"])
    if raw:
        merged["raw_payload"] = raw
    return merged


def _annotate_quote_sources(detail: dict[str, Any], source: str) -> dict[str, Any]:
    if not detail:
        return detail
    annotated = {**detail}
    raw = annotated.get("raw_payload") if isinstance(annotated.get("raw_payload"), dict) else {}
    quote_sources = raw.get("quote_sources") if isinstance(raw.get("quote_sources"), dict) else {}
    if annotated.get("quote_price") not in (None, ""):
        quote_sources["quote_price"] = source
    if annotated.get("video_quote_price") not in (None, ""):
        quote_sources["video_quote_price"] = source
    if quote_sources:
        annotated["raw_payload"] = {**raw, "quote_sources": quote_sources}
    return annotated


def _finalize_detail_payload(detail: dict[str, Any]) -> dict[str, Any]:
    if not detail:
        return detail
    finalized = {**detail}
    raw = finalized.get("raw_payload") if isinstance(finalized.get("raw_payload"), dict) else {}
    summary = _build_detail_collection_summary(finalized)
    finalized["detail_collection_summary"] = summary
    finalized["information_completeness"] = {
        "module_count": summary.get("module_count", 0),
        "note_case_count": summary.get("note_case_count", 0),
        "has_audience_chart": summary.get("has_audience_chart", False),
        "has_region_distribution": summary.get("has_region_distribution", False),
        "has_device_distribution": summary.get("has_device_distribution", False),
    }
    finalized["raw_payload"] = {**raw, "detail_collection_summary": summary}
    return finalized


def _collect_first_detail_for_creator(context: Any, row: Any, detail_url: str = "") -> dict[str, Any]:
    if detail_url and "/blogger-detail/" in detail_url:
        return _collect_detail_by_url(context, detail_url)

    existing_pages = set(context.pages)
    trigger = row.locator(".profile").first
    if not trigger.count():
        trigger = row.locator(".kol-name").first
    if not trigger.count():
        return {}
    try:
        with context.expect_page(timeout=4000) as page_info:
            trigger.click(timeout=2000)
        detail_page = page_info.value
    except Exception:
        try:
            trigger.click(timeout=2000)
            detail_page = next((page for page in context.pages if page not in existing_pages and "/blogger-detail/" in page.url), None)
        except Exception:
            detail_page = None
    if not detail_page:
        return {}
    try:
        _install_detail_api_capture(detail_page)
        _install_detail_notes_capture(detail_page)
        detail_page.wait_for_load_state("domcontentloaded", timeout=10000)
        _wait_for_detail_ready(detail_page)
        text = detail_page.locator("body").inner_text(timeout=5000)
        detail = _extract_detail_fields(text, detail_page.url)
        detail = {**_detail_url_fields(detail_page.url, source="detail_page"), **detail}
        detail = _detail_from_api_cache(detail_page, detail)
        detail = _merge_recent_notes_into_detail(detail, detail_page)
        detail = _merge_detail_payload(detail, _collect_audience_profile_chart_metrics(detail_page))
        detail = _merge_detail_payload(detail, _collect_detail_interaction_states(detail_page))
        detail = _merge_detail_payload(detail, _capture_audience_profile_screenshot(detail_page, detail))
        detail = _merge_note_details_into_payload(detail)
        detail = _annotate_note_cases_with_traffic_reference(detail)
        detail = _annotate_quote_sources(detail, "pgy_detail_one_price")
        return _finalize_detail_payload(detail)
    finally:
        try:
            detail_page.close()
        except Exception:
            pass


def _collect_detail_by_url(context: Any, url: str) -> dict[str, Any]:
    if not url or "pgy.xiaohongshu.com" not in url:
        return {}
    detail_page = context.new_page()
    try:
        _install_detail_api_capture(detail_page)
        _install_detail_notes_capture(detail_page)
        detail_page.goto(url, wait_until="domcontentloaded", timeout=30000)
        _wait_for_detail_ready(detail_page)
        text = detail_page.locator("body").inner_text(timeout=5000)
        detail = _extract_detail_fields(text, detail_page.url)
        detail = {**_detail_url_fields(detail_page.url, source="detail_url"), **detail}
        detail = _detail_from_api_cache(detail_page, detail)
        detail = _merge_recent_notes_into_detail(detail, detail_page)
        detail = _merge_detail_payload(detail, _collect_audience_profile_chart_metrics(detail_page))
        detail = _merge_detail_payload(detail, _collect_detail_interaction_states(detail_page))
        detail = _merge_detail_payload(detail, _capture_audience_profile_screenshot(detail_page, detail))
        detail = _merge_note_details_into_payload(detail)
        detail = _annotate_note_cases_with_traffic_reference(detail)
        detail = _annotate_quote_sources(detail, "pgy_detail_one_price")
        return _finalize_detail_payload(detail)
    finally:
        try:
            detail_page.close()
        except Exception:
            pass


def _detail_ready_from_cache(page: Any) -> bool:
    cache = _detail_api_cache(page)
    if not cache:
        return False
    has_profile = isinstance(cache.get("blogger_profile"), dict) and bool(cache.get("blogger_profile"))
    has_fans = isinstance(cache.get("fans_summary"), dict) and bool(cache.get("fans_summary"))
    has_notes = False
    notes_detail = cache.get("notes_detail")
    if isinstance(notes_detail, dict):
        has_notes = any(
            isinstance(page_map, dict) and any(isinstance(payload, dict) and payload for payload in page_map.values())
            for page_map in notes_detail.values()
        )
    data_summary = cache.get("data_summary")
    has_summary = isinstance(data_summary, dict) and any(isinstance(item, dict) and item for item in data_summary.values())
    if has_profile and (has_notes or has_summary or has_fans):
        return True
    return False


def _detail_ready_from_dom(page: Any) -> bool:
    try:
        return bool(
            page.locator("body").evaluate(
                """
                body => {
                  const text = body?.innerText || '';
                  const lines = text.split('\\n').map(item => item.trim()).filter(Boolean);
                  const index = lines.indexOf('笔记主页');
                  const nickname = index >= 0 && lines[index + 1] === '直播主页' ? lines[index + 2] : '';
                  const hasNickname = Boolean(nickname)
                    && !['小红书号：', '小红书号', '数据概览', '笔记数据', '粉丝分析'].includes(nickname);
                  return hasNickname
                    && text.includes('小红书号：')
                    && text.includes('粉丝数')
                    && (text.includes('数据概览') || text.includes('笔记数据') || text.includes('粉丝分析'));
                }
                """
            )
        )
    except Exception:
        return False


def _wait_for_detail_ready(page: Any, timeout_ms: int = 4500) -> None:
    deadline = time.perf_counter() + max(0.5, timeout_ms / 1000)
    while time.perf_counter() < deadline:
        has_notes = bool(getattr(page, "_pgy_latest_detail_notes", []) or [])
        cache = _detail_api_cache(page)
        notes_cache = cache.get("notes_detail") if isinstance(cache.get("notes_detail"), dict) else {}
        if isinstance(notes_cache, dict) and notes_cache:
            has_notes = True
        if (_detail_ready_from_cache(page) or _detail_ready_from_dom(page)) and has_notes:
            page.wait_for_timeout(350)
            return
        page.wait_for_timeout(250)
    page.wait_for_timeout(500)


def _target_detail_url(target: dict[str, Any]) -> str:
    url = str(target.get("pgy_url") or target.get("profile_url") or "").strip()
    return url if "/blogger-detail/" in url else ""


def _blogger_id_from_detail_url(url: str) -> str:
    match = re.search(r"/blogger-detail/([^?/#]+)", str(url or ""))
    return match.group(1) if match else ""


def _detail_payload_matches_target(target: dict[str, Any], detail: dict[str, Any], url: str) -> bool:
    expected = str(target.get("pgy_blogger_id") or _blogger_id_from_detail_url(url) or "").strip()
    actual = str(detail.get("pgy_blogger_id") or "").strip()
    if expected and actual and expected != actual:
        return False
    summary = detail.get("detail_collection_summary") if isinstance(detail.get("detail_collection_summary"), dict) else {}
    if summary and int(summary.get("module_count") or 0) <= 1:
        return False
    return bool(
        detail.get("nickname")
        or detail.get("followers_count") is not None
        or int(summary.get("module_count") or 0) >= 3
    )


def _collect_detail_target_by_url(target: dict[str, Any]) -> dict[str, Any]:
    creator_id = str(target.get("creator_id") or "")
    url = _target_detail_url(target)
    item_started_at = _now_text()
    item_started_perf = time.perf_counter()
    if not url:
        return {
            "ok": False,
            "creator_id": creator_id,
            "nickname": target.get("nickname"),
            "message": "本地没有可打开的详情页链接",
            "detail_collection_timing": _elapsed_timing(item_started_at, item_started_perf),
        }
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        return {
            "ok": False,
            "creator_id": creator_id,
            "nickname": target.get("nickname"),
            "message": "当前环境未安装 Playwright，无法执行真实页面采集",
            "detail_collection_timing": _elapsed_timing(item_started_at, item_started_perf),
        }
    try:
        with sync_playwright() as p:
            browser = p.chromium.connect_over_cdp("http://127.0.0.1:9222")
            context = browser.contexts[0] if browser.contexts else browser.new_context()
            detail = _collect_detail_by_url(context, url)
        item_timing = _elapsed_timing(item_started_at, item_started_perf)
        if detail and _detail_payload_matches_target(target, detail, url):
            return {
                "ok": True,
                "creator": {**target, **detail, "creator_id": creator_id},
                "detail_collection_timing": item_timing,
            }
        return {
            "ok": False,
            "creator_id": creator_id,
            "nickname": target.get("nickname"),
            "message": "详情页解析失败",
            "detail_collection_timing": item_timing,
        }
    except Exception as error:
        return {
            "ok": False,
            "creator_id": creator_id,
            "nickname": target.get("nickname"),
            "message": f"详情页解析失败：{str(error)[:180]}",
            "detail_collection_timing": _elapsed_timing(item_started_at, item_started_perf),
        }


def _creator_list_rows(page: Any) -> Any:
    return page.locator(".blogger-list_list .d-new-table tbody tr").filter(has_not=page.locator(".skeleton-block"))


def _creator_collection_key(creator: dict[str, Any]) -> str:
    for field in ["pgy_blogger_id", "pgy_url", "profile_url", "xiaohongshu_id", "creator_id"]:
        value = str(creator.get(field) or "").strip()
        if value:
            return f"{field}:{value}"
    return "|".join(
        str(creator.get(field) or "").strip()
        for field in ["nickname", "followers_count", "quote_price", "ip_city"]
    )


def _creator_match_keys(creator: dict[str, Any]) -> set[str]:
    keys: set[str] = set()
    for field in ["pgy_blogger_id", "pgy_url", "profile_url", "xiaohongshu_id"]:
        value = str(creator.get(field) or "").strip()
        if value:
            keys.add(f"{field}:{value}")
    nickname = _clean_text(str(creator.get("nickname") or ""))
    if nickname:
        keys.add(f"nickname:{nickname}")
        location = _clean_text(str(creator.get("ip_city") or ""))
        if location:
            keys.add(f"nickname_location:{nickname}:{location}")
    return keys


def _creator_table_signature(page: Any) -> str:
    try:
        rows = _creator_list_rows(page)
        count = min(rows.count(), 8)
        parts = []
        for index in range(count):
            parts.append(_clean_text(rows.nth(index).inner_text(timeout=800))[:240])
        return "||".join(parts)
    except Exception:
        return ""


def _has_valid_creator_rows(page: Any) -> bool:
    try:
        rows = _creator_list_rows(page)
        count = min(rows.count(), 6)
    except Exception:
        return False
    for index in range(count):
        try:
            text = rows.nth(index).inner_text(timeout=800)
        except Exception:
            continue
        if _parse_row_text(text, page.url):
            return True
    return False


def _reset_creator_list_scroll(page: Any) -> None:
    try:
        page.evaluate(
            """
            () => {
              const body = document.querySelector('.solar_body');
              if (body) body.scrollTop = 0;
              const list = document.querySelector('.blogger-list_list');
              if (list) list.scrollTop = 0;
              window.scrollTo(0, 0);
            }
            """
        )
    except Exception:
        pass


def _wait_for_creator_table_change(page: Any, before_signature: str, timeout_ms: int = 8000) -> bool:
    rounds = max(1, int(timeout_ms / 500))
    for _ in range(rounds):
        page.wait_for_timeout(500)
        try:
            rows = _creator_list_rows(page)
            if rows.count() and _has_valid_creator_rows(page) and _creator_table_signature(page) != before_signature:
                _reset_creator_list_scroll(page)
                return True
        except Exception:
            continue
    return False


def _pagination_page_number(text: str) -> int | None:
    digits = re.findall(r"\d+", text or "")
    if not digits:
        return None
    if len(digits) >= 2 and len(set(digits)) == 1:
        digits = [digits[0]]
    try:
        return int(digits[0])
    except ValueError:
        return None


def _click_pagination_element(page: Any, candidate: Any) -> bool:
    try:
        candidate.scroll_into_view_if_needed(timeout=1200)
    except Exception:
        pass
    try:
        candidate.click(timeout=1800)
    except Exception:
        try:
            candidate.evaluate("node => node.click()")
        except Exception:
            try:
                box = candidate.bounding_box(timeout=1000)
                if not box:
                    return False
                page.mouse.click(box["x"] + box["width"] / 2, box["y"] + box["height"] / 2)
            except Exception:
                return False
    page.wait_for_timeout(900)
    _reset_creator_list_scroll(page)
    return True


def _click_next_creator_page(page: Any, target_page: int | None = None) -> bool:
    page_buttons = page.locator(".d-pagination .d-pagination-page")
    try:
        button_count = min(page_buttons.count(), 80)
    except Exception:
        button_count = 0
    if target_page is not None:
        for index in range(button_count):
            candidate = page_buttons.nth(index)
            try:
                if not _is_visible(candidate):
                    continue
                class_name = (candidate.get_attribute("class") or "").lower()
                if "disabled" in class_name:
                    continue
                number = _pagination_page_number(candidate.inner_text(timeout=500))
                if number == target_page and _click_pagination_element(page, candidate):
                    return True
            except Exception:
                continue
    active_index = -1
    for index in range(button_count):
        candidate = page_buttons.nth(index)
        try:
            class_name = (candidate.get_attribute("class") or "").lower()
            text = _clean_text(candidate.inner_text(timeout=500))
            if ("primary-light" in class_name or "actived" in class_name or "active" in class_name) and text:
                active_index = index
                break
        except Exception:
            continue
    if active_index >= 0:
        for index in range(active_index + 1, button_count):
            candidate = page_buttons.nth(index)
            try:
                if not _is_visible(candidate):
                    continue
                class_name = (candidate.get_attribute("class") or "").lower()
                if "disabled" in class_name:
                    continue
                text = _clean_text(candidate.inner_text(timeout=500))
                if not text or text == "...":
                    continue
                if _click_pagination_element(page, candidate):
                    return True
            except Exception:
                continue

    candidates = page.locator("button, a, li, div[role=button], span[role=button]")
    try:
        count = min(candidates.count(), 300)
    except Exception:
        return False
    for index in range(count):
        candidate = candidates.nth(index)
        try:
            if not _is_visible(candidate):
                continue
            text = _clean_text(candidate.inner_text(timeout=500))
            title = _clean_text(candidate.get_attribute("title") or "")
            aria = _clean_text(candidate.get_attribute("aria-label") or "")
            class_name = (candidate.get_attribute("class") or "").lower()
            descriptor = f"{text}|{title}|{aria}|{class_name}"
            is_next = (
                "下一页" in descriptor
                or "next" in descriptor
                or "pagination-next" in descriptor
                or text in {">", "›", "»"}
            )
            if not is_next or "上一页" in descriptor or "prev" in descriptor:
                continue
            disabled = (
                candidate.get_attribute("disabled") is not None
                or str(candidate.get_attribute("aria-disabled") or "").lower() == "true"
                or "disabled" in class_name
            )
            if disabled:
                continue
            if _click_pagination_element(page, candidate):
                return True
        except Exception:
            continue
    return False


def _scroll_creator_list(page: Any) -> bool:
    try:
        result = page.evaluate(
            """
            () => {
              const root = document.querySelector('.blogger-list_list') || document.body;
              const scrollables = Array.from(root.querySelectorAll('*'))
                .filter(el => el.scrollHeight > el.clientHeight + 8);
              const target = scrollables
                .sort((a, b) => (b.clientHeight * b.clientWidth) - (a.clientHeight * a.clientWidth))[0];
              if (!target) {
                const before = Math.round(window.scrollY || document.documentElement.scrollTop || 0);
                window.scrollBy(0, Math.max(480, Math.floor(window.innerHeight * 0.75)));
                const after = Math.round(window.scrollY || document.documentElement.scrollTop || 0);
                return { changed: after > before, before, after };
              }
              const before = Math.round(target.scrollTop);
              target.scrollTop = Math.min(
                target.scrollTop + Math.max(480, Math.floor(target.clientHeight * 0.85)),
                target.scrollHeight
              );
              const after = Math.round(target.scrollTop);
              return { changed: after > before, before, after };
            }
            """
        )
        return bool(result and result.get("changed"))
    except Exception:
        return False


def _extract_current_creator_page(
    page: Any,
    creators: list[dict[str, Any]],
    seen: set[str],
    limit: int,
    include_details: bool = False,
    collect_profile_urls: bool = False,
    detail_limit: int | None = None,
    page_number: int = 1,
) -> int:
    rows = _creator_list_rows(page)
    try:
        count = rows.count()
    except Exception:
        count = 0
    context = page.context
    headers = _extract_table_headers(page)
    added = 0
    for index in range(min(count, max(limit * 2, 120))):
        row = rows.nth(index)
        try:
            text = row.inner_text(timeout=3000)
        except Exception:
            continue
        table_payload = _extract_row_table_payload(row, headers)
        creator = _parse_row_text(text, page.url, table_payload=table_payload)
        if not creator:
            continue
        creator.update(_extract_row_link_fields(row))
        if not creator.get("pgy_url"):
            creator.update(_api_kol_link_fields(page, index, creator))
        key = _creator_collection_key(creator)
        if key in seen:
            continue
        seen.add(key)
        raw_payload = creator.get("raw_payload") if isinstance(creator.get("raw_payload"), dict) else {}
        creator["raw_payload"] = {**raw_payload, "collection_page": page_number, "collection_row_index": index + 1}
        max_detail_count = detail_limit if detail_limit is not None else limit
        _attach_recent_note_payload(
            page,
            index,
            creator,
            include_details=include_details and len(creators) < max_detail_count,
            max_notes=2,
        )
        avatar = row.locator("img.head-photo").first
        if avatar.count():
            creator["avatar_url"] = avatar.get_attribute("src") or ""
        if include_details and len(creators) < max_detail_count:
            detail = _collect_first_detail_for_creator(context, row, str(creator.get("pgy_url") or creator.get("profile_url") or ""))
            if detail:
                raw_payload = creator.get("raw_payload") if isinstance(creator.get("raw_payload"), dict) else {}
                detail_raw = detail.pop("raw_payload", {})
                for price_field in ("quote_price", "video_quote_price"):
                    if price_field in detail and detail.get(price_field) not in (None, ""):
                        raw_payload[f"list_{price_field}"] = creator.get(price_field)
                creator.update({key: value for key, value in detail.items() if value not in ("", None)})
                creator["raw_payload"] = {**raw_payload, "detail": detail_raw}
        creators.append(creator)
        added += 1
        if len(creators) >= limit:
            break
    return added


def _extract_visible_creators(
    page: Any,
    limit: int,
    include_details: bool = False,
    collect_profile_urls: bool = True,
    detail_limit: int | None = None,
    api_first: bool = True,
    progress_callback: Any = None,
    stop_callback: Any = None,
) -> list[dict[str, Any]]:
    if _stop_requested(stop_callback):
        return []
    if api_first:
        for mode, prepare in [
            ("api_snapshot", lambda: bool(_kol_request_from_snapshot(getattr(page, "_pgy_latest_kol_request", {}) or {}))),
            ("api_browser_capture", lambda: _capture_kol_api_from_browser(page)),
            ("api_template", lambda: _restore_kol_request_from_template(page)),
        ]:
            if _stop_requested(stop_callback):
                return []
            _emit_collection_progress(
                progress_callback,
                {"mode": f"{mode}_start", "collected_count": 0, "limit": limit},
            )
            if not prepare():
                _emit_collection_progress(
                    progress_callback,
                    {"mode": f"{mode}_failed", "collected_count": 0, "limit": limit},
                )
                continue
            api_creators = _collect_creators_from_api_pool(page, limit, progress_callback=progress_callback, mode=mode, stop_callback=stop_callback)
            if api_creators:
                return api_creators
            _emit_collection_progress(
                progress_callback,
                {"mode": f"{mode}_failed", "collected_count": 0, "limit": limit},
            )
        _emit_collection_progress(
            progress_callback,
            {"mode": "dom_fallback", "collected_count": 0, "limit": limit},
        )
    creators: list[dict[str, Any]] = []
    seen: set[str] = set()
    page_number = 1
    idle_rounds = 0
    max_rounds = max(4, min(120, (limit // 20) + 8))
    api_matches_rows = _api_pool_matches_current_rows(page)
    if api_matches_rows:
        _append_api_creators(page, creators, seen, limit)
    for _ in range(max_rounds):
        if _stop_requested(stop_callback):
            break
        before_count = len(creators)
        _extract_current_creator_page(
            page,
            creators,
            seen,
            limit,
            include_details=include_details,
            collect_profile_urls=collect_profile_urls,
            detail_limit=detail_limit,
            page_number=page_number,
        )
        if api_matches_rows:
            _append_api_creators(page, creators, seen, limit)
        if len(creators) >= limit:
            break
        added = len(creators) - before_count
        api_appended = _append_expanded_api_creators(page, creators, seen, limit, stop_callback=stop_callback) if api_matches_rows else 0
        if api_appended:
            idle_rounds = 0
            if len(creators) >= limit:
                break
            continue
        before_signature = _creator_table_signature(page)
        clicked_next = _click_next_creator_page(page, target_page=page_number + 1)
        if clicked_next:
            if _wait_for_creator_table_change(page, before_signature, timeout_ms=12000):
                page_number += 1
                idle_rounds = 0
                _emit_collection_progress(
                    progress_callback,
                    {"mode": "dom", "page_number": page_number, "collected_count": len(creators), "limit": limit},
                )
                api_matches_rows = _api_pool_matches_current_rows(page)
                if api_matches_rows:
                    _append_api_creators(page, creators, seen, limit)
                continue
            page.wait_for_timeout(2500)
            if _has_valid_creator_rows(page) and _creator_table_signature(page) != before_signature:
                page_number += 1
                idle_rounds = 0
                _emit_collection_progress(
                    progress_callback,
                    {"mode": "dom", "page_number": page_number, "collected_count": len(creators), "limit": limit},
                )
                api_matches_rows = _api_pool_matches_current_rows(page)
                if api_matches_rows:
                    _append_api_creators(page, creators, seen, limit)
                continue
            if added > 0:
                page_number += 1
                idle_rounds = 0
                _emit_collection_progress(
                    progress_callback,
                    {"mode": "dom_no_change", "page_number": page_number, "collected_count": len(creators), "limit": limit},
                )
                continue
        if _scroll_creator_list(page):
            page.wait_for_timeout(1000)
            if _creator_table_signature(page) != before_signature:
                idle_rounds = 0
                continue
        api_added = _append_expanded_api_creators(page, creators, seen, limit, stop_callback=stop_callback) if api_matches_rows else 0
        if api_added:
            idle_rounds = 0
            if len(creators) >= limit:
                break
            continue
        idle_rounds = idle_rounds + 1 if added == 0 else 0
        if idle_rounds >= 2 or added == 0:
            break
    if creators:
        _emit_collection_progress(
            progress_callback,
            {"mode": "dom_done", "page_number": page_number, "collected_count": len(creators), "limit": limit},
        )
    return creators


def _click_text_if_visible(page: Any, text: str) -> bool:
    locator = page.locator(".blogger-list_filter").get_by_text(text, exact=True)
    if not locator.count():
        locator = page.get_by_text(text, exact=True)
    clicked = _click_first_visible(locator)
    if clicked:
        page.wait_for_timeout(500)
    return clicked


def _click_filter_checkbox(page: Any, text: str) -> bool:
    checkbox = page.locator(".blogger-list_filter .d-checkbox").filter(has_text=text)
    try:
        count = min(checkbox.count(), 20)
    except Exception:
        return False
    for index in range(count):
        candidate = checkbox.nth(index)
        if "disabled" in ((candidate.get_attribute("class") or "").lower()):
            continue
        simulator = candidate.locator(".d-checkbox-simulator").first
        simulator_class = simulator.get_attribute("class") if simulator.count() else ""
        if "checked" in (simulator_class or ""):
            return True
        try:
            candidate.scroll_into_view_if_needed(timeout=1200)
            candidate.click(timeout=1500)
            page.wait_for_timeout(500)
            return True
        except Exception:
            continue
    return False


def _click_filter_tag(page: Any, text: str) -> bool:
    tag = page.locator(".blogger-list_filter .tag, .blogger-list_filter .selector-body-options *").filter(has_text=re.compile(f"^{re.escape(text)}$"))
    try:
        count = min(tag.count(), 30)
    except Exception:
        return False
    for index in range(count):
        candidate = tag.nth(index)
        class_name = candidate.get_attribute("class") or ""
        if "disabled" in class_name:
            continue
        if "active" in class_name:
            return True
        try:
            candidate.scroll_into_view_if_needed(timeout=1200)
            candidate.click(timeout=1500)
            page.wait_for_timeout(700)
            return True
        except Exception:
            continue
    return False


def _click_blogger_category_subcategory(page: Any, main_value: str, sub_value: str) -> bool:
    if not main_value or not sub_value:
        return False
    if _blogger_category_parent_all_selected(page, main_value):
        _remove_selected_filter_chip(page, ["博主类目", main_value, "全部"])
    main_pattern = re.compile(f"^\\s*{re.escape(main_value)}\\s*$")
    main_locators = [
        page.locator(".blogger-list_filter .tag, .blogger-list_filter .selector-body-options *").filter(has_text=main_pattern),
        page.get_by_text(main_value, exact=True),
    ]
    for main_locator in main_locators:
        try:
            count = min(main_locator.count(), 20)
        except Exception:
            continue
        for index in range(count):
            main = main_locator.nth(index)
            try:
                if not _is_visible(main):
                    continue
                main.scroll_into_view_if_needed(timeout=1200)
                try:
                    main.hover(timeout=1500)
                except Exception:
                    pass
                try:
                    box = main.bounding_box(timeout=800)
                    if box:
                        page.mouse.move(box["x"] + box["width"] / 2, box["y"] + box["height"] / 2)
                except Exception:
                    pass
                try:
                    main.evaluate(
                        """
                        node => {
                          for (const type of ['mouseenter', 'mouseover', 'mousemove']) {
                            node.dispatchEvent(new MouseEvent(type, { bubbles: true, view: window }));
                          }
                        }
                        """
                    )
                except Exception:
                    pass
                page.wait_for_timeout(500)
                for selector in [
                    ".d-popover .tag",
                    ".d-popover .d-checkbox",
                    ".d-popover *",
                    ".filter-select-popover .tag",
                    ".filter-select-popover *",
                ]:
                    sub = page.locator(selector).filter(has_text=re.compile(f"^\\s*{re.escape(sub_value)}\\s*$")).first
                    if not sub.count() or not _is_visible(sub):
                        continue
                    if _click_locator(page, sub, timeout=1800):
                        page.wait_for_timeout(700)
                        return True
            except Exception:
                continue
    return False


def _apply_blogger_category_filter(page: Any, item: dict[str, Any]) -> tuple[bool, str]:
    main_value = str(item.get("value") or "").strip()
    sub_value = str(item.get("sub_value") or item.get("subValue") or "").strip()
    if not main_value:
        return False, "博主类目缺少主类目"
    if sub_value:
        valid_subcategories = PGY_BLOGGER_CATEGORY_TAXONOMY.get(main_value) or []
        if sub_value not in valid_subcategories:
            return False, f"博主类目二级类目不在白名单：{main_value}-{sub_value}"
        if _click_blogger_category_subcategory(page, main_value, sub_value):
            return True, f"已选择博主类目：{main_value}-{sub_value}"
        return False, f"页面未找到或未能选择博主二级类目：{main_value}-{sub_value}"
    if _click_filter_tag(page, main_value) or _click_text_if_visible(page, main_value):
        return True, f"已选择博主类目：{main_value}"
    return False, f"页面未找到博主类目：{main_value}"


def _find_filter_trigger(page: Any, field: str) -> Any:
    selectors = [
        ".blogger-list_filter .custom-selector__button",
        ".blogger-list_filter button.dropdown-button",
        ".blogger-list_filter .tag",
    ]
    field_pattern = re.compile(f"^\\s*(?:新\\s*)?{re.escape(field)}\\s*$")
    for selector in selectors:
        locator = page.locator(selector).filter(has_text=field_pattern)
        try:
            count = min(locator.count(), 10)
        except Exception:
            continue
        for index in range(count):
            candidate = locator.nth(index)
            if _is_visible(candidate):
                return candidate
    return None


def _popover_for_trigger(page: Any, trigger: Any) -> Any:
    try:
        box = trigger.bounding_box(timeout=1000) or {}
        trigger_x = float(box.get("x") or 0)
        trigger_y = float(box.get("y") or 0)
    except Exception:
        trigger_x = 0
        trigger_y = 0
    popovers = page.locator(".d-popover, .filter-select-popover")
    try:
        count = min(popovers.count(), 120)
    except Exception:
        return None
    best = None
    best_score = 10**9
    for index in reversed(range(count)):
        popover = popovers.nth(index)
        try:
            box = popover.bounding_box(timeout=800) or {}
            x = float(box.get("x") or 0)
            y = float(box.get("y") or 0)
            width = float(box.get("width") or 0)
            height = float(box.get("height") or 0)
            if width <= 0 or height <= 0:
                continue
        except Exception:
            continue
        score = abs(x - trigger_x) + abs(y - trigger_y)
        if score < best_score:
            best = popover
            best_score = score
    return best


def _open_filter_popover(page: Any, field: str) -> tuple[Any, str]:
    trigger = _find_filter_trigger(page, field)
    if trigger is None:
        return None, "页面未找到筛选控件"
    try:
        _reset_creator_list_scroll(page)
        page.wait_for_timeout(150)
        if not _click_locator(page, trigger, timeout=1800):
            return None, "打开筛选控件失败：点击未生效"
        page.wait_for_timeout(600)
    except Exception as error:
        return None, f"打开筛选控件失败：{error}"
    popover = _popover_for_trigger(page, trigger)
    if popover is None:
        return None, "筛选弹层未出现"
    return popover, "已打开筛选弹层"


def _click_popover_confirm(page: Any, popover: Any) -> bool:
    confirm = popover.locator("button, .d-button").filter(has_text=re.compile("^\\s*确定\\s*$")).last
    try:
        if confirm.count():
            confirm.click(timeout=1500)
            page.wait_for_timeout(800)
            return True
    except Exception:
        pass
    try:
        page.keyboard.press("Escape")
    except Exception:
        pass
    return False


def _marketing_goal_parent(item: dict[str, Any]) -> str:
    value = _item_selection_value(item)
    parent = str(item.get("goal") or item.get("parent_value") or item.get("parentValue") or "").strip()
    if not parent and value in PGY_MARKETING_GOAL_DEFAULT_METRIC:
        parent = value
    if not parent:
        parent = PGY_MARKETING_GOAL_METRIC_PARENT.get(value) or ""
    return parent


def _marketing_goal_metric(item: dict[str, Any]) -> str:
    value = _item_selection_value(item)
    return PGY_MARKETING_GOAL_DEFAULT_METRIC.get(value) or value


def _click_marketing_goal_parent(page: Any, parent: str) -> tuple[Any, str]:
    if not parent:
        return None, "营销目标缺少父级目标"
    for selector in [
        ".blogger-list_filter .tag",
        ".blogger-list_filter .selector-body-options *",
        ".blogger-list_filter button",
        ".blogger-list_filter *",
    ]:
        candidates = page.locator(selector).filter(has_text=re.compile(f"^\\s*{re.escape(parent)}\\s*$"))
        try:
            count = min(candidates.count(), 30)
        except Exception:
            continue
        for index in range(count):
            candidate = candidates.nth(index)
            try:
                if not _is_visible(candidate):
                    continue
                if not _click_locator(page, candidate, timeout=1800):
                    continue
                page.wait_for_timeout(650)
                popover = _popover_for_trigger(page, candidate) or _last_visible_popover(page)
                if popover is not None:
                    return popover, "已打开营销目标弹层"
            except Exception:
                continue
    return None, f"页面未找到营销目标父级：{parent}"


def _apply_marketing_goal_group(page: Any, parent: str, items: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    applied: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    pending = [item for item in items if not _filter_already_selected(page, item)]
    for item in items:
        if item not in pending:
            applied.append({**item, "message": "页面已存在该筛选条件"})
    if not pending:
        return applied, skipped

    popover, message = _click_marketing_goal_parent(page, parent)
    if popover is None:
        for item in pending:
            skipped.append({**item, "message": message})
        return applied, skipped

    successes: list[dict[str, Any]] = []
    failures: list[dict[str, Any]] = []
    for item in pending:
        metric = _marketing_goal_metric(item)
        selected = False
        for candidate in _filter_value_candidates({**item, "value": metric}):
            if _select_popover_checkbox(page, popover, candidate):
                selected = True
                break
        if selected:
            successes.append({**item, "value": metric, "message": f"已选择{parent}下的{metric}"})
        else:
            failures.append({**item, "value": metric, "message": f"营销目标弹层内未找到指标：{metric}"})

    _click_popover_confirm(page, popover)
    return [*applied, *successes], [*skipped, *failures]


def _select_popover_checkbox(page: Any, popover: Any, text: str) -> bool:
    for selector in [".d-checkbox", ".d-dropdown-option", ".d-options *", ".tag", "label", "button", "*"]:
        option = popover.locator(selector).filter(has_text=re.compile(f"^\\s*{re.escape(text)}\\s*$")).first
        try:
            if not option.count() or not _is_visible(option):
                continue
            if selector == ".d-checkbox" and _checkbox_checked(option):
                return True
            if not _click_locator(page, option, timeout=1500):
                continue
            page.wait_for_timeout(300)
            return True
        except Exception:
            continue
    return False


def _click_popover_group(page: Any, popover: Any, group_label: str) -> bool:
    if not group_label:
        return False
    selectors = [
        ".range-select-content__left .range-select-item",
        ".d-new-cascader__option-list__item:first-child .d-new-cascader__option-wrapper",
        ".d-cascader-menu:first-child .d-cascader-menu-item",
        ".d-new-cascader__option-wrapper",
        ".range-select-item",
    ]
    pattern = re.compile(f"^\\s*{re.escape(group_label)}\\s*$")
    for selector in selectors:
        candidates = popover.locator(selector).filter(has_text=pattern)
        try:
            count = min(candidates.count(), 20)
        except Exception:
            continue
        for index in range(count):
            candidate = candidates.nth(index)
            try:
                if not _is_visible(candidate):
                    continue
                class_name = candidate.get_attribute("class") or ""
                if "--active" in class_name or "active" in class_name:
                    return True
                if _click_locator(page, candidate, timeout=1500):
                    page.wait_for_timeout(450)
                    return True
            except Exception:
                continue
    return False


def _select_grouped_popover_value(page: Any, popover: Any, item: dict[str, Any]) -> bool:
    field = str(item.get("field") or "")
    value = _item_selection_value(item)
    group = str(item.get("sub_field") or item.get("subField") or "")
    if not group:
        group = _infer_option_group(field, value)
    if group:
        _click_popover_group(page, popover, group)
    for candidate in _filter_value_candidates({**item, "value": value}):
        if _select_popover_checkbox(page, popover, candidate):
            return True
    if group:
        return False
    for inferred_group in [
        str(group_item.get("label") or "")
        for group_item in _catalog_option_groups(field)
        if value in (group_item.get("options") or [])
    ]:
        if _click_popover_group(page, popover, inferred_group) and _select_popover_checkbox(page, popover, value):
            return True
    return False


PGY_FOREIGN_REGION_OPTIONS = {"美国", "日本", "澳大利亚", "英国", "加拿大", "韩国", "法国", "德国", "新加坡", "其他"}
PGY_CHINA_REGION_OPTIONS = {
    "北京", "上海", "天津", "重庆", "河北", "山西", "内蒙古", "辽宁", "吉林", "黑龙江",
    "江苏", "浙江", "安徽", "福建", "江西", "山东", "河南", "湖北", "湖南", "广东",
    "广西", "海南", "四川", "贵州", "云南", "西藏", "陕西", "甘肃", "青海", "宁夏",
    "新疆", "香港", "澳门", "台湾",
}
PGY_REGION_ALIASES = {
    "北京/上海优先": ["北京", "上海"],
    "北京上海优先": ["北京", "上海"],
    "北上广深": ["北京", "上海", "广州", "深圳"],
    "一线": ["北京", "上海", "广州", "深圳"],
    "一线城市": ["北京", "上海", "广州", "深圳"],
}


def _region_targets_from_item(item: dict[str, Any]) -> list[str]:
    raw_values: list[str] = []
    for key in ("regions", "region_values", "input_values"):
        values = item.get(key)
        if isinstance(values, list):
            raw_values.extend(str(value) for value in values)
    for key in ("city", "province", "country"):
        value = str(item.get(key) or "").strip()
        if value:
            raw_values.append(value)
    raw_value = str(item.get("value") or "")
    raw_values.extend(PGY_REGION_ALIASES.get(raw_value) or [])
    raw_values.extend(re.split(r"[、,，/|｜\s]+", raw_value))
    targets: list[str] = []
    for value in raw_values:
        clean = value.strip(" ：:;；。优先重点必须地域要求地区城市IPip")
        if not clean or clean in {"中国", "国内", "全国", "不限"}:
            continue
        if clean in PGY_REGION_ALIASES:
            targets.extend(PGY_REGION_ALIASES[clean])
            continue
        if clean.endswith("市") and len(clean) <= 4:
            clean = clean[:-1]
        if clean.endswith("省") and len(clean) <= 4:
            clean = clean[:-1]
        if clean:
            targets.append(clean)
    return list(dict.fromkeys(targets))


def _click_cascade_option(page: Any, popover: Any, text: str, column_index: int | None = None, checkbox: bool | None = None) -> bool:
    if not text:
        return False
    column_selectors = [
        ".d-new-cascader__option-list__item",
        ".d-cascader-menu",
        ".range-select-content__column",
        ".range-select-content__left, .range-select-content__middle, .range-select-content__right",
    ]
    scopes: list[Any] = []
    if column_index is not None:
        for selector in column_selectors:
            columns = popover.locator(selector)
            try:
                if columns.count() > column_index:
                    scopes.append(columns.nth(column_index))
            except Exception:
                continue
    scopes.append(popover)
    selectors = [".d-checkbox", "label", ".d-new-cascader__option-wrapper", ".d-cascader-menu-item", ".range-select-item", "button", "*"]
    pattern = re.compile(f"^\\s*{re.escape(text)}\\s*$")
    for scope in scopes:
        for selector in selectors:
            candidates = scope.locator(selector).filter(has_text=pattern)
            try:
                count = min(candidates.count(), 30)
            except Exception:
                continue
            for index in range(count):
                candidate = candidates.nth(index)
                try:
                    if not _is_visible(candidate):
                        continue
                    if checkbox is True and _checkbox_checked(candidate):
                        return True
                    if _click_locator(page, candidate, timeout=1500):
                        page.wait_for_timeout(450)
                        return True
                except Exception:
                    continue
    return False


def _apply_region_cascade(page: Any, item: dict[str, Any]) -> tuple[bool, str]:
    field = str(item.get("field") or "地域")
    popover, message = _open_filter_popover(page, field)
    if popover is None:
        return False, message
    targets = _region_targets_from_item(item)
    if not targets:
        _click_popover_confirm(page, popover)
        return False, "地域条件未解析到可勾选国家/省市，已保留在采集计划中"

    selected: list[str] = []
    failed: list[str] = []
    china_targets = [target for target in targets if target not in PGY_FOREIGN_REGION_OPTIONS]
    foreign_targets = [target for target in targets if target in PGY_FOREIGN_REGION_OPTIONS]

    if china_targets:
        _click_cascade_option(page, popover, "全部", column_index=0, checkbox=True)
        _click_cascade_option(page, popover, "全部", column_index=0)
        if _click_cascade_option(page, popover, "中国", column_index=0, checkbox=True):
            for target in china_targets:
                clicked = False
                candidates = [target]
                if target in {"广州", "深圳"}:
                    candidates = ["广东", target]
                for candidate in candidates:
                    if _click_cascade_option(page, popover, candidate, column_index=1, checkbox=True):
                        clicked = True
                        break
                    if _click_cascade_option(page, popover, candidate, column_index=2, checkbox=True):
                        clicked = True
                        break
                if clicked:
                    selected.append(target)
                else:
                    failed.append(target)
        else:
            failed.extend(china_targets)

    for target in foreign_targets:
        if _click_cascade_option(page, popover, target, column_index=0, checkbox=True):
            selected.append(target)
        else:
            failed.append(target)

    _click_popover_confirm(page, popover)
    if selected:
        suffix = f"；未找到：{'、'.join(failed)}" if failed else ""
        return True, f"已选择地域：{'、'.join(selected)}{suffix}"
    return False, f"地域弹层内未找到匹配选项：{'、'.join(failed or targets)}"


def _visible_popovers(page: Any, selector: str = ".d-popover, .filter-select-popover") -> list[Any]:
    popovers = page.locator(selector)
    try:
        count = min(popovers.count(), 160)
    except Exception:
        return []
    visible: list[Any] = []
    for index in range(count):
        candidate = popovers.nth(index)
        if _is_visible(candidate):
            visible.append(candidate)
    return visible


def _last_visible_popover(page: Any, selector: str = ".d-popover, .filter-select-popover") -> Any:
    visible = _visible_popovers(page, selector)
    return visible[-1] if visible else None


def _subfield_aliases(field: str, sub_field: str) -> list[str]:
    aliases = [sub_field]
    if field == "合作报价" and sub_field == "图文笔记":
        aliases.extend(["图文报价", "图文笔记报价", "图文笔记一口价"])
    if field == "合作报价" and sub_field == "视频笔记":
        aliases.extend(["视频报价", "视频笔记报价", "视频笔记一口价"])
    if field == "预估互动单价" and sub_field == "图文笔记互动单价":
        aliases.append("预估图文互动单价")
    if field == "预估互动单价" and sub_field == "视频笔记互动单价":
        aliases.append("预估视频互动单价")
    if field == "预估阅读单价" and sub_field == "图文预估阅读单价":
        aliases.append("图文笔记阅读单价")
    if field == "预估阅读单价" and sub_field == "视频预估阅读单价":
        aliases.append("视频笔记阅读单价")
    return list(dict.fromkeys([item for item in aliases if item]))


def _find_subfield_selector(popover: Any, field: str, sub_field: str) -> Any:
    items = popover.locator(".filters-item")
    aliases = [_clean_text(item) for item in _subfield_aliases(field, sub_field)]
    try:
        count = min(items.count(), 20)
    except Exception:
        count = 0
    for index in range(count):
        item = items.nth(index)
        try:
            text = _clean_text(item.inner_text(timeout=800))
        except Exception:
            text = ""
        if aliases and not any(alias and alias in text for alias in aliases):
            continue
        selector = item.locator(".custom-selector__button, .d-input-wrapper, input[readonly]").first
        try:
            if selector.count() and _is_visible(selector):
                return selector
        except Exception:
            continue
    selectors = popover.locator(".custom-selector__button, input[readonly]")
    try:
        if selectors.count():
            return selectors.first
    except Exception:
        pass
    return None


def _fill_nested_number_range(nested: Any, min_value: Any, max_value: Any) -> int:
    inputs = nested.locator("input[type=text], input[type=number], input:not([type])")
    try:
        count = min(inputs.count(), 8)
    except Exception:
        return 0
    editable: list[Any] = []
    for index in range(count):
        input_box = inputs.nth(index)
        try:
            readonly = input_box.get_attribute("readonly")
            disabled = input_box.get_attribute("disabled")
            if readonly is not None or disabled is not None or not _is_visible(input_box):
                continue
            editable.append(input_box)
        except Exception:
            continue
    target_inputs = editable[-2:]
    values = [min_value, max_value] if len(target_inputs) >= 2 else [max_value]
    filled = 0
    for input_box, value in zip(target_inputs, values):
        if value is None or value == "":
            continue
        text_value = _format_filter_number(value)
        try:
            input_box.fill(text_value, timeout=1500)
            filled += 1
        except Exception:
            try:
                input_box.click(timeout=1000)
                input_box.press("Control+A")
                input_box.type(text_value, timeout=1500)
                filled += 1
            except Exception:
                try:
                    input_box.evaluate(
                        """
                        (node, value) => {
                          node.removeAttribute('readonly');
                          node.removeAttribute('disabled');
                          node.value = value;
                          node.dispatchEvent(new Event('input', { bubbles: true }));
                          node.dispatchEvent(new Event('change', { bubbles: true }));
                        }
                        """,
                        text_value,
                    )
                    filled += 1
                except Exception:
                    continue
    return filled


def _click_nested_confirm(page: Any, nested: Any) -> bool:
    confirm = nested.locator("button, .d-button").filter(has_text=re.compile("^\\s*确定\\s*$")).last
    try:
        if confirm.count() and _click_locator(page, confirm, timeout=1500):
            page.wait_for_timeout(700)
            return True
    except Exception:
        pass
    return False


def _apply_single_subfield_number_range(page: Any, popover: Any, item: dict[str, Any], sub_field: str) -> tuple[bool, str]:
    field = item.get("field") or ""
    value = str(item.get("value") or "")
    selector = _find_subfield_selector(popover, field, sub_field)
    if selector is None:
        return False, "弹层内未找到子字段选择框，已保留在采集计划中"
    if not _click_locator(page, selector, timeout=1500):
        return False, "子字段下拉未打开，已保留在采集计划中"
    page.wait_for_timeout(600)
    nested = _last_visible_popover(page, ".d-popover.filters-item-custom, .d-popover .filters-item-custom, .filter-select-popover")
    if nested is None:
        return False, "子字段区间弹层未出现，已保留在采集计划中"
    min_value, max_value = _range_for_subfield(item, sub_field)
    if _range_policy_for_item(item, sub_field) != "min_only" and max_value in (None, ""):
        match = re.search(r"[≤<]\s*([0-9]+(?:\.[0-9]+)?)", value)
        max_value = match.group(1) if match else ""
    filled = _fill_nested_number_range(nested, min_value, max_value)
    required_fills = 2 if min_value not in (None, "") and max_value not in (None, "") else 1
    if filled < required_fills:
        _click_nested_confirm(page, nested)
        return False, "子字段区间输入未完整填写，已保留在采集计划中"
    if not _click_nested_confirm(page, nested):
        return False, "子字段区间确认失败，已保留在采集计划中"
    if min_value not in (None, "") and max_value not in (None, ""):
        range_text = f"{_format_filter_number(min_value)}～{_format_filter_number(max_value)}"
    elif min_value not in (None, ""):
        range_text = f"{_format_filter_number(min_value)}以上"
    else:
        range_text = f"{_format_filter_number(max_value)}以下"
    return True, f"{sub_field or field} {range_text}"


def _apply_subfield_number_range(page: Any, item: dict[str, Any]) -> tuple[bool, str]:
    field = item.get("field") or ""
    popover, message = _open_filter_popover(page, field)
    if popover is None:
        return False, message
    default_subfields: str | list[str] = ["图文笔记", "视频笔记"] if field == "合作报价" else str(item.get("sub_field") or "")
    sub_fields = _subfield_names_from_item(item, default=default_subfields)
    if not sub_fields:
        sub_fields = [str(item.get("sub_field") or "")]
    successes: list[str] = []
    failures: list[str] = []
    for sub_field in sub_fields:
        success, detail = _apply_single_subfield_number_range(page, popover, item, sub_field)
        if success:
            successes.append(detail)
        else:
            failures.append(f"{sub_field or field}：{detail}")
    if not successes:
        _click_popover_confirm(page, popover)
        return False, "；".join(failures) or "子字段区间输入失败，已保留在采集计划中"
    _click_popover_confirm(page, popover)
    suffix = f"；部分失败：{'；'.join(failures)}" if failures else ""
    return True, f"已填写{field}自定义区间：{'；'.join(successes)}{suffix}"


def _apply_preset_or_number_range(page: Any, item: dict[str, Any]) -> tuple[bool, str]:
    field = item.get("field") or ""
    value = str(item.get("value") or "")
    popover, message = _open_filter_popover(page, field)
    if popover is None:
        return False, message
    if _range_policy_for_item(item) != "min_only":
        candidates = _filter_value_candidates(item)
        for candidate in [candidate for candidate in dict.fromkeys(candidates) if candidate]:
            if _select_popover_checkbox(page, popover, candidate):
                _click_popover_confirm(page, popover)
                return True, "已在弹层中选择区间筛选项"
    parsed_min, parsed_max = _range_numbers_from_text(value)
    min_value = item.get("min", parsed_min if parsed_min is not None else "")
    max_value = item.get("max", parsed_max if parsed_max is not None else "")
    min_value, max_value = _apply_range_policy(item, min_value, max_value)
    if min_value in (None, "") and max_value in (None, ""):
        _click_popover_confirm(page, popover)
        return False, "弹层内未找到匹配区间，已保留在采集计划中"
    filled = _fill_nested_number_range(popover, min_value, max_value)
    if filled < (2 if min_value not in (None, "") and max_value not in (None, "") else 1):
        _click_popover_confirm(page, popover)
        return False, "区间输入未完整填写，已保留在采集计划中"
    if not _click_popover_confirm(page, popover):
        return False, "区间筛选确认失败，已保留在采集计划中"
    if min_value not in (None, "") and max_value not in (None, ""):
        range_text = f"{_format_filter_number(min_value)}～{_format_filter_number(max_value)}"
    elif min_value not in (None, ""):
        range_text = f"{_format_filter_number(min_value)}以上"
    else:
        range_text = f"{_format_filter_number(max_value)}以下"
    return True, f"已填写{field}自定义区间 {range_text}"


def _fill_popover_text_inputs(popover: Any, values: list[str]) -> int:
    filled = 0
    if not values:
        return filled
    inputs = popover.locator("input[type=text], textarea")
    try:
        count = min(inputs.count(), len(values))
    except Exception:
        return filled
    for index in range(count):
        value = values[index]
        if not value:
            continue
        try:
            input_box = inputs.nth(index)
            input_box.fill(str(value), timeout=1500)
            filled += 1
        except Exception:
            continue
    return filled


def _apply_popover_filter_item(page: Any, item: dict[str, Any]) -> tuple[bool, str]:
    field = item.get("field") or ""
    value = str(item.get("value") or "")
    control_type = item.get("control_type") or ""
    if control_type in {"marketing_goal_metric"}:
        parent = _marketing_goal_parent(item)
        applied, skipped = _apply_marketing_goal_group(page, parent, [item])
        if applied:
            return True, applied[-1].get("message") or "已选择营销目标指标"
        return False, (skipped[-1].get("message") if skipped else "营销目标筛选失败，已保留在采集计划中")
    if control_type in {"text_multi_with_exclude", "searchable_multi_select_with_exclude"}:
        popover, message = _open_filter_popover(page, field)
        if popover is None:
            return False, message
        values = item.get("input_values")
        if not isinstance(values, list):
            values = [part.strip() for part in re.split(r"[、,，/]+", value) if part.strip()]
        filled = _fill_popover_text_inputs(popover, [str(part) for part in values])
        if item.get("exclude") or "剔除" in value:
            _select_popover_checkbox(page, popover, "剔除上述品牌已合作博主")
        _click_popover_confirm(page, popover)
        if filled:
            suffix = "；品牌不足3个，已保留待补充" if filled < int(item.get("min_items") or 0) else ""
            return True, f"已填写文本筛选项{suffix}"
        return False, "该筛选项需要填写品牌名称，已保留待补充"
    if control_type in {"checkbox_popover", "dropdown", "select_popover", "single_select_popover", "dropdown_single", "nested_select_popover"}:
        popover, message = _open_filter_popover(page, field)
        if popover is None:
            return False, message
        if control_type == "nested_select_popover":
            _click_popover_confirm(page, popover)
            return False, "该条件打开后仍有下拉选项，需要继续选择，已保留在采集计划中"
        if _select_grouped_popover_value(page, popover, item):
            _click_popover_confirm(page, popover)
            return True, "已在弹层中选择筛选项"
        _click_popover_confirm(page, popover)
        return False, "弹层内未找到匹配选项，已保留在采集计划中"
    if control_type in {"brand_search_recommendation"}:
        return False, "该条件需要在右上角合作品牌/竞品搜索入口选择品牌，已保留在采集计划中"
    if control_type in {"subfield_preset_or_number_range", "subfield_preset_or_percent_range", "multi_subfield_preset_or_number_range"}:
        return _apply_subfield_number_range(page, item)
    if control_type in {"range_select_pair", "number_range", "preset_or_number_range", "preset_or_percent_range"}:
        return _apply_preset_or_number_range(page, item)
    if control_type in {"cascade_checkbox_popover", "three_level_cascade_checkbox_popover"}:
        if field in {"地域", "粉丝地域"}:
            return _apply_region_cascade(page, item)
        return False, "该条件需要区间/子筛选/级联细分，已保留在采集计划中"
    return False, ""


def _apply_filter_item(page: Any, item: dict[str, Any]) -> tuple[bool, str]:
    value = item.get("value") or ""
    field = item.get("field") or ""
    if _filter_already_selected(page, item):
        return True, "页面已存在该筛选条件"
    if field == "博主类目":
        return _apply_blogger_category_filter(page, item)
    popover_success, popover_message = _apply_popover_filter_item(page, item)
    if popover_success or popover_message:
        return popover_success, popover_message
    candidates = _filter_value_candidates(item)
    if field in {"平台推荐", "常规剔除"}:
        for candidate in candidates:
            if _click_filter_checkbox(page, candidate):
                return True, "已勾选复选筛选项"
        return False, "未找到可勾选的复选筛选项"
    for candidate in candidates:
        if _click_filter_tag(page, candidate) or _click_text_if_visible(page, candidate):
            return True, "已点击页面筛选项"
    if value.endswith("优先") or "/" in value or value in {"预估阅读/互动单价"}:
        return False, "该条件需要下拉/区间细分，已保留在采集计划中"
    return False, "页面未找到可直接点击的筛选项"


def _checkbox_checked(checkbox: Any) -> bool:
    input_box = checkbox.locator("input[type=checkbox]").first
    try:
        if input_box.count():
            return bool(input_box.is_checked(timeout=500))
    except Exception:
        pass
    for attr in ["aria-checked", "data-checked", "checked"]:
        try:
            value = input_box.get_attribute(attr) if input_box.count() else checkbox.get_attribute(attr)
        except Exception:
            value = None
        if str(value).lower() in {"true", "checked", "1"}:
            return True
    simulator = checkbox.locator(".d-checkbox-simulator").first
    try:
        if simulator.count() and "checked" in (simulator.get_attribute("class") or "").lower():
            return True
    except Exception:
        pass
    try:
        return "checked" in (checkbox.get_attribute("class") or "").lower()
    except Exception:
        return False


def _radio_checked(radio: Any) -> bool:
    simulator = radio.locator(".d-radio-simulator").first
    if not simulator.count():
        return False
    return "checked" in (simulator.get_attribute("class") or "")


def _metric_candidates(metric: str) -> list[str]:
    aliases = PGY_METRIC_ALIASES.get(metric) or []
    if metric.endswith("（日常）"):
        aliases.append(metric.replace("（日常）", ""))
    return list(dict.fromkeys([metric, *aliases]))


def _visible_table_text(page: Any) -> str:
    try:
        table = page.locator(".blogger-list_list .d-new-table, .blogger-list_list").first
        if table.count():
            return table.inner_text(timeout=2500)
    except Exception:
        pass
    try:
        return page.locator("body").inner_text(timeout=2500)
    except Exception:
        return ""


def _metric_already_displayed(page: Any, metric: str) -> bool:
    text = _clean_text(_visible_table_text(page))
    if not text:
        return False
    if metric in {PGY_ALL_NON_LIVE_METRICS, "__ALL_NON_LIVE__"}:
        return all(_clean_text(item) in text for item in PGY_REQUIRED_NON_LIVE_METRICS)
    return any(_clean_text(candidate) in text for candidate in _metric_candidates(metric))


def _display_metrics_already_ready(page: Any, metrics: list[str]) -> bool:
    return bool(metrics) and all(_metric_already_displayed(page, metric) for metric in metrics)


def _selected_modal_metrics(modal: Any) -> set[str]:
    try:
        text = modal.evaluate(
            """
            node => {
              const blocks = Array.from(node.querySelectorAll('*'))
                .filter(el => /已添加/.test(el.textContent || ''));
              const right = blocks
                .sort((a, b) => (b.getBoundingClientRect().left - a.getBoundingClientRect().left))[0] || node;
              return right.innerText || right.textContent || '';
            }
            """
        )
    except Exception:
        try:
            text = modal.inner_text(timeout=1200)
        except Exception:
            text = ""
    lines = _visible_text_lines(str(text or ""))
    ignore = {"已添加", "清空", "以上为横向固定列", "博主信息", "操作", "近期笔记"}
    return {_clean_text(line) for line in lines if line and line not in ignore and not line.startswith("已添加")}


def _required_modal_metrics(metrics: list[str]) -> list[str]:
    required: list[str] = []
    for metric in metrics:
        if metric in {PGY_ALL_NON_LIVE_METRICS, "__ALL_NON_LIVE__"}:
            required.extend(PGY_REQUIRED_NON_LIVE_METRICS)
        else:
            required.append(metric)
    return list(dict.fromkeys(required))


def _modal_metrics_ready(modal: Any, metrics: list[str]) -> bool:
    selected_text = "".join(_selected_modal_metrics(modal))
    if not selected_text:
        return False
    for metric in _required_modal_metrics(metrics):
        candidates = _metric_candidates(metric)
        if not any(_clean_text(candidate) in selected_text for candidate in candidates):
            return False
    return True


def _ensure_modal_option(page: Any, modal: Any, metric: str) -> tuple[bool, str]:
    for candidate in _metric_candidates(metric):
        checkbox = modal.locator(".d-checkbox").filter(has_text=re.compile(f"^{re.escape(candidate)}$")).first
        if checkbox.count():
            if _checkbox_checked(checkbox):
                return True, "已选中"
            try:
                checkbox.scroll_into_view_if_needed(timeout=2000)
                checkbox.click(timeout=2000)
                page.wait_for_timeout(250)
                if _checkbox_checked(checkbox):
                    return True, "已勾选"
                return False, "点击后未变为选中"
            except Exception as error:
                return False, f"勾选失败：{error}"
        radio = modal.locator(".d-radio-group .d-radio, .d-radio-group > *").filter(has_text=re.compile(f"^{re.escape(candidate)}$")).first
        if radio.count():
            if _radio_checked(radio):
                return True, "已选中"
            try:
                radio.scroll_into_view_if_needed(timeout=2000)
                radio.click(timeout=2000)
                page.wait_for_timeout(250)
                return (True, "已选择") if _radio_checked(radio) else (False, "点击后未变为选中")
            except Exception as error:
                return False, f"选择失败：{error}"
    return False, "弹窗内未找到该指标"


def _modal_scroll_state(modal: Any) -> tuple[int, int]:
    try:
        payload = modal.evaluate(
            """
            node => {
              const scrollables = Array.from(node.querySelectorAll('*'))
                .filter(el => el.scrollHeight > el.clientHeight + 8);
              const target = scrollables
                .sort((a, b) => (b.clientHeight * b.clientWidth) - (a.clientHeight * a.clientWidth))[0];
              if (!target) return { top: 0, max: 0 };
              return { top: Math.round(target.scrollTop), max: Math.round(target.scrollHeight - target.clientHeight) };
            }
            """
        )
        return int(payload.get("top") or 0), int(payload.get("max") or 0)
    except Exception:
        return 0, 0


def _scroll_modal_options(modal: Any) -> bool:
    try:
        before, max_top = _modal_scroll_state(modal)
        after = modal.evaluate(
            """
            node => {
              const scrollables = Array.from(node.querySelectorAll('*'))
                .filter(el => el.scrollHeight > el.clientHeight + 8);
              const target = scrollables
                .sort((a, b) => (b.clientHeight * b.clientWidth) - (a.clientHeight * a.clientWidth))[0];
              if (!target) return 0;
              target.scrollTop = Math.min(target.scrollTop + Math.max(320, Math.floor(target.clientHeight * 0.75)), target.scrollHeight);
              return Math.round(target.scrollTop);
            }
            """
        )
        return int(after or 0) > before and before < max_top
    except Exception:
        return False


def _ensure_all_non_live_modal_options(page: Any, modal: Any) -> dict[str, list[dict[str, str]]]:
    selected: list[dict[str, str]] = []
    skipped: list[dict[str, str]] = []
    seen: set[str] = set()
    idle_rounds = 0
    for _ in range(25):
        before_seen_count = len(seen)
        checkboxes = modal.locator(".d-checkbox")
        try:
            count = min(checkboxes.count(), 500)
        except Exception:
            count = 0
        for index in range(count):
            checkbox = checkboxes.nth(index)
            try:
                label = " ".join(_visible_text_lines(checkbox.inner_text(timeout=800)))
            except Exception:
                label = f"指标{index + 1}"
            if not label or label in seen:
                continue
            seen.add(label)
            class_name = (checkbox.get_attribute("class") or "").lower()
            try:
                checked = _checkbox_checked(checkbox)
            except Exception:
                checked = False
            if "直播" in label:
                skipped.append({"metric": label, "message": "直播数据已排除；保持当前选择状态" if checked else "直播数据已排除"})
                continue
            if "disabled" in class_name and not checked:
                skipped.append({"metric": label, "message": "指标不可选"})
                continue
            try:
                checkbox.scroll_into_view_if_needed(timeout=1500)
                if checked:
                    selected.append({"metric": label, "message": "已选中，未重复点击"})
                    continue
                checkbox.click(timeout=1500)
                page.wait_for_timeout(150)
                if _checkbox_checked(checkbox):
                    selected.append({"metric": label, "message": "已勾选"})
                else:
                    skipped.append({"metric": label, "message": "点击后未变为选中，已跳过以避免误取消"})
            except Exception as error:
                skipped.append({"metric": label, "message": f"勾选失败：{error}"})
        if not _scroll_modal_options(modal):
            break
        page.wait_for_timeout(300)
        idle_rounds = idle_rounds + 1 if len(seen) == before_seen_count else 0
        if idle_rounds >= 3:
            break
    if not selected and not skipped:
        skipped.append({"metric": PGY_ALL_NON_LIVE_METRICS, "message": "弹窗内没有识别到指标项"})
    return {"selected_metrics": selected, "skipped_metrics": skipped}


def _ensure_display_metrics(page: Any, metrics: list[str]) -> dict[str, list[dict[str, str]]]:
    selected: list[dict[str, str]] = []
    skipped: list[dict[str, str]] = []
    if not metrics:
        return {"selected_metrics": selected, "skipped_metrics": skipped}
    if _display_metrics_already_ready(page, metrics):
        return {
            "selected_metrics": [{"metric": metric, "message": "页面已展示，跳过重复勾选"} for metric in metrics],
            "skipped_metrics": skipped,
        }
    button = page.locator("button").filter(has_text="选择展示指标").last
    if not button.count():
        return {
            "selected_metrics": selected,
            "skipped_metrics": [{"metric": metric, "message": "页面未找到选择展示指标按钮"} for metric in metrics],
        }
    try:
        button.scroll_into_view_if_needed(timeout=2000)
        button.click(timeout=3000)
        page.wait_for_timeout(800)
    except Exception as error:
        return {
            "selected_metrics": selected,
            "skipped_metrics": [{"metric": metric, "message": f"打开展示指标弹窗失败：{error}"} for metric in metrics],
        }
    modal = page.locator(".d-modal.custom-data-model, .d-modal").last
    try:
        modal.wait_for(state="visible", timeout=5000)
    except Exception:
        return {
            "selected_metrics": selected,
            "skipped_metrics": [{"metric": metric, "message": "展示指标弹窗未出现"} for metric in metrics],
        }
    if _modal_metrics_ready(modal, metrics):
        try:
            page.keyboard.press("Escape")
        except Exception:
            pass
        return {
            "selected_metrics": [{"metric": metric, "message": "自定义列已保留完整配置，跳过重复勾选"} for metric in metrics],
            "skipped_metrics": skipped,
        }
    if any(metric in {PGY_ALL_NON_LIVE_METRICS, "__ALL_NON_LIVE__"} for metric in metrics):
        result = _ensure_all_non_live_modal_options(page, modal)
        selected.extend(result["selected_metrics"])
        skipped.extend(result["skipped_metrics"])
    else:
        for metric in metrics:
            success, message = _ensure_modal_option(page, modal, metric)
            target = selected if success else skipped
            target.append({"metric": metric, "message": message})
    confirm = modal.locator("button").filter(has_text="确定").last
    try:
        if confirm.count():
            confirm.click(timeout=3000)
            page.wait_for_timeout(1500)
        else:
            page.keyboard.press("Escape")
    except Exception:
        try:
            page.keyboard.press("Escape")
        except Exception:
            pass
    return {"selected_metrics": selected, "skipped_metrics": skipped}


def apply_collection_plan(page: Any, plan: dict[str, Any], stop_callback: Any = None, ensure_metrics: bool = True) -> dict[str, Any]:
    _reset_kol_response_capture_state(page)
    applied: list[dict[str, str]] = []
    skipped: list[dict[str, str]] = []
    grouped_marketing_goals: dict[str, list[dict[str, Any]]] = {}
    regular_filters: list[dict[str, Any]] = []
    before_signature = _creator_table_signature(page)
    for item in plan.get("filters") or []:
        if not isinstance(item, dict):
            continue
        if str(item.get("field") or "") == "营销目标" or str(item.get("control_type") or "") == "marketing_goal_metric":
            parent = _marketing_goal_parent(item)
            grouped_marketing_goals.setdefault(parent, []).append(item)
        else:
            regular_filters.append(item)
    for parent, items in grouped_marketing_goals.items():
        if _stop_requested(stop_callback):
            return {"applied_filters": applied, "skipped_filters": skipped, "selected_metrics": [], "skipped_metrics": [], "stopped": True}
        group_applied, group_skipped = _apply_marketing_goal_group(page, parent, items)
        applied.extend(group_applied)
        skipped.extend(group_skipped)
    for item in regular_filters:
        if _stop_requested(stop_callback):
            return {"applied_filters": applied, "skipped_filters": skipped, "selected_metrics": [], "skipped_metrics": [], "stopped": True}
        success, message = _apply_filter_item(page, item)
        if success and _filter_acceptance_required(item) and not _verify_filter_selected(page, item):
            success = False
            message = f"{message}；但已选筛选栏未确认该条件，已阻止作为有效筛选"
        if success:
            applied.append({**item, "message": message})
        else:
            skipped.append({**item, "message": message})
    if _stop_requested(stop_callback):
        return {"applied_filters": applied, "skipped_filters": skipped, "selected_metrics": [], "skipped_metrics": [], "stopped": True}
    if applied:
        if not _wait_for_creator_table_change(page, before_signature, timeout_ms=12000):
            page.wait_for_timeout(2500)
    _reset_kol_response_capture_state(page)
    if _stop_requested(stop_callback):
        return {"applied_filters": applied, "skipped_filters": skipped, "selected_metrics": [], "skipped_metrics": [], "stopped": True}
    _prime_kol_api_capture(page, reload_if_empty=True, stop_callback=stop_callback)
    if _stop_requested(stop_callback):
        return {"applied_filters": applied, "skipped_filters": skipped, "selected_metrics": [], "skipped_metrics": [], "stopped": True}
    metric_result = _ensure_display_metrics(page, [str(item) for item in plan.get("display_metrics") or []]) if ensure_metrics else {"selected_metrics": [], "skipped_metrics": []}
    return {"applied_filters": applied, "skipped_filters": skipped, **metric_result}


def _read_visible_collection_state(page: Any) -> dict[str, Any]:
    rows = _creator_list_rows(page)
    try:
        rows.first.wait_for(state="visible", timeout=8000)
    except Exception:
        pass
    try:
        raw_count = rows.count()
    except Exception:
        raw_count = 0
    count = 0
    for index in range(min(raw_count, 30)):
        try:
            row_text = rows.nth(index).inner_text(timeout=800)
        except Exception:
            continue
        if _parse_row_text(row_text, page.url):
            count += 1
    return {
        "rows": rows,
        "count": count,
        "empty_hint": _page_empty_result_hint(page),
    }


def _click_first_visible_text(page: Any, scope_selector: str, text_pattern: re.Pattern[str]) -> bool:
    candidates = page.locator(scope_selector).filter(has_text=text_pattern)
    try:
        count = min(candidates.count(), 80)
    except Exception:
        return False
    for index in range(count):
        candidate = candidates.nth(index)
        try:
            if _is_visible(candidate) and _click_locator(page, candidate, timeout=1200):
                page.wait_for_timeout(700)
                return True
        except Exception:
            continue
    return False


def _clear_current_pgy_filters(page: Any) -> dict[str, Any]:
    clicked: list[str] = []
    before_signature = _creator_table_signature(page)
    clear_pattern = re.compile(r"^\s*(清空|重置|清除|清空筛选|重置筛选|清除筛选|全部清除)\s*$")
    for selector in [
        ".blogger-list_filter button",
        ".blogger-list_filter .d-button",
        ".blogger-list_filter [role=button]",
        ".blogger-list_filter *",
    ]:
        if _click_first_visible_text(page, selector, clear_pattern):
            clicked.append("clear_button")
            break

    close_selectors = [
        ".blogger-list_filter .d-tag .d-icon-close",
        ".blogger-list_filter .d-tag-close",
        ".blogger-list_filter .tag .close",
        ".blogger-list_filter [class*='close']",
    ]
    for selector in close_selectors:
        close_buttons = page.locator(selector)
        try:
            count = min(close_buttons.count(), 40)
        except Exception:
            continue
        for index in range(count):
            button = close_buttons.nth(index)
            try:
                if not _is_visible(button):
                    continue
                if _click_locator(page, button, timeout=1000):
                    clicked.append("filter_tag_close")
                    page.wait_for_timeout(250)
            except Exception:
                continue

    if clicked:
        _wait_for_creator_table_change(page, before_signature, timeout_ms=5000)
    return {
        "cleared_filters": bool(clicked),
        "clear_actions": clicked,
    }


def _apply_relaxed_plan_after_empty_result(page: Any, plan: dict[str, Any], plan_result: dict[str, Any], stop_callback: Any = None) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    current_plan = plan
    current_result = plan_result
    current_state = _read_visible_collection_state(page)
    if current_state["count"] > 0 or not current_state["empty_hint"]:
        return current_plan, current_result, current_state

    attempts = [("broad", False), ("unfiltered", True)]
    for stage, reset_filters in attempts:
        if _stop_requested(stop_callback):
            break
        relaxed_plan = _relaxed_collection_plan(plan, stage)
        if relaxed_plan is None:
            continue
        try:
            if reset_filters or "/solar/pre-trade/note/kol" not in page.url:
                page.goto(PGY_KOL_URL, wait_until="domcontentloaded", timeout=12000)
                page.wait_for_timeout(1500)
        except Exception:
            pass
        retry_result = apply_collection_plan(page, relaxed_plan, stop_callback=stop_callback) if relaxed_plan.get("filters") else {
            "applied_filters": [],
            "skipped_filters": [],
            "selected_metrics": [],
            "skipped_metrics": [],
        }
        if not relaxed_plan.get("filters"):
            metrics_result = _ensure_display_metrics(page, [str(item) for item in relaxed_plan.get("display_metrics") or []])
            retry_result = {**retry_result, **metrics_result}
        retry_state = _read_visible_collection_state(page)
        merged_result = _merge_plan_results(current_result, retry_result, relaxed_plan.get("relaxation") or {})
        if retry_state["count"] > 0 or not retry_state["empty_hint"]:
            return relaxed_plan, merged_result, retry_state
        current_plan = relaxed_plan
        current_result = merged_result
        current_state = retry_state
    return current_plan, current_result, current_state


def _extract_recommendation_count(page: Any) -> dict[str, Any]:
    try:
        text = page.locator("body").inner_text(timeout=3000)
    except Exception as error:
        return {"actual_recommend_count": None, "actual_count_text": "", "actual_count_error": f"读取推荐博主数失败：{error}"}
    patterns = [
        r"推荐\s*([0-9][0-9,，.]*)\s*(万|\+)?\s*位?博主",
        r"推荐博主\s*([0-9][0-9,，.]*)\s*(万|\+)?",
        r"共\s*([0-9][0-9,，.]*)\s*(万|\+)?\s*位?博主",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if not match:
            continue
        number_text = match.group(1).replace(",", "").replace("，", "")
        suffix = match.group(2) or ""
        try:
            value = float(number_text)
        except ValueError:
            continue
        if suffix == "万":
            value *= 10000
        return {
            "actual_recommend_count": int(value),
            "actual_count_text": match.group(0),
            "actual_count_is_lower_bound": suffix == "+",
        }
    return {"actual_recommend_count": None, "actual_count_text": "", "actual_count_error": "页面未识别到“推荐 N 位博主”文案"}


def _export_current_table(page: Any) -> dict[str, Any]:
    export_button = page.locator("button").filter(has_text=re.compile("导出|下载")).last
    if not export_button.count():
        return {"status": "skipped", "message": "页面未找到导出按钮"}
    export_dir = ROOT / "runtime" / "exports" / "pgy"
    export_dir.mkdir(parents=True, exist_ok=True)
    try:
        export_button.scroll_into_view_if_needed(timeout=2000)
        with page.expect_download(timeout=10000) as download_info:
            export_button.click(timeout=3000)
        download = download_info.value
        filename = download.suggested_filename or "pgy-export.csv"
        path = export_dir / filename
        download.save_as(str(path))
        parsed = parse_export_file(path)
        return {"status": "success", "message": "蒲公英列表已导出", "path": str(path), "parsed": parsed}
    except Exception as error:
        return {"status": "failed", "message": f"导出未完成：{error}"}


def browser_status() -> dict[str, Any]:
    try:
        response = requests.get(CDP_URL, timeout=2)
        if response.ok:
            payload = response.json()
            return {
                "connected": True,
                "message": "已连接 Chrome 调试端口",
                "webSocketDebuggerUrl": payload.get("webSocketDebuggerUrl"),
                "browser": payload.get("Browser"),
            }
    except requests.RequestException:
        pass
    return {
        "connected": False,
        "message": "未检测到 Chrome 调试端口，请先启动调试浏览器并登录蒲公英",
        "start_command": 'Start-Process "<chrome.exe>" -ArgumentList "--remote-debugging-port=9222","--user-data-dir=<project>\\runtime\\chrome-pgy-profile"',
    }


def _find_chrome_exe() -> Path | None:
    candidates = [
        Path("C:/Program Files/Google/Chrome/Application/chrome.exe"),
        Path("C:/Program Files (x86)/Google/Chrome/Application/chrome.exe"),
    ]
    local_app_data = Path.home() / "AppData/Local/Google/Chrome/Application/chrome.exe"
    candidates.append(local_app_data)
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return None


def start_browser() -> dict[str, Any]:
    status = browser_status()
    if status["connected"]:
        return status
    chrome = _find_chrome_exe()
    if not chrome:
        return {"connected": False, "message": "未找到 Chrome，请手动用调试端口启动浏览器"}
    profile = PGY_CHROME_PROFILE_DIR
    profile.mkdir(parents=True, exist_ok=True)
    _prepare_chrome_profile_for_clean_start(profile)
    subprocess.Popen(
        [
            str(chrome),
            "--remote-debugging-port=9222",
            f"--user-data-dir={profile}",
            "--no-first-run",
            "--no-default-browser-check",
            "--disable-session-crashed-bubble",
            "--hide-crash-restore-bubble",
            PGY_KOL_URL,
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        creationflags=subprocess.CREATE_NO_WINDOW,
    )
    return {**browser_status(), "message": "已尝试启动独立 Chrome，请在打开的页面登录蒲公英"}


def _prepare_chrome_profile_for_clean_start(profile: Path) -> None:
    profile.mkdir(parents=True, exist_ok=True)
    default_profile = profile / "Default"
    for name in ["Last Session", "Last Tabs", "Current Session", "Current Tabs"]:
        try:
            (default_profile / name).unlink(missing_ok=True)
        except Exception:
            pass
    preferences_path = default_profile / "Preferences"
    if not preferences_path.exists():
        return
    try:
        preferences = json.loads(preferences_path.read_text(encoding="utf-8"))
        if not isinstance(preferences, dict):
            return
        profile_prefs = preferences.setdefault("profile", {})
        if isinstance(profile_prefs, dict):
            profile_prefs["exit_type"] = "Normal"
            profile_prefs["exited_cleanly"] = True
        session_prefs = preferences.setdefault("session", {})
        if isinstance(session_prefs, dict):
            session_prefs["restore_on_startup"] = 0
        preferences_path.write_text(json.dumps(preferences, ensure_ascii=False), encoding="utf-8")
    except Exception:
        pass


def _select_pgy_list_page(context: Any, *, require_existing: bool = False) -> Any:
    pages = []
    for page in context.pages:
        try:
            if page.is_closed():
                continue
        except Exception:
            continue
        pages.append(page)
    list_page = next((item for item in pages if "/solar/pre-trade/note/kol" in (item.url or "")), None)
    if list_page is not None:
        try:
            list_page.bring_to_front()
        except Exception:
            pass
        return list_page
    if require_existing:
        return None
    pgy_page = next((item for item in pages if "pgy.xiaohongshu.com" in (item.url or "")), None)
    page = pgy_page or (pages[0] if pages else context.new_page())
    try:
        page.bring_to_front()
    except Exception:
        pass
    return page


def _is_recoverable_page_error(error: Any) -> bool:
    text = str(error or "").lower()
    return any(
        marker in text
        for marker in [
            "target closed",
            "page closed",
            "browser closed",
            "context closed",
            "crash",
            "crashed",
            "net::err_aborted",
            "execution context was destroyed",
        ]
    )


def _ensure_pgy_page_alive(context: Any, page: Any, *, preserve_existing_filters: bool = False) -> Any:
    try:
        if page is not None and not page.is_closed():
            return page
    except Exception:
        pass
    if preserve_existing_filters:
        return None
    page = context.new_page()
    page.goto(PGY_KOL_URL, wait_until="domcontentloaded", timeout=12000)
    page.wait_for_timeout(1500)
    return page


def collect_visible_list(
    brief: str = "",
    screening_plan: dict[str, Any] | None = None,
    apply_filters: bool = True,
    limit: int = 5000,
    include_details: bool = True,
    collect_profile_urls: bool = True,
    export_metrics: bool = True,
    reset_filters: bool = False,
    preflight_only: bool = False,
    kol_request_snapshot: dict[str, Any] | None = None,
    progress_callback: Any = None,
    stop_callback: Any = None,
    _recovery_attempt: int = 0,
) -> dict[str, Any]:
    if not browser_status()["connected"]:
        return {"ok": False, "message": "未连接 Chrome 调试端口"}
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        return {"ok": False, "message": "当前环境未安装 Playwright，无法执行真实页面采集"}

    def stopped_payload(page: Any | None = None, **extra: Any) -> dict[str, Any]:
        current_url = ""
        try:
            current_url = str(getattr(page, "url", "") or "") if page is not None else ""
        except Exception:
            current_url = ""
        return {
            "ok": False,
            "stopped": True,
            "message": "用户已停止当前采集",
            "current_url": current_url,
            "detail_collection": "stopped",
            "export_result": {"status": "skipped", "message": "用户已停止当前采集"},
            "creators": [],
            **extra,
        }

    if _stop_requested(stop_callback):
        return stopped_payload()

    try:
        with sync_playwright() as p:
            browser = p.chromium.connect_over_cdp("http://127.0.0.1:9222")
            context = browser.contexts[0] if browser.contexts else browser.new_context()
            preserve_existing_filters = bool(not apply_filters and not reset_filters)
            page = _select_pgy_list_page(context, require_existing=preserve_existing_filters)
            page = _ensure_pgy_page_alive(context, page, preserve_existing_filters=preserve_existing_filters)
            try:
                page.set_default_timeout(8000)
                page.set_default_navigation_timeout(12000)
            except Exception:
                pass
            if _stop_requested(stop_callback):
                return stopped_payload(page)
            if page is None:
                return {
                    "ok": False,
                    "message": "预检后的蒲公英列表页不存在，已停止以避免跳回空筛选列表误采泛达人池",
                    "current_url": "",
                }
            _install_kol_response_capture(page)
            if not preserve_existing_filters:
                _reset_kol_response_capture_state(page)
            restored_kol_snapshot = _restore_kol_request_snapshot(
                page,
                kol_request_snapshot,
                clear_pool=True,
            )
            if _stop_requested(stop_callback):
                return stopped_payload(page)
            if (
                "pgy.xiaohongshu.com" not in page.url
                or "/solar/pre-trade/note/kol" not in page.url
            ):
                if preserve_existing_filters:
                    return {
                        "ok": False,
                        "message": "当前页不是预检后的博主广场列表页，已停止以避免重置筛选条件",
                        "current_url": page.url,
                    }
                page.goto(PGY_KOL_URL, wait_until="domcontentloaded", timeout=12000)
                if _stop_requested(stop_callback):
                    return stopped_payload(page)
            if "/solar/pre-trade/note/kol" not in page.url:
                return {"ok": False, "message": "请先打开蒲公英博主广场 / 找博主页面", "current_url": page.url}
            page.wait_for_timeout(1500)
            if _stop_requested(stop_callback):
                return stopped_payload(page)
            login_hint = page.locator("text=登录").first
            if login_hint.count() and "login" in page.url.lower():
                return {"ok": False, "message": "蒲公英尚未登录，请在打开的 Chrome 页面完成登录", "current_url": page.url}
            if reset_filters and apply_filters:
                _clear_current_pgy_filters(page)
                if _stop_requested(stop_callback):
                    return stopped_payload(page)
            plan = build_collection_plan(brief, screening_plan)
            plan_result = (
                apply_collection_plan(page, plan, stop_callback=stop_callback, ensure_metrics=not preflight_only)
                if apply_filters
                else {"applied_filters": [], "skipped_filters": [], "selected_metrics": [], "skipped_metrics": []}
            )
            if plan_result.get("stopped") or _stop_requested(stop_callback):
                return stopped_payload(page, collection_plan=plan, **plan_result)
            allow_api_reload = _should_reload_for_api_prime(
                apply_filters=apply_filters,
                preserve_existing_filters=preserve_existing_filters,
                preflight_only=preflight_only,
            )
            _prime_kol_api_capture(page, reload_if_empty=allow_api_reload and not restored_kol_snapshot, stop_callback=stop_callback)
            if _stop_requested(stop_callback):
                return stopped_payload(page, collection_plan=plan, **plan_result)
            if (collect_profile_urls or include_details) and not getattr(page, "_pgy_latest_api_kols", []):
                _prime_kol_api_capture(page, reload_if_empty=allow_api_reload and not restored_kol_snapshot, stop_callback=stop_callback)
                if _stop_requested(stop_callback):
                    return stopped_payload(page, collection_plan=plan, **plan_result)
            active_plan = plan
            collection_state = _read_visible_collection_state(page)
            if apply_filters and collection_state["count"] == 0 and collection_state["empty_hint"]:
                active_plan, plan_result, collection_state = _apply_relaxed_plan_after_empty_result(page, plan, plan_result, stop_callback=stop_callback)
                if plan_result.get("stopped") or _stop_requested(stop_callback):
                    return stopped_payload(page, collection_plan=active_plan, **plan_result)
            recommendation_count = _extract_recommendation_count(page)
            request_snapshot = _current_kol_request_snapshot(page)
            request_validation = validate_kol_request_snapshot_against_plan(request_snapshot, active_plan)
            repair_result: dict[str, Any] = {}
            repaired_capture = False
            if apply_filters and request_validation.get("expected_constraints") and not request_validation.get("request_valid"):
                if _stop_requested(stop_callback):
                    return stopped_payload(page, collection_plan=active_plan, kol_request_snapshot=request_snapshot, kol_request_validation=request_validation, **plan_result, **recommendation_count)
                original_request_issues = request_validation.get("issues") or []
                repair_result = _repair_kol_capture_for_plan(page, request_snapshot, active_plan)
                if _stop_requested(stop_callback):
                    return stopped_payload(page, collection_plan=active_plan, kol_request_snapshot=request_snapshot, kol_request_validation=request_validation, kol_request_repair=repair_result, **plan_result, **recommendation_count)
                if not repair_result.get("ok"):
                    return {
                        "ok": False,
                        "message": "蒲公英页面筛选与真实 API 请求不一致，自动修复失败，已停止采集以避免误入库："
                        + "；".join(request_validation.get("issues") or [])
                        + f"；{repair_result.get('message') or ''}",
                        "current_url": page.url,
                        "collection_plan": active_plan,
                        "kol_request_snapshot": repair_result.get("kol_request_snapshot") or request_snapshot,
                        "kol_request_validation": repair_result.get("kol_request_validation") or request_validation,
                        "kol_request_repair": repair_result,
                        **plan_result,
                        **recommendation_count,
                        "detail_collection": "stopped_request_repair_failed",
                        "export_result": {"status": "skipped", "message": "真实请求自动修复失败，未导出列表"},
                        "creators": [],
                    }
                repaired_capture = True
                request_snapshot = repair_result.get("kol_request_snapshot") or _current_kol_request_snapshot(page)
                request_validation = repair_result.get("kol_request_validation") or validate_kol_request_snapshot_against_plan(request_snapshot, active_plan, require_clean_pool=False)
                total_count = request_snapshot.get("total_count") if isinstance(request_snapshot, dict) else None
                if total_count not in (None, ""):
                    recommendation_count = {
                        "actual_recommend_count": int(total_count),
                        "actual_count_text": f"修复后 API 推荐 {int(total_count)} 位博主",
                        "actual_count_is_lower_bound": False,
                    }
                elif repair_result.get("first_page_count") is not None:
                    recommendation_count = {
                        "actual_recommend_count": int(repair_result.get("first_page_count") or 0),
                        "actual_count_text": f"修复后 API 首页返回 {int(repair_result.get('first_page_count') or 0)} 位博主",
                        "actual_count_is_lower_bound": True,
                    }
                skipped = plan_result.get("skipped_filters") if isinstance(plan_result.get("skipped_filters"), list) else []
                plan_result["skipped_filters"] = [
                    *skipped,
                    {
                        "field": "真实请求自动修复",
                        "value": "蒲公英达人列表 API",
                        "message": repair_result.get("message") or "已重组请求并继续采集",
                        "issues_before_repair": original_request_issues,
                    },
                ]
            if preflight_only:
                return {
                    "ok": True,
                    "current_url": page.url,
                    "collection_plan": active_plan,
                    "kol_request_snapshot": request_snapshot,
                    "kol_request_validation": request_validation,
                    "kol_request_repair": repair_result,
                    **plan_result,
                    **recommendation_count,
                    "detail_collection": "preflight_only",
                    "export_result": {"status": "skipped", "message": "预检阶段不导出列表"},
                    "creators": [],
                }
            export_result = (
                {"status": "skipped", "message": "已自动修复为 API 采集，页面导出不代表修复后请求，跳过导出"}
                if repaired_capture
                else _export_current_table(page)
                if export_metrics
                else {"status": "skipped", "message": "本次未请求导出"}
            )
            if _stop_requested(stop_callback):
                return stopped_payload(page, collection_plan=active_plan, export_result=export_result, kol_request_snapshot=request_snapshot, kol_request_validation=request_validation, kol_request_repair=repair_result, **plan_result, **recommendation_count)
            rows = collection_state["rows"]
            count = collection_state["count"]
            if count == 0 and not repaired_capture:
                hint = collection_state.get("empty_hint")
                message = (
                    f"蒲公英页面提示「{hint}」，系统已尝试自动放宽筛选但仍无可采集列表"
                    if hint
                    else "未识别到博主列表，请确认已登录并停留在博主广场列表页"
                )
                return {
                    "ok": False,
                    "message": message,
                    "current_url": page.url,
                    "collection_plan": active_plan,
                    "export_result": export_result,
                    "kol_request_snapshot": request_snapshot,
                    "kol_request_validation": request_validation,
                    **plan_result,
                    **recommendation_count,
                }
            collect_limit = max(1, min(limit, 5000))
            creators = _extract_visible_creators(
                page,
                collect_limit,
                include_details=include_details and not repaired_capture,
                collect_profile_urls=collect_profile_urls,
                detail_limit=collect_limit if include_details and not repaired_capture else 0,
                api_first=repaired_capture or not include_details,
                progress_callback=progress_callback,
                stop_callback=stop_callback,
            )
            if _stop_requested(stop_callback):
                return stopped_payload(page, collection_plan=active_plan, export_result=export_result, kol_request_snapshot=request_snapshot, kol_request_validation=request_validation, kol_request_repair=repair_result, **plan_result, **recommendation_count)
            if not creators:
                return {
                    "ok": False,
                    "message": "已连接蒲公英页面，但未读取到有效达人内容。请确认列表加载完成后再采集，必要时滚动列表或刷新页面。",
                    "current_url": page.url,
                    "collection_plan": active_plan,
                    "export_result": export_result,
                    "kol_request_snapshot": request_snapshot,
                    "kol_request_validation": request_validation,
                    "kol_request_repair": repair_result,
                    **plan_result,
                    **recommendation_count,
                }
            return {
                "ok": True,
                "creators": creators,
                "current_url": page.url,
                "collection_plan": active_plan,
                "detail_collection": "repaired_api" if repaired_capture else ("planned" if include_details else "skipped"),
                "export_result": export_result,
                "kol_request_snapshot": request_snapshot,
                "kol_request_validation": request_validation,
                "kol_request_repair": repair_result,
                **recommendation_count,
                **plan_result,
            }
    except Exception as error:
        if _is_recoverable_page_error(error) and _recovery_attempt < 1 and not (not apply_filters and not reset_filters):
            return collect_visible_list(
                brief=brief,
                screening_plan=screening_plan,
                apply_filters=apply_filters,
                limit=limit,
                include_details=include_details,
                collect_profile_urls=collect_profile_urls,
                export_metrics=export_metrics,
                reset_filters=True,
                preflight_only=preflight_only,
                kol_request_snapshot=kol_request_snapshot,
                progress_callback=progress_callback,
                stop_callback=stop_callback,
                _recovery_attempt=_recovery_attempt + 1,
            )
        return {"ok": False, "message": f"蒲公英采集失败：{error}"}


def collect_details_for_targets(targets: list[dict[str, Any]], limit: int = 500, progress_callback: Any = None) -> dict[str, Any]:
    if not targets:
        return {"ok": False, "message": "没有可补全详情页的初筛通过达人", "creators": []}
    if not browser_status()["connected"]:
        return {"ok": False, "message": "未连接 Chrome 调试端口"}
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        return {"ok": False, "message": "当前环境未安装 Playwright，无法执行真实页面采集"}

    target_limit = max(1, min(limit, 500))
    pending = targets[:target_limit]
    by_nickname = {str(item.get("nickname") or "").strip(): item for item in pending if item.get("nickname")}
    completed: list[dict[str, Any]] = []
    failed: list[dict[str, Any]] = []
    started_at = _now_text()
    started_perf = time.perf_counter()

    def attach_timing(creator: dict[str, Any], timing: dict[str, Any]) -> dict[str, Any]:
        raw_payload = creator.get("raw_payload") if isinstance(creator.get("raw_payload"), dict) else {}
        creator["raw_payload"] = {**raw_payload, "detail_collection_timing": timing}
        creator["detail_collection_timing"] = timing
        return creator

    def finish_payload(ok: bool, message: str) -> dict[str, Any]:
        timing = _elapsed_timing(started_at, started_perf)
        average_seconds = round(timing["duration_seconds"] / max(1, len(completed)), 2) if completed else 0
        return {
            "ok": ok,
            "creators": completed,
            "failed": failed,
            "message": message,
            "started_at": timing["started_at"],
            "finished_at": timing["finished_at"],
            "duration_seconds": timing["duration_seconds"],
            "duration_text": timing["duration_text"],
            "average_seconds_per_creator": average_seconds,
            "average_duration_text": _duration_text(average_seconds) if average_seconds else "",
        }

    def emit_progress(stage: str, current: dict[str, Any] | None = None) -> None:
        if not progress_callback:
            return
        try:
            progress_callback(
                {
                    "stage": stage,
                    "total_count": len(pending),
                    "completed_count": len(completed),
                    "failed_count": len(failed),
                    "current_creator_id": (current or {}).get("creator_id") or "",
                    "current_nickname": (current or {}).get("nickname") or "",
                }
            )
        except Exception:
            pass

    try:
        emit_progress("starting")
        url_targets = [target for target in pending if _target_detail_url(target)]
        fallback_targets = [target for target in pending if not _target_detail_url(target)]
        matched_ids: set[str] = set()

        if url_targets:
            max_workers = min(PGY_DETAIL_COLLECT_CONCURRENCY, len(url_targets))
            indexed_results: dict[int, dict[str, Any]] = {}
            with ThreadPoolExecutor(max_workers=max_workers) as executor:
                future_map = {
                    executor.submit(_collect_detail_target_by_url, target): index
                    for index, target in enumerate(url_targets)
                }
                for future in as_completed(future_map):
                    indexed_results[future_map[future]] = future.result()
            for index in sorted(indexed_results):
                result = indexed_results[index]
                if result.get("ok") and isinstance(result.get("creator"), dict):
                    creator = attach_timing(result["creator"], result.get("detail_collection_timing") or {})
                    completed.append(creator)
                    matched_ids.add(str(creator.get("creator_id") or ""))
                    emit_progress("collecting", creator)
                else:
                    failed.append({
                        "creator_id": result.get("creator_id"),
                        "nickname": result.get("nickname"),
                        "message": result.get("message") or "详情页解析失败",
                        "detail_collection_timing": result.get("detail_collection_timing"),
                    })
                    emit_progress("collecting", result)

        if fallback_targets:
            by_nickname = {
                str(item.get("nickname") or "").strip(): item
                for item in fallback_targets
                if item.get("nickname")
            }
            with sync_playwright() as p:
                browser = p.chromium.connect_over_cdp("http://127.0.0.1:9222")
                context = browser.contexts[0] if browser.contexts else browser.new_context()
                page = _select_pgy_list_page(context)
                if "pgy.xiaohongshu.com" not in page.url:
                    page.goto(PGY_KOL_URL, wait_until="domcontentloaded", timeout=30000)
                _install_kol_response_capture(page)
                _reset_kol_response_capture_state(page)
                page.wait_for_timeout(300)
                _prime_kol_api_capture(page, reload_if_empty=False)
                rows = page.locator(".blogger-list_list .d-new-table tbody tr").filter(has_not=page.locator(".skeleton-block"))
                try:
                    row_count = min(rows.count(), 300)
                except Exception:
                    row_count = 0
                for index in range(row_count):
                    row = rows.nth(index)
                    creator = _parse_row_text(row.inner_text(timeout=2500), page.url)
                    if not creator:
                        continue
                    target = by_nickname.get(str(creator.get("nickname") or "").strip())
                    if not target:
                        continue
                    api_kol = _match_api_kol(page, index, creator)
                    api_recent_note_briefs = _recent_note_briefs_from_kol(api_kol, max_notes=2) if api_kol else []
                    api_recent_notes = _collect_recent_note_details_from_api(page, api_kol, max_notes=2) if api_kol else []
                    link_fields = _extract_row_link_fields(row) or _api_kol_link_fields(page, index, creator)
                    detail_url = str(link_fields.get("pgy_url") or link_fields.get("profile_url") or "")
                    item_started_at = _now_text()
                    item_started_perf = time.perf_counter()
                    detail = _collect_first_detail_for_creator(context, row, detail_url)
                    item_timing = _elapsed_timing(item_started_at, item_started_perf)
                    if detail:
                        merged = {**target, **link_fields, **detail, "creator_id": target.get("creator_id")}
                        raw_payload = merged.get("raw_payload") if isinstance(merged.get("raw_payload"), dict) else {}
                        if api_recent_note_briefs:
                            raw_payload["recent_note_briefs"] = api_recent_note_briefs
                        if api_recent_notes:
                            raw_payload["recent_notes"] = api_recent_notes
                        if raw_payload:
                            merged["raw_payload"] = raw_payload
                        attach_timing(merged, item_timing)
                        completed.append(merged)
                        matched_ids.add(str(target.get("creator_id")))
                        emit_progress("collecting", merged)
                    else:
                        failed.append({
                            "creator_id": target.get("creator_id"),
                            "nickname": target.get("nickname"),
                            "message": "详情页打开失败",
                            "detail_collection_timing": item_timing,
                        })
                        emit_progress("collecting", target)

            for target in fallback_targets:
                creator_id = str(target.get("creator_id") or "")
                if creator_id in matched_ids:
                    continue
                failed.append({
                    "creator_id": creator_id,
                    "nickname": target.get("nickname"),
                    "message": "当前列表未找到该达人，且本地没有可打开的详情页链接",
                })
                emit_progress("collecting", target)
        payload = finish_payload(True, "")
        emit_progress("finished")
        payload["message"] = (
            f"详情页补全完成 {len(completed)} 个，失败 {len(failed)} 个"
            f"；总耗时 {payload['duration_text']}"
            f"{f'，单个约 {payload['average_duration_text']}' if payload.get('average_duration_text') else ''}"
        )
        return payload
    except Exception as error:
        return finish_payload(False, f"详情页补全失败：{error}")
